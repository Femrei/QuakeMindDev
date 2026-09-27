# -*- coding: utf-8 -*-
"""Revizyon 9: son okuma (sayfa duzeni).

1. Tablo hucreleri Normal stilinden IKI YANA YASLI hizayi miras aliyordu; dar
   sutunlarda sozcukler arasi bosluk aciliyordu ("Defne   /   Elektrik",
   Tablo 15, 17, 21). Izgara tablolarinda hucre paragraflari sola yaslandi.
2. Sekil 11'in gorseli bir sayfada, altyazisi sonraki sayfadaydi. Altyazisi
   olan her gorsel paragrafina ve her tablo basligina "sonrakiyle birlikte
   tut" verildi.
3. Tablo 21'de ROC-AUC sutunu ("0,6370 ± 0,0637") iki satira bolunuyordu;
   sutun genislikleri icerige gore yeniden dagitildi (toplam ayni, 16,4 cm).
"""
import re
import shutil
import sys

sys.path.insert(0, r"C:\Users\utkuc\AppData\Local\Temp\claude\D--quakemind-final-QuakeMind\4779ac37-74b5-4aeb-8690-74d7792eedc5\scratchpad")
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph
from yardimci import Duzenleyici

SRC = r"D:\quakemind_final\QuakeMind\makale\arsiv\QuakeMind_Makale_YEDEK_pre_revize9_2026-09-27.docx"
DST = r"D:\quakemind_final\QuakeMind\makale\QuakeMind_Makale_Taslak_2026-09-16_sekillieklendi_v3_LOCAL_2026-09-21.docx"
shutil.copyfile(SRC, DST)
doc = Document(DST)
D = Duzenleyici(doc)

# 1. izgara tablolarinda sola yaslama (acikca ortalanmis hucrelere dokunulmaz)
n = 0
for t in doc.tables:
    if t.style is None or t.style.name != "Table Grid":
        continue
    for tr in t._tbl.iter(qn("w:tr")):
        for p_el in tr.iter(qn("w:p")):
            p = Paragraph(p_el, doc)
            if p.alignment is None:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                n += 1
D.log.append(f"HIZA {n} hucre paragrafi sola")

# 2. gorsel + altyazi / tablo basligi + tablo birlikte
body = doc.element.body
g = tb = 0


def _sonraki_altyazi_mi(el):
    nxt = el.getnext()
    while nxt is not None and nxt.tag == qn("w:p") and not Paragraph(nxt, doc).text.strip() \
            and not nxt.findall(".//" + qn("w:drawing")):
        nxt = nxt.getnext()
    return nxt is not None and nxt.tag == qn("w:p") and bool(re.match(r"^Şekil \d+\.", Paragraph(nxt, doc).text))


from docx.oxml import OxmlElement
for el in body.iterchildren():
    # yan yana panelli sekiller gorunmez bir duzen tablosunda (Sekil 11 gibi):
    # tablonun butun paragraflari sonrakiyle tutulur, satirlar bolunmez
    if el.tag == qn("w:tbl") and el.findall(".//" + qn("w:drawing")) and _sonraki_altyazi_mi(el):
        for p_el in el.iter(qn("w:p")):
            Paragraph(p_el, doc).paragraph_format.keep_with_next = True
        for tr in el.iter(qn("w:tr")):
            trPr = tr.get_or_add_trPr()
            if trPr.find(qn("w:cantSplit")) is None:
                trPr.append(OxmlElement("w:cantSplit"))
        g += 1
        continue
    if el.tag != qn("w:p"):
        continue
    p = Paragraph(el, doc)
    if re.match(r"^Tablo \d+\.", p.text):
        p.paragraph_format.keep_with_next = True
        tb += 1
        continue
    if el.findall(".//" + qn("w:drawing")):
        nxt = el.getnext()
        while nxt is not None and nxt.tag == qn("w:p") and not Paragraph(nxt, doc).text.strip() \
                and not nxt.findall(".//" + qn("w:drawing")):
            nxt = nxt.getnext()
        if nxt is not None and nxt.tag == qn("w:p") and re.match(r"^Şekil \d+\.", Paragraph(nxt, doc).text):
            p.paragraph_format.keep_with_next = True
            g += 1
D.log.append(f"BIRLIKTE TUT {g} gorsel, {tb} tablo basligi")

# 2b. tablolar: satir sayfa arasinda bolunmez, baslik satiri yinelenir; kisa
# tablolar (<= 15 satir) tek parca tutulur (Sekil 11 duzeltmesi Tablo 17'yi
# ikiye bolmustu)
kisa = 0
for t in doc.tables:
    if t.style is None or t.style.name != "Table Grid":
        continue
    satirlar = t._tbl.findall(qn("w:tr"))
    for k, tr in enumerate(satirlar):
        trPr = tr.get_or_add_trPr()
        if trPr.find(qn("w:cantSplit")) is None:
            trPr.append(OxmlElement("w:cantSplit"))
        if k == 0 and trPr.find(qn("w:tblHeader")) is None:
            trPr.append(OxmlElement("w:tblHeader"))
    if len(satirlar) <= 15:
        for tr in satirlar[:-1]:
            for p_el in tr.iter(qn("w:p")):
                Paragraph(p_el, doc).paragraph_format.keep_with_next = True
        kisa += 1
D.log.append(f"TABLO satir bolunmez + baslik yinelenir; {kisa} kisa tablo tek parca")

# 3. Tablo 21 sutunlari
t21 = [t for t in doc.tables if t.rows[0].cells[0].text == "Model" and "ROC-AUC" in t.rows[0].cells[1].text][0]
genislik = [3.0, 2.7, 1.6, 1.6, 1.9, 1.8, 1.6, 2.2]
assert len(t21.columns) == len(genislik) and abs(sum(genislik) - 16.4) < 1e-9
for i, w in enumerate(genislik):
    t21.columns[i].width = Cm(w)
    for c in t21.columns[i].cells:
        c.width = Cm(w)
D.log.append("TABLO 21 sutun genislikleri")

doc.save(DST)
print("\n".join(D.log))
