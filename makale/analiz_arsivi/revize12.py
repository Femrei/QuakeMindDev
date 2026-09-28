# -*- coding: utf-8 -*-
"""Revizyon 12: "olculmemis ve kapatilabilir" eksiklerin kapatilmasi (28 Eylul 2026).

Kaynaklar (hepsi QuakeMindBackend/benchmark/reports/ ya da ekip kayitlari):
  - A* capraz dogrulama: tatbikat_b0/astar_capraz.json (araclar_astar_capraz.py)
  - Sekiz alan: tatbikat_b0/sekiz_alan_*.json (araclar_tatbikat_kos.py)
  - Metin ust kategori + J-123: metin_p5_olcumu.json (araclar_metin_p5_olc.py)
  - YOLO: apps/camera_detection/models/{bina,catlak}.pt icindeki train_metrics
  - Son isleme seviyeleri: Kaggle quakemind-rapor-paketi kosu_kayitlari/kaggle/P2/*/tohum_42.json
  - Kapanma etiketleyicisi: benchmark/kapanma_yer_gercegi/etiketler_defne_elektrik.json ("Claude")
Sekil 24 (orneK veri kipinde cekilmis arayuz goruntusu) kaldirildi; modulun gercek
ciktilari Sekil 5 ve 10'da.
"""
import json
import shutil
import sys

sys.path.insert(0, r"C:\Users\utkuc\AppData\Local\Temp\claude\D--quakemind-final-QuakeMind\4779ac37-74b5-4aeb-8690-74d7792eedc5\scratchpad")
sys.path.insert(0, r"D:\quakemind_final\QuakeMind\QuakeMindBackend")
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from yardimci import Duzenleyici
from araclar_tatbikat_ozet import alan_ozeti

KOK = r"D:\quakemind_final\QuakeMind\QuakeMindBackend"
REP = KOK + r"\benchmark\reports"
SRC = r"D:\quakemind_final\QuakeMind\makale\arsiv\QuakeMind_Makale_YEDEK_pre_revize12_2026-09-28.docx"
DST = r"D:\quakemind_final\QuakeMind\makale\QuakeMind_Makale_Taslak_2026-09-16_sekillieklendi_v3_LOCAL_2026-09-21.docx"
v = lambda x, b=3: f"{x:.{b}f}".replace(".", ",")

# ---------------------------------------------------------------- dogrulamalar
A = json.load(open(REP + r"\tatbikat_b0\astar_capraz.json", encoding="utf-8"))
sorgu = sum(x["cift"] for x in A)
ayni = sum(x["ayni"] for x in A)
ulas = sum(x["ulasilamazUyumlu"] for x in A)
assert all(not x["uyumsuz"] for x in A) and (sorgu, ayni, ulas) == (248, 240, 8)
assert max(x["enBuyukFarkSn"] for x in A) < 1e-13
assert {x["vmaxKmh"] for x in A} == {45.0, 70.0} and all(x["eskiSezgiselFarkli"] == 0 for x in A)

sekiz = {}
for m in ("segformer_mitb0", "segformer_mitb2", "unet_mitb4"):
    for s in json.load(open(REP + rf"\tatbikat_b0\sekiz_alan_{m}.json", encoding="utf-8")):
        if s.get("hata"):
            continue
        r = alan_ozeti(s)
        if r["gecerliTohum"] and r["kazanc"][0] and r["kazanc"][0] > 0.05:
            sekiz[(s["alan"], m)] = r
assert set(sekiz) == {("Merkez / Şubat Stadyumu çevresi", "segformer_mitb2"),
                      ("Şekeroba (orta)", "segformer_mitb2"),
                      ("Antakya / Türünçlü – Akdeniz – Armutlu", "unet_mitb4")}, sekiz.keys()
st = sekiz[("Merkez / Şubat Stadyumu çevresi", "segformer_mitb2")]["kazanc"]
se = sekiz[("Şekeroba (orta)", "segformer_mitb2")]["kazanc"]
tu = sekiz[("Antakya / Türünçlü – Akdeniz – Armutlu", "unet_mitb4")]
assert (v(st[0], 1), v(st[1], 1), v(se[0], 1), v(se[1], 1), v(tu["kazanc"][0], 1), tu["ortakGorev"][0]) == \
       ("40,4", "7,6", "13,7", "12,5", "32,5", 1.0)

M = json.load(open(REP + r"\metin_p5_olcumu.json", encoding="utf-8"))
b = M["modeller"]["berturk"]["sinama"]
k2 = M["modeller"]["2kveri"]
assert (b["yeni"]["mikroF1"], b["yeni"]["makroF1"], b["yeni"]["dogruluk"]) == (0.797, 0.7583, 0.882)
assert (b["model"]["mikroF1"], b["model"]["makroF1"], b["eski"]["mikroF1"], b["eski"]["makroF1"]) == \
       (0.8189, 0.7874, 0.7219, 0.7098)
acil = b["model"]["siniflar"]["Acil Yardım İhtiyacı"]["duyarlilik"]
alak = (b["model"]["siniflar"]["Alakasız"]["duyarlilik"], b["yeni"]["siniflar"]["Alakasız"]["duyarlilik"])
assert (round(acil, 2), round(alak[0], 2), round(alak[1], 2)) == (0.35, 0.80, 0.67)
assert (round(k2["sinama"]["altEtiketMikroF1"], 3), round(k2["egitim"]["altEtiketMikroF1"], 3)) == (0.916, 0.912)
J = M["j123"]
assert (J["toplamKayit"], J["lojistikEtiketli"], J["ulasimSozcukluKayit"], J["lojistikEtiketliVeUlasimSozcuklu"]) == \
       (2028, 38, 69, 3)

E = json.load(open(KOK + r"\benchmark\kapanma_yer_gercegi\etiketler_defne_elektrik.json", encoding="utf-8"))
assert E["etiketleyen"].startswith("Claude")

# ---------------------------------------------------------------- belge
shutil.copyfile(SRC, DST)
doc = Document(DST)
D = Duzenleyici(doc)
R = D.degistir
YZ = "bir büyük dil modeli (Claude, Anthropic)"

# OZ / Abstract: kapanma etiketleyicisi
R("Kapanma katmanının doğruluğu bir mahallede kör etiketlemeyle ayrıca sınanmış, karar verilebilir parçalarda",
  "Kapanma katmanının doğruluğu bir mahallede kör etiketlemeyle ayrıca sınanmış (etiketleme bir büyük dil "
  "modeliyle yapılmıştır; insan doğrulaması bekleniyor), karar verilebilir parçalarda")
R("Blockage-layer accuracy was separately probed by blind labeling in one neighborhood, yielding",
  "Blockage-layer accuracy was separately probed by blind labeling in one neighborhood (performed by a large "
  "language model; human verification is pending), yielding")

# 3.4 son isleme seviyesi
R("Seviyenin sonuç ölçütleri üzerindeki etkisi bu sürümde ölçülmemiştir.",
  "Seviyenin etkisi, Tablo 12'deki on modelin eğitim koşularında 91 karoluk sınama kümesinde ölçülmüştür: "
  "seviye 1 ve seviye 2, IoU_hasar'ı en çok 0,0003 ve 0,0009 değiştirmekte, kesinliği binde birkaç artırıp "
  "duyarlılığı binde birkaç azaltmaktadır; son işlemenin bu kümede başarıma pratik bir etkisi yoktur.")

# 3.5 YOLO
D.yaz("Çatlak modeli Roboflow Universe üzerinde barındırılan",
      "Her iki model de Roboflow Universe üzerinde barındırılan ve yalnızca sınırlayıcı kutu (bounding box) "
      "etiketi içeren kümelerle eğitilmiştir; model dosyalarında veri klasörleri bina modeli için 'maras-2', "
      "çatlak modeli için 'crack-2' olarak kayıtlıdır. Aynı dosyalardaki eğitim bilgisine göre iki model de "
      "YOLOv8s ön eğitimli ağırlığından 640 piksellik girdiyle eğitilmiştir (Ultralytics 8.2.103): bina modeli "
      "40 devirlik planla ve 5 devir sabırlı erken durdurmayla 21. devirde, çatlak modeli 150 devirde durmuştur. "
      "Eğitim sırasında ayrılan doğrulama bölmesinde ölçülen değerler bina modeli için mAP@0,5 = 0,592 ve "
      "mAP@0,5:0,95 = 0,394 (kesinlik 0,596, duyarlılık 0,585), çatlak modeli için mAP@0,5 = 0,731 ve "
      "mAP@0,5:0,95 = 0,584'tür (kesinlik 0,831, duyarlılık 0,647). Bu bölme en iyi ağırlığın seçiminde de "
      "kullanıldığından değerler iyimser bir üst kestirimdir ve bağımsız bir saha sınamasının yerine geçmez; "
      "bölmelerin görüntü sayıları dosyalarda kayıtlı değildir (Tablo 22).")

# 3.6 metin kumesi: kisisel veri
R("bu iki alan açık bir eksiklik olarak kaydedilmektedir.",
  "bu iki alan açık bir eksiklik olarak kaydedilmektedir. Bu çalışmadaki denetimde, 'anonymized' adına karşın "
  "kümede kişi adlarının maskelendiği ancak telefon numarası ve açık adres içeren kayıtların bulunduğu "
  "görülmüştür.")

# 4.1 olculen/olculmeyen
R("kapanma katmanının bir mahallede kör etiketlemeyle ön doğrulaması (Bölüm 4.3, Tablo 13)",
  "kapanma katmanının bir mahallede kör etiketlemeyle (yapay zekâ etiketleyici) ön doğrulaması (Bölüm 4.3, "
  "Tablo 13)")

# 4.3 etiketleyici aciklamasi
R("yol parçaları, modelin ne dediği görülmeden (kör) gerçek ortofoto üzerinden elle işaretlenmiştir.",
  "yol parçaları, modelin ne dediği görülmeden (kör) gerçek ortofoto kesitleri üzerinden işaretlenmiştir. "
  f"Etiketleme bu sürümde {YZ} tarafından, kesit görüntüleri üzerinden yapılmıştır; aynı kesitlerin bir insan "
  "etiketleyici tarafından etiketlenmesi ve etiketleyiciler arası uyumun raporlanması bekleyen bir iştir "
  "(etiketleme kiti ve uyum betiği depodadır).")
R("(Hatay / Defne – Elektrik Mah., U-Net MiT-B4; kör etiketleme, ön deneme).",
  "(Hatay / Defne – Elektrik Mah., U-Net MiT-B4; kör etiketleme, yapay zekâ etiketleyici, ön deneme).")
R("Sınırlar: tek etiketleyici, ikinci görüş yok;",
  "Sınırlar: tek etiketleyici ve bu etiketleyici bir yapay zekâ modelidir, insan görüşü yok;")

# 4.5 A* capraz dogrulama
R("Bu, ayrı ve bağımsız bir yol bulucuyla çapraz doğrulanmamıştır; önceki sürümde anılan '90 görevde en fazla "
  "0,00 m fark' denetimi bu sürümde yinelenmemiştir.",
  "A* aramasının kendisi bağımsız bir çözücüyle çapraz doğrulanmıştır: üç alanın beş tohumundaki bütün ekip–"
  "afetzede çiftlerinde, kapanmasız ve kapanmalı iki görünümde (248 ayrı sorgu) A*'ın bulduğu süre, aynı maliyet "
  "modeliyle açıkça kurulan kenar grafiği üzerinde networkx Dijkstra'sının bulduğu süreyle birebir aynıdır (en "
  "büyük fark 5,7·10⁻¹⁴ sn; 8 sorguda ikisi de hedefe yol bulamamıştır). Tatbikat mahallelerinde en yüksek kenar "
  "hızı 45–70 km/sa olduğundan taslağın önceki hâlindeki sabit 80 km/sa sezgiseli bu alanlarda da kabul edilebilir "
  "kalmış ve sonuçları etkilememiştir.")

# 4.7 kisitlar: sekiz alan
R("Üç alan on bir ölçülen alandan seçilmiştir: kalan sekizinde kollar ayrışmamıştır, çünkü model ya kapanma "
  "bulmamış ya da bulduğu kapanmalar hiçbir rotanın üzerine düşmemiştir. Bu seçim bir kiraz toplama olarak "
  "okunmamalı, ayrışmanın hangi koşulda ortaya çıktığının bir bulgusu olarak alınmalıdır.",
  "Üç alan on bir ölçülen alandan, kolların en az iki modelde ayrıştığı alanlar olarak seçilmiştir. Kalan sekiz "
  "alan da düzeltilmiş kodla üç modelin üçüyle ve beş tohumla koşulmuştur: beşinde hiçbir modelde ayrışma yoktur "
  "(model ya kapanma bulmamış ya da bulduğu kapanmalar hiçbir rotaya düşmemiştir); üçünde ayrışma yalnız tek bir "
  f"modelde görülmüştür: Kahramanmaraş Şubat Stadyumu çevresinde MiT-B2 ile %{v(st[0], 1)} ± {v(st[1], 1)}, "
  f"Kahramanmaraş Şekeroba'nın orta kesiminde MiT-B2 ile %{v(se[0], 1)} ± {v(se[1], 1)} ve Hatay Türünçlü'de "
  f"U-Net MiT-B4 ile tek ortak görevde %{v(tu['kazanc'][0], 1)}. Tek modelde ayrışan alanların kazancı raporlanan "
  "bandın (%8–60) içinde kalmaktadır; seçim bir kiraz toplama değil, ayrışmanın hangi koşulda (kapanmanın hangi "
  "modelde rotaya düştüğü) ortaya çıktığının bir bulgusudur.")

# 4.10 kisit (i) ve ust kategori paragrafi
R("(i) etiket uzayı beş üst kategori değil alt kategori kümesidir, dolayısıyla beş üst kategori üzerindeki ölçüm "
  "(Tablo 22, üçüncü satır) bu tabloyla karşılanmamaktadır;",
  "(i) etiket uzayı alt kategori kümesidir; üst kategori düzeyindeki ölçüm aşağıda ayrıca verilmiştir;")
kisit = D.bul("Karşılaştırma üç noktada sınırlıdır: (i) etiket uzayı")
D.sonrasina_ekle(
    kisit._p,
    "Üst kategori düzeyi. BERTurk karşılaştırma ağırlığı (yalnız eğitim bölmesiyle eğitilmiş) aynı 305 kayıtlık "
    "sınama kümesinde, üretimdeki üst kategori katmanından geçirilerek ayrıca ölçülmüştür. Altın alt etiketler "
    "Tablo 8 eşlemesiyle üst kategorilere indirilmiş; 'Lojistik' alt etiketinin iki üst kategoriye ayrılması "
    "metindeki anahtar sözcüğe dayandığından ve altında bu bilgi bulunmadığından iki lojistik kategorisi "
    "değerlendirmede birleştirilmiştir (dört sınıf; tahmin tek kategori, altın çok etiketli, 'Belirsiz' çekimser "
    f"sayılır). Düzeltilmiş katmanla mikro-F₁ {v(b['yeni']['mikroF1'])}, makro-F₁ {v(b['yeni']['makroF1'])} ve "
    f"tahminin altın kategoriler arasında bulunma oranı {v(b['yeni']['dogruluk'])}'dir; katmansız eşleme "
    f"{v(b['model']['mikroF1'])} / {v(b['model']['makroF1'])}, katmanın 27 Eylül öncesi hâli "
    f"{v(b['eski']['mikroF1'])} / {v(b['eski']['makroF1'])} vermektedir. Acil Yardım İhtiyacı sınıfında duyarlılık "
    f"{v(acil, 2)}'te kalmaktadır: altını birden çok kategori taşıyan kayıtlarda tek tahmin çoğunlukla Enkaz "
    "Bildirimi'ne düşmektedir. Anahtar sözcük kurtarması bu kümede hayati sınıfların duyarlılığını artırmamış, "
    f"'Alakasız' sınıfının duyarlılığını {v(alak[0], 2)}'den {v(alak[1], 2)}'ye düşürmüştür. Üretimde ve tatbikatta "
    "koşan ağırlık sınama bölmesinde eğitim bölmesindekinden daha yüksek başarım verdiğinden (13 alt etiket "
    f"mikro-F₁ {v(k2['sinama']['altEtiketMikroF1'])} ve {v(k2['egitim']['altEtiketMikroF1'])}) sınama kayıtlarını "
    "eğitimde görmüş olduğu düşünülmektedir; bu nedenle başarım sayıları karşılaştırma ağırlığıyla verilmiştir.",
    kisit, kalin_bas="Üst kategori düzeyi.")

# 5.1 / 5.3
R("Kapanma katmanının doğruluğu, bir mahallede kör etiketlemeyle sınanmış",
  "Kapanma katmanının doğruluğu, bir mahallede kör etiketlemeyle (yapay zekâ etiketleyici) sınanmış")
R("Kapanma katmanının doğruluğu yalnızca bir mahallede, tek etiketleyiciyle ve küçük örneklemle sınanmıştır",
  "Kapanma katmanının doğruluğu yalnızca bir mahallede, tek (yapay zekâ) etiketleyiciyle ve küçük örneklemle "
  "sınanmıştır")
R("Sınıflandırıcının beş üst kategori üzerindeki kesinlik, duyarlılık ve F₁ değerleri henüz ölçülmemiştir.",
  f"Etiketli sınama kümesinde üst kategori düzeyinde ölçülen mikro-F₁ {v(b['yeni']['mikroF1'])}'dir (Bölüm "
  "4.10); anahtar sözcük kurtarması bu kümede başarımı katmansız eşlemeye göre "
  f"{v(100 * (b['model']['mikroF1'] - b['yeni']['mikroF1']), 1)} puan düşürmektedir. Eğitim kümesinde 'Lojistik' "
  f"alt etiketi {J['toplamKayit']:,} kaydın {J['lojistikEtiketli']}'inde geçmekte ve yakıt, tahliye, iş makinesi "
  f"gibi taşıma ihtiyaçlarını anlatmaktadır; ulaşım sözcüğü geçen {J['ulasimSozcukluKayit']} kaydın yalnız "
  f"{J['lojistikEtiketliVeUlasimSozcuklu']}'ü bu etiketi taşır ve sözcük çoğunlukla bir adresin parçasıdır. Yol "
  "kapanması ihbarları kümede neredeyse bulunmadığından model onları tanıyamamaktadır.".replace("2,028", "2.028"))
R("üç alan on bir ölçülen alandan, karşılaştırma kollarının ayrıştığı alanlar olarak seçilmiştir (kalan sekizde "
  "kapanmalar rotalara düşmediğinden kollar ayrışmamıştır).",
  "üç alan on bir ölçülen alandan, kolların en az iki modelde ayrıştığı alanlar olarak seçilmiştir (kalan sekizin "
  "beşinde ayrışma yoktur, üçünde yalnız tek modelde; Bölüm 4.7).")
R("(l) Yer seviyesi (kamera) ve Türkçe metin kanallarının başarımı bu sürümde ölçülmemiştir: kamera kanalında "
  "fotoğraflarda kutu düzeyinde yer gerçeği bulunmadığından raporlanan değerler tespit sayısıdır, mAP değildir; "
  "metin kanalında bağımsız etiket bulunmadığından raporlanan değer yazar niyetiyle örtüşmedir, F₁ değildir.",
  "(l) Tatbikattaki yer seviyesi (kamera) ve Türkçe metin kanallarının başarımı ölçülmemiştir: tatbikat "
  "fotoğraflarında kutu düzeyinde yer gerçeği bulunmadığından raporlanan değerler tespit sayısıdır, mAP değildir; "
  "tatbikat metinlerinde bağımsız etiket bulunmadığından raporlanan değer yazar niyetiyle örtüşmedir, F₁ değildir. "
  "Modellerin etiketli kümelerdeki başarımı ayrıca verilmiştir: YOLO modelleri için yalnız eğitimdeki doğrulama "
  "bölmesinde (Bölüm 3.5), metin sınıflandırıcı için ayrılmış sınama bölmesinde (Bölüm 4.10).")
R("için nesne tespiti başarımı (mAP@0,5 ve mAP@0,5:0,95) bu sürümde ölçülmemiştir; ilgili eksen Tablo 22'de "
  "ölçülmedi olarak işaretlenmiştir.",
  "için bağımsız bir sınama kümesinde nesne tespiti başarımı ölçülmemiştir; raporlanan mAP değerleri (Bölüm 3.5) "
  "en iyi ağırlığın seçiminde de kullanılan doğrulama bölmesine aittir ve iyimser bir üst kestirimdir.")

# 5.4
R("İkincisi, kapanma katmanının doğruluğunun tek mahalle ve tek etiketleyicinin ötesinde, çoklu etiketleyiciyle",
  "İkincisi, kapanma katmanının doğruluğunun tek mahalle ve tek (yapay zekâ) etiketleyicinin ötesinde, insan "
  "etiketleyicilerle ve etiketleyiciler arası uyum raporlanarak")
R("Üçüncüsü, kamera ve metin kanallarının başarımının etiketli kümelerle (nesne tespiti için saha test kümesi, "
  "Türkçe metin için DepremKit/TSEqD gibi bir küme) mAP ve F₁ cinsinden ölçülmesi, 13 alt kategoriyi beş üst "
  "kategoriye indiren eşlemenin (Tablo 8) aynı etiketli kümeyle sınanması ve",
  "Üçüncüsü, kamera kanalının bağımsız bir saha sınama kümesinde mAP cinsinden ve metin kanalının DepremKit/TSEqD "
  "gibi dış bir kümede F₁ cinsinden ölçülmesi ve")
R("yol kapanması ihbarlarının anahtar sözcük yerine modelden tanınabilmesi için eğitim kümesinin genişletilmesidir.",
  "yol kapanması ihbarlarının anahtar sözcük yerine modelden tanınabilmesi için eğitim kümesine bu tür "
  "ihbarların eklenmesidir.")
R("Risk skorlama, çok mimarili bina hasarı kıyaslaması, görev tabanlı değerlendirme ve kapanma katmanının kendi "
  "yer gerçeğimizle doğruluğu bu sürümde ölçülmüştür; kalan eksenler (yer seviyesi hasar tespiti mAP, çok "
  "mimarili yol kıyaslaması, bölgeler arası genelleme) bu sürümde ölçülmemiştir.",
  "Risk skorlama, çok mimarili bina hasarı kıyaslaması, görev tabanlı değerlendirme ve metin sınıflandırmanın "
  "üst kategori başarımı ölçülmüştür; kapanma katmanının doğruluğu ve yer seviyesi tespit başarımı yalnız ön "
  "değerlerle (yapay zekâ etiketleyici; eğitimdeki doğrulama bölmesi) verilmiştir; çok mimarili yol kıyaslaması "
  "ve bölgeler arası genelleme ölçülmemiştir.")

# 6
R("kapanma katmanının doğruluğunun bir mahallede kör etiketlemeyle sınanması ve yüksek kesinlik (3/3)",
  "kapanma katmanının doğruluğunun bir mahallede kör etiketlemeyle (bir büyük dil modeliyle; insan doğrulaması "
  "bekleniyor) sınanması ve yüksek kesinlik (3/3)")
R("Yer seviyesi tespit başarımı (mAP) ve bölgeler arası genelleme bu sürümde ölçülmemiş; kapanma katmanının yer "
  "gerçeğine göre değerlendirilmesi yalnızca bir mahallede, tek etiketleyiciyle ve ön deneme olarak yapılabilmiştir.",
  "Yer seviyesi tespit başarımı bağımsız bir sınama kümesinde ve bölgeler arası genelleme bu sürümde ölçülmemiş; "
  "kapanma katmanının yer gerçeğine göre değerlendirilmesi yalnızca bir mahallede, tek (yapay zekâ) "
  "etiketleyiciyle ve ön deneme olarak yapılabilmiştir.")
R("metinde beş üst kategori yerine çok etiketli alt kategori kümesi,", "metinde tek rastgele bölme ve tek koşum,")
R("mahalleler on bir ölçülen alandan, kıyas kollarının ayrıştığı alanlar olarak seçilmiştir.",
  "mahalleler on bir ölçülen alandan, kolların en az iki modelde ayrıştığı alanlar olarak seçilmiştir.")

# Etik ve veri erisilebilirligi
R("Bu çalışmada kullanılan tüm veriler açık erişimli kaynaklardan elde edilmiştir.",
  "Metin sınıflandırma kümesi dışındaki bütün veriler açık erişimli kaynaklardan elde edilmiştir. Metin kümesi, 6 "
  "Şubat 2023 depremlerinin ardından paylaşılmış gerçek ihbar metinlerinden oluşan özel bir kümedir; kişi adları "
  "maskelenmiş olsa da telefon numarası ve adres gibi kişisel veriler içeren kayıtlar barındırmaktadır (Bölüm 3.6) "
  "ve bu çalışmada yalnız model eğitimi ve değerlendirmesi için kullanılmıştır.")
R("küme, makul bir talep hâlinde sorumlu yazardan temin edilebilir.",
  "küme, telefon numarası ve adres gibi kişisel veri alanları maskelendikten sonra makul bir talep hâlinde "
  "sorumlu yazardan temin edilebilir.")

# 4.12 ve Sekil 24
R("Şekil 21–24, sırasıyla yer seviyesi hasar tespiti, deprem risk skorlaması, Türkçe afet metni analizi ve uydu "
  "tabanlı yol hasarı/güzergâh modüllerinin arayüz görünümlerini sunmaktadır. Şekil 21–23 modellerin gerçek "
  "çıktısını gösterir; Şekil 24 ise arayüzün yerleşik örnek veri (simülasyon) kipinde çekilmiştir ve gerçek bir "
  "analiz çıktısı değildir.",
  "Şekil 21–23, sırasıyla yer seviyesi hasar tespiti, deprem risk skorlaması ve Türkçe afet metni analizi "
  "modüllerinin arayüz görünümlerini sunmakta ve modellerin gerçek çıktısını göstermektedir. Uydu tabanlı yol "
  "hasarı ve güzergâh modülünün gerçek çıktıları Şekil 5 ve Şekil 10'da verilmiştir; bu modülün arayüz görüntüsü, "
  "yalnız yerleşik örnek veri kipinde çekilebildiğinden bu sürümde gösterilmemiştir.")
t = next(tb for tb in doc.tables if any("Şekil 24." in c.text for r in tb.rows for c in r.cells))
c24 = t.cell(1, 1)
assert "Şekil 24." in c24.text and "Şekil 23." in t.cell(1, 0).text
for p in list(c24.paragraphs):
    p._p.getparent().remove(p._p)
c24._tc.append(c24._tc.makeelement(qn("w:p"), {}))
yeni = t.cell(1, 0).merge(c24)
for p in list(yeni.paragraphs):
    if not p.text.strip() and not p._p.findall(".//" + qn("w:drawing")):
        p._p.getparent().remove(p._p)
D.log.append("SEKIL 24 kaldirildi; Sekil 23 alt satirda tek hucre")
# Gorsel iceren duzen tablolarinda satir sayfa arasinda bolunmesin (Sekil 21-22'nin
# gorselleri bir sayfada, altyazilari sonrakinde kaliyordu)
from docx.oxml import OxmlElement
n = 0
for tb in doc.tables:
    if tb._tbl.findall(".//" + qn("w:drawing")) and (tb.style is None or tb.style.name != "Table Grid"):
        for tr in tb._tbl.findall(qn("w:tr")):
            trPr = tr.get_or_add_trPr()
            if trPr.find(qn("w:cantSplit")) is None:
                trPr.append(OxmlElement("w:cantSplit"))
                n += 1
D.log.append(f"DUZEN TABLOSU {n} satir bolunmez")

# Tablo 22
t22 = next(tb for tb in doc.tables if tb.rows[0].cells[0].text == "Değerlendirme ekseni")
for row in t22.rows:
    c0 = row.cells[0].text
    if c0.startswith("Kapanma katmanı"):
        D.hucre_yaz(row.cells[2], "Kör etiketleme (yapay zekâ etiketleyici; insan doğrulaması bekleniyor)")
    elif c0.startswith("Yer seviyesi"):
        D.hucre_yaz(row.cells[3], "Bağımsız sınama yok; eğitimdeki doğrulama bölmesinde (en iyi ağırlığın seçiminde "
                                  "de kullanıldı): bina mAP@0,5 = 0,592, mAP@0,5:0,95 = 0,394; çatlak 0,731 / 0,584")
    elif c0.startswith("Metin sınıflandırma"):
        D.hucre_yaz(row.cells[2], "Ayrılmış sınama bölmesi (intent_test_v13, n = 305)")
        D.hucre_yaz(row.cells[3], f"Ölçülmüştür (karşılaştırma ağırlığı; iki lojistik kategorisi birleşik): üst "
                                  f"kategori mikro-F₁ = {v(b['yeni']['mikroF1'])}, makro-F₁ = "
                                  f"{v(b['yeni']['makroF1'])} (düzeltilmiş katman); alt kategori mikro-F₁ = 0,825, "
                                  "makro-F₁ = 0,472")
D.log.append("TABLO 22 uc satir")

doc.save(DST)
print("\n".join(D.log))
