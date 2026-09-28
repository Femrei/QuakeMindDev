# -*- coding: utf-8 -*-
"""Revizyon 15: bölgeler arası genelleme ölçüldü (§4.4, yeni Tablo 14).

Kaynak: makale/analiz_arsivi/kaggle/sonuclar/
  birlesik_P2/*.json            birleşik eğitim (Tablo 12'nin kayıtları, rapor paketi v2)
  oturum2_bolgeler_arasi/*.json il dışarıda bırakma (Kaggle, 28 Eylül; günlükten gunluk_ayikla.py)
  oturum3_segformer_haric_maras/*.json  (varsa) SegFormer, Kahramanmaraş dışarıda
  oturum1_coklu_tohum/*.json    (varsa) U-Net ek tohumları -> oynama bandı

Yeni tablo §4.4'e girdiği için eski Tablo 14–22 bir kayar (15–23); metindeki bütün atıflar, Türkçe
ekleri ve iki şekildeki tablo numarası (Şekil 6, Şekil 12) birlikte güncellenir.
Betik yedekten başlar; yeni sonuç geldikçe yeniden çalıştırılır.
"""
import copy
import glob
import json
import shutil
import sys

sys.path.insert(0, r"C:\Users\utkuc\AppData\Local\Temp\claude\D--quakemind-final-QuakeMind\4779ac37-74b5-4aeb-8690-74d7792eedc5\scratchpad")
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from yardimci import Duzenleyici

KOK = r"D:\quakemind_final\QuakeMind"
SON = KOK + r"\makale\analiz_arsivi\kaggle\sonuclar"
SEK = KOK + r"\QuakeMindBackend\benchmark\makale_sekilleri"
SRC = KOK + r"\makale\arsiv\QuakeMind_Makale_YEDEK_pre_revize15_2026-09-28.docx"
DST = KOK + r"\makale\QuakeMind_Makale_Taslak_2026-09-16_sekillieklendi_v3_LOCAL_2026-09-21.docx"
v = lambda x, b=3: f"{x:.{b}f}".replace(".", ",")
fark = lambda x: ("−" if x < 0 else "+") + v(abs(x))
oku = lambda y: json.load(open(y, encoding="utf-8"))
r3 = lambda x: float(f"{x:.3f}")   # farklar tabloda gorunen (yuvarlanmis) degerlerden alinir

AL = {"Hatay": "sinama_hatay", "Kahramanmaraş": "sinama_kahramanmaras"}
AD = {"unet_mitb4": "U-Net (MiT-B4)", "segformer_mitb4": "SegFormer-B4"}
# Tablo 12'nin kaynagi olan kayitlar (artirim tohumlu, yeni kod); eski SegFormer tohum_42/tohum_7 eski koddur
BIR = {"unet_mitb4": oku(f"{SON}/birlesik_P2/unet_mitb4__tohum_42.json"),
       "segformer_mitb4": oku(f"{SON}/birlesik_P2/segformer_mitb4__tohum_42__2.json")}
assert [round(BIR[m]["sinama_sonrasi_tumu"]["iou_hasar"], 3) for m in AD] == [0.433, 0.407]  # Tablo 12
HAR = {}
for f in glob.glob(f"{SON}/oturum*/P2_haric_*__*__tohum_42.json"):
    k = oku(f)
    HAR[(k["model"], {"P2_haric_hatay": "Hatay", "P2_haric_maras": "Kahramanmaraş"}[k["deney"]])] = k
assert {("unet_mitb4", "Hatay"), ("unet_mitb4", "Kahramanmaraş"), ("segformer_mitb4", "Hatay")} <= set(HAR)
SEG_MARAS = ("segformer_mitb4", "Kahramanmaraş") in HAR
for k in list(HAR.values()) + list(BIR.values()):  # ayni sinama kumesi, ayni ayarlar
    assert k["sinama_hatay"]["n_karo"] == 34 and k["sinama_kahramanmaras"]["n_karo"] == 32
    assert k["sinama_sonrasi_tumu"]["n_karo"] == 91 and k["ayarlar"]["donem"] == 50
    assert k["yazilim"]["torch"] == "2.10.0+cu128" and k["ayarlar"]["ogrenme_orani"] == 2e-4
    assert k["ayarlar"]["artirim_tohumlu"] is True and k["ayarlar"]["girdi"] == "sonrasi"

# gorulmemis ildeki kayip (birlesik - il disarida)
KAYIP = {(m, i): r3(BIR[m][AL[i]]["iou_hasar"]) - r3(HAR[(m, i)][AL[i]]["iou_hasar"]) for (m, i) in HAR}
ORAN = {k: KAYIP[k] / r3(BIR[k[0]][AL[k[1]]]["iou_hasar"]) for k in KAYIP}
KMIN, KMAX = min(KAYIP.values()), max(KAYIP.values())
OTEKI = {(m, i): r3(HAR[(m, i)][AL[j]]["iou_hasar"]) - r3(BIR[m][AL[j]]["iou_hasar"])
         for (m, i) in HAR for j in AL if j != i}
ORUNTU = [k for k in HAR if abs(OTEKI[k]) <= 0.005 and KAYIP[k] >= 0.05]
SAYI = {1: "bir", 2: "iki", 3: "üç", 4: "dört", 5: "beş", 6: "altı"}
SAYIDA = {1: "birinde", 2: "ikisinde", 3: "üçünde", 4: "dördünde", 5: "beşinde", 6: "altısında"}
OMIN, OMAX = min(ORAN.values()), max(ORAN.values())
# birlesik egitimin il bazindaki tohum oynamasi: yalniz ayni kod ayarindaki (artirim tohumlu) kosular
TEKRAR = {}
IL_FARK = None   # tohum degisiminin iki il arasinda yarattigi en buyuk fark (ayni kod, tohum 42'ye gore)
for m in AD:
    ks = [BIR[m]] + [oku(f) for f in glob.glob(f"{SON}/oturum*/P2__{m}__tohum_*.json")]
    ks = [k for k in ks if k["ayarlar"].get("artirim_tohumlu") is True]
    if len(ks) > 1:
        b = BIR[m]
        for k in ks[1:]:
            d = abs((r3(k[AL["Hatay"]]["iou_hasar"]) - r3(b[AL["Hatay"]]["iou_hasar"])) -
                    (r3(k[AL["Kahramanmaraş"]]["iou_hasar"]) - r3(b[AL["Kahramanmaraş"]]["iou_hasar"])))
            IL_FARK = d if IL_FARK is None else max(IL_FARK, d)
        TEKRAR[m] = (sorted(k["tohum"] for k in ks), {i: max(k[a]["iou_hasar"] for k in ks) -
                                                      min(k[a]["iou_hasar"] for k in ks) for i, a in AL.items()})
if TEKRAR:
    BANT = max(max(d.values()) for _, d in TEKRAR.values())
    _toh = sorted({t for tt, _ in TEKRAR.values() for t in tt})
    bant_cumle = (f"Aynı ayarlarla üç tohumla ({', '.join(map(str, _toh))}) yinelenen birleşik "
                  f"{' ve '.join(AD[m] for m in TEKRAR)} eğitimlerinde il bazındaki oynama en çok {v(BANT, 2)} "
                  f"kadardır; ")
else:  # ara surum: ayni kodla ikinci tohum henuz yok; eski kodla yinelenen SegFormer-B4 kosulari
    ks = [oku(f"{SON}/birlesik_P2/segformer_mitb4__tohum_{t}.json") for t in ("42", "42__2", "7")]
    BANT = max(max(k[a]["iou_hasar"] for k in ks) - min(k[a]["iou_hasar"] for k in ks) for a in AL.values())
    bant_cumle = (f"Tohum ve artırım ayarı değiştirilerek yinelenen üç birleşik SegFormer-B4 eğitiminde il bazındaki "
                  f"oynama en çok {v(BANT, 2)} kadardır; ")

u = lambda i, j: r3(HAR[("unet_mitb4", i)][AL[j]]["iou_hasar"]) - r3(BIR["unet_mitb4"][AL[j]]["iou_hasar"])
s = lambda i, j: r3(HAR[("segformer_mitb4", i)][AL[j]]["iou_hasar"]) - r3(BIR["segformer_mitb4"][AL[j]]["iou_hasar"])
karo = lambda k: k["egitim"]["karo"]["egitim"]
assert karo(BIR["unet_mitb4"]) == 1013 and karo(HAR[("unet_mitb4", "Hatay")]) == 521
assert karo(HAR[("unet_mitb4", "Kahramanmaraş")]) == 798
assert s("Hatay", "Hatay") < -0.09 and abs(s("Hatay", "Kahramanmaraş")) < 0.01   # cumlelerdeki nitelemeler
assert abs(u("Hatay", "Hatay") - u("Hatay", "Kahramanmaraş")) < 0.01 and abs(u("Kahramanmaraş", "Hatay")) < 0.01
print("kayip", {k: round(x, 4) for k, x in KAYIP.items()}, "bant", round(BANT, 4), TEKRAR)

shutil.copyfile(SRC, DST)
doc = Document(DST)
D = Duzenleyici(doc)
R = D.degistir

# ---------------------------------------------------------------- 1) tablo numaralari 14-22 -> 15-23
for n in range(22, 13, -1):
    R(f"Tablo {n}", f"Tablo {n + 1}", hepsi=True)
R("Tablo 16–19", "Tablo 16–20")          # eski 15–19
R("Tablo 15, 16", "Tablo 15, 17")        # eski 14, 16
for eski, yeni in (("Tablo 16'teki", "Tablo 16'daki"), ("Tablo 16'te", "Tablo 16'da"), ("Tablo 16'in", "Tablo 16'nın"),
                   ("Tablo 17'daki", "Tablo 17'deki"), ("Tablo 23'de", "Tablo 23'te")):
    R(eski, yeni, hepsi=True)
metin = "\n".join(p.text for p in D.paragraflar())
for yanlis in ("Tablo 16'te", "Tablo 16'in", "Tablo 17'da", "Tablo 23'de", "Tablo 14"):
    assert yanlis not in metin, yanlis
D.log.append("TABLO NUMARALARI 14-22 -> 15-23")

# ---------------------------------------------------------------- 2) sekiller (numara iceren iki sekil)
for alt, dosya in (("Şekil 6. Üç düzeyli değerlendirme", "sekil_protokol.png"),
                   ("Şekil 12. QuakeMind kolunun kazancı", "sekil_kesif_kazanc.png")):
    p = D.basliktan_onceki_gorsel(alt)
    ext = p._p.find(".//" + qn("wp:extent"))
    D.resim_koy(p, f"{SEK}\\{dosya}", int(ext.get("cx")) / 360000)

# ---------------------------------------------------------------- 3) §4.4 metin + Tablo 14
R("Hatay ve Kahramanmaraş yerleşim dokusu, yol genişlikleri ve yapı stoku bakımından farklılık gösterdiğinden iki il "
  "arasında çift yönlü bir genelleme deneyi tasarlanmıştır.",
  "Hatay ve Kahramanmaraş yerleşim dokusu, yol genişlikleri ve yapı stoku bakımından farklılık gösterdiğinden, "
  "kümede en çok karosu bulunan bu iki il için çift yönlü bir il dışarıda bırakma deneyi yapılmıştır.")
yontem = D.yaz(
    "İl-ayrık eğitim (modelin yalnız bir ilin verisiyle eğitilip",
    "Deneyde bir ilin bütün blokları eğitim ve doğrulama kümelerinden çıkarılmış, sınama kümesi değiştirilmemiştir. "
    "Böylece karşılaştırma birleşik eğitimle aynı 91 deprem sonrası sınama karosunda (Hatay 34, Kahramanmaraş 32, "
    "diğer iller 25) yapılmakta; en iyi dönem de dışarıda bırakılan ili içermeyen doğrulama kümesiyle seçildiğinden o "
    "il hiçbir aşamada görülmemektedir. Eğitim kümesi yedi ilden oluştuğu için Hatay'ın çıkarıldığı model "
    "Kahramanmaraş ile birlikte Gaziantep, Malatya, Osmaniye, Adıyaman ve Kilis karolarıyla eğitilmiştir. Deney, "
    "sınamada en yüksek IoU_hasar değerini veren U-Net (MiT-B4) ile işletimdeki modelin mimarisi SegFormer-B4 üzerinde, "
    "Tablo 12 ile aynı ayarlar, yazılım sürümleri ve tohumla (42) koşulmuştur; sonuçlar Tablo 14'tedir.")
cap13 = D.bul("Tablo 13. Kapanma katmanının kendi yer gerçeğimizle")
cap = D.sonrasina_ekle(
    yontem._p,
    "Tablo 14. Bir ilin eğitimden çıkarılmasıyla bölgeler arası genelleme (sınama IoU_hasar, tohum 42). Kalın değer "
    "eğitimde hiç görülmemiş ilde ölçülmüştür; parantez içi birleşik eğitime göre farktır. Sınama kümesi bütün "
    "satırlarda aynıdır (Hatay 34, Kahramanmaraş 32, toplam 91 deprem sonrası karo).", cap13, kalin_bas="Tablo 14.")
cap.paragraph_format.keep_with_next = True

t13 = [t for t in doc.tables if t.rows[0].cells[0].text == "Ölçüt" and "Wilson" in
       " ".join(c.text for r in t.rows for c in r.cells)][0]
satirlar = [["Mimari", "Eğitimden çıkarılan il", "Eğitim karosu", "Hatay", "Kahramanmaraş", "Tüm sınama"]]
kalin = []
for m in AD:
    b = BIR[m]
    satirlar.append([AD[m], "— (birleşik eğitim)", f"{karo(b):,}".replace(",", " "), v(b[AL["Hatay"]]["iou_hasar"]),
                     v(b[AL["Kahramanmaraş"]]["iou_hasar"]), v(b["sinama_sonrasi_tumu"]["iou_hasar"])])
    for i in AL:
        if (m, i) not in HAR:
            continue
        h = HAR[(m, i)]
        hucre = {j: f"{v(h[AL[j]]['iou_hasar'])} ({fark(r3(h[AL[j]]['iou_hasar']) - r3(b[AL[j]]['iou_hasar']))})"
                 for j in AL}
        satirlar.append([AD[m], i, f"{karo(h):,}".replace(",", " "), hucre["Hatay"], hucre["Kahramanmaraş"],
                         v(h["sinama_sonrasi_tumu"]["iou_hasar"])])
        kalin.append((len(satirlar) - 1, 3 if i == "Hatay" else 4))
t = doc.add_table(rows=len(satirlar), cols=6)
t._tbl.remove(t._tbl.tblPr)
t._tbl.insert(0, copy.deepcopy(t13._tbl.tblPr))
cap._p.addnext(t._tbl)
ref_ppr = t13.rows[1].cells[1].paragraphs[0]._p.pPr
ref_rpr = t13.rows[1].cells[1].paragraphs[0].runs[0]._element.rPr
gen = [3.0, 3.1, 1.8, 2.8, 2.9, 2.4]
for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), gen):
    gc.set(qn("w:w"), str(int(w * 567)))
for ri, (row, degerler) in enumerate(zip(t.rows, satirlar)):
    trpr = row._tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit")
    trpr.append(cs)
    for ci, (c, deger) in enumerate(zip(row.cells, degerler)):
        c.width = Cm(gen[ci])
        p = c.paragraphs[0]
        if ref_ppr is not None:
            p._p.insert(0, copy.deepcopy(ref_ppr))
        r = p.add_run(deger)
        if ref_rpr is not None:
            r._element.insert(0, copy.deepcopy(ref_rpr))
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.keep_with_next = ri < len(satirlar) - 1
D.tablo_bicimle(t, punto=8.5)
for ri, ci in kalin:
    for r in t.rows[ri].cells[ci].paragraphs[0].runs:
        r.bold = True
D.log.append(f"TABLO 14 eklendi ({len(satirlar) - 1} satir)")

# yorum
seg_maras = ""
if SEG_MARAS:
    seg_maras = (f" Kahramanmaraş çıkarıldığında SegFormer-B4'te düşüş Kahramanmaraş'ta "
                 f"{fark(s('Kahramanmaraş', 'Kahramanmaraş'))}, Hatay'da {fark(s('Kahramanmaraş', 'Hatay'))} olmuştur.")
bulgu = D.sonrasina_ekle(
    t._tbl,
    f"Eğitimde görülmemiş ilde IoU_hasar birleşik eğitime göre {v(KMIN, 2)}–{v(KMAX, 2)} düşmüştür "
    f"(%{OMIN * 100:.0f}–{OMAX * 100:.0f}). Kahramanmaraş çıkarıldığında eğitim kümesi yalnız %21 küçülmekte "
    f"(1 013 → 798 karo) ve U-Net'te düşüş yalnız Kahramanmaraş'ta görülmektedir ({fark(u('Kahramanmaraş', 'Kahramanmaraş'))}; "
    f"Hatay'da {fark(u('Kahramanmaraş', 'Hatay'))}); bu kayıp büyük ölçüde görülmemiş ilin kendisine bağlanabilir. "
    f"Hatay çıkarıldığında ise eğitim kümesi yarıya inmekte (521 karo) ve U-Net'te düşüş iki ilde benzer kalmaktadır "
    f"({fark(u('Hatay', 'Hatay'))} ve {fark(u('Hatay', 'Kahramanmaraş'))}); bu durumda kaybın ile özgü kısmı veri "
    f"azalmasından ayırt edilememektedir. SegFormer-B4'te ise Hatay çıkarıldığında düşüş neredeyse yalnız Hatay'da "
    f"görülmektedir ({fark(s('Hatay', 'Hatay'))}; Kahramanmaraş'ta {fark(s('Hatay', 'Kahramanmaraş'))}); bu mimaride "
    f"de görülmemiş ilin etkisi öne çıkmaktadır.{seg_maras} " + bant_cumle
    + ("görülmemiş ilde ölçülen kayıpların hepsi bu oynamanın üstündedir." if KMIN > BANT else
       f"görülmemiş ildeki kayıplar bu oynamayla aynı mertebededir. Dolayısıyla tek tohumla ölçülen kaybın büyüklüğü "
       f"kesin değildir; tutarlı olan örüntüdür: {SAYI[len(HAR)]} koşunun {SAYIDA[len(ORUNTU)]} dışarıda bırakılan ilin değeri "
       f"{v(min(KAYIP[k] for k in ORUNTU), 2)}–{v(max(KAYIP[k] for k in ORUNTU), 2)} düşerken aynı modelin öteki ildeki "
       f"değeri değişmemiştir (en çok ±0,005). "
       + (f"Tohum değişimi birleşik eğitimde de iki il arasında {v(IL_FARK, 2)} kadar fark yaratabildiğinden "
          f"bu örüntü de tek tohumla kesinleşmez; " if IL_FARK is not None and IL_FARK >= min(KAYIP[k] for k in ORUNTU) - 0.005
          else (f"Bu fark, tohum değişiminin birleşik eğitimde iki il arasında yarattığı farkın (en çok {v(IL_FARK, 2)}) "
                f"üzerindedir; " if IL_FARK is not None else ""))
       + "kaybın büyüklüğünü kesinleştirmek için deneyin çoklu tohumla yinelenmesi gerekir."),
    yontem)
D.sonrasina_ekle(
    bulgu._p,
    "Sınırlar: deney tek tohumludur ve il başına 32–34 sınama karosuna dayanır; çıkarılan ilin payı eğitim kümesinin "
    "boyutunu da değiştirdiğinden (Hatay %49, Kahramanmaraş %21), aynı sayıda karonun rastgele çıkarıldığı, veri "
    "miktarıyla eşlenmiş bir kıyas gelecek çalışmaya bırakılmıştır. Operasyonel sonucu şudur: model yeni bir deprem "
    "bölgesinde kullanılacaksa Tablo 12'deki sınama değerleri iyimser okunmalı, o bölgeden etiketli az sayıda karoyla "
    "ince ayar yapılmadan bu değerlere güvenilmemelidir.",
    yontem)

# ---------------------------------------------------------------- 4) oteki yerler
R("kapanma katmanının bir mahallede kör insan etiketlemesiyle ön doğrulaması (Bölüm 4.3, Tablo 13), kamera",
  "kapanma katmanının bir mahallede kör insan etiketlemesiyle ön doğrulaması (Bölüm 4.3, Tablo 13), bir ilin "
  "eğitimden çıkarıldığı bölgeler arası genelleme deneyi (Bölüm 4.4, Tablo 14), kamera")
R("Bölüm 4.4'teki il-ayrık genelleme deneyi bu sürümde koşulmamıştır.",
  "Bölüm 4.4'teki genelleme deneyi iki mimaride ve tek tohumla koşulmuştur.")
R("İl bazında Hatay 0,267-0,425, Kahramanmaraş 0,324-0,456 aralığındadır.",
  "İl bazında Hatay 0,267-0,425, Kahramanmaraş 0,324-0,456 aralığındadır; bir ilin eğitimden çıkarıldığı genelleme "
  "deneyi Bölüm 4.4'tedir.")
R("çok mimarili yol kıyaslaması ve bölgeler arası genelleme ölçülmemiştir.",
  "çok mimarili yol kıyaslaması ölçülmemiş, bölgeler arası genelleme iki mimaride tek tohumla ölçülmüştür.")
R("Yer seviyesi tespit başarımı bağımsız bir sınama kümesinde ve bölgeler arası genelleme bu sürümde ölçülmemiş;",
  "Yer seviyesi tespit başarımı bağımsız bir sınama kümesinde ölçülmemiş, bölgeler arası genelleme ise iki mimaride "
  "ve tek tohumla ölçülebilmiştir;")
ICINDE = KMIN <= BANT
R("değeri U-Net ile elde edilmiştir (0,433).",
  f"değeri U-Net ile elde edilmiştir (0,433). Bir ilin eğitimden çıkarıldığı genelleme deneyinde görülmemiş ilde "
  f"IoU_hasar {v(KMIN, 2)}–{v(KMAX, 2)} düşmüştür" + (" (tek tohum; tohum oynaması mertebesinde)." if ICINDE else "."))
R("was obtained by U-Net (0.433).",
  f"was obtained by U-Net (0.433). When one province was held out of training, IoU on that unseen province "
  f"dropped by {KMIN:.2f}–{KMAX:.2f}" + (" (single seed; within seed variation)." if ICINDE else "."))

t23 = next(tb for tb in doc.tables if tb.rows[0].cells[0].text == "Değerlendirme ekseni")
for row in t23.rows:
    if row.cells[0].text.startswith("Bölgeler arası genelleme"):
        D.hucre_yaz(row.cells[2], "Eğitimden çıkarılan il (aynı sınama kümesi)")
        D.hucre_yaz(row.cells[3], f"Ölçülmüştür (Tablo 14): görülmemiş ilde IoU_hasar {v(KMIN, 2)}–{v(KMAX, 2)} düşük "
                                  f"(U-Net ve SegFormer-B4, tek tohum" +
                                  ("; tohum oynaması mertebesinde, çoklu tohumla yinelenecek)" if ICINDE else ")"))
        for c in row.cells:
            for rpr in c._tc.iter(qn("w:rPr")):
                for h in rpr.findall(qn("w:highlight")):
                    rpr.remove(h)
D.log.append("TABLO 23 genelleme satiri")

doc.save(DST)
print("\n".join(D.log))
