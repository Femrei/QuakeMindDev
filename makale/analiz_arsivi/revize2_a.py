# -*- coding: utf-8 -*-
"""Tarama revizyonu - BOLUM A: kodla dogrulanan veri/altyapi/metin/kamera/risk
duzeltmeleri ve bayat (eski senaryo) kalintilarinin temizligi."""


def uygula(D, doc):
    deg, yaz, sil = D.degistir, D.yaz, D.sil

    # ================================================================ OZ/ABSTRACT
    yaz("An Architecture and a Task-Based Evaluation Protocol Toward Integrating Road Damage",
        "An Architecture and a Task-Based Evaluation Protocol Toward Integrating Satellite-Imagery "
        "Damage Segmentation and Road-Network Intersection with Blockage-Aware Team Dispatch: The Case "
        "of the 6 February 2023 Kahramanmaraş Earthquakes")

    # ================================================================ KISALTMALAR
    t0 = doc.tables[0]
    for row in t0.rows:
        k = row.cells[0].text
        if k.startswith('KATE-CD / MEDBFS'):
            D.hucre_yaz(row.cells[0], 'KATE-PD / KATE-CD / MEDBFS')
            D.hucre_yaz(row.cells[1], 'Bölütleme karşılaştırmasında kullanılan açık bina hasarı veri kümeleri '
                        '(KATE-PD ve KATE-CD eğitim/sınama, MEDBFS yalnız dış sınama) [56, 57]')
        elif k == 'v_ort':
            D.hucre_yaz(row.cells[0], 'v(e)')
            D.hucre_yaz(row.cells[1], 'Kenar hızı (maxspeed etiketinden, yoksa yol sınıfından; düz bir ortalama hız kullanılmaz)')
        elif k == 't_u':
            D.hucre_yaz(row.cells[1], 'Kapanmaya varan ekibin keşif başına kaybettiği süre (bu sürümde değerlendirme 1,0 dk + manevra 25 sn = 1,42 dk)')
        elif k == 'k':
            D.hucre_yaz(row.cells[1], 'Kapanmayı bilmeyen ekibin yolda keşfettiği kapanma (yeniden yönlenme) sayısı')
        elif k.startswith('R, φ, λ'):
            D.hucre_yaz(row.cells[1], 'A* sezgiselindeki büyük çember (haversine) uzaklığında Dünya yarıçapı, enlem ve boylam')

    # ================================================================ 1. GIRIS
    deg("Uydu tabanlı hızlı haritalama servisleri bu boşluğu",
        "Uydu tabanlı hızlı haritalama servisleri [8] bu boşluğu")
    deg("beş rastgelelik tohumu, üç model ve üç etkin veri kanalı üzerinde değerlendirilmiştir.",
        "beş rastgelelik tohumu ve üç bölütleme modeli üzerinde değerlendirilmiştir. Yer seviyesi (kamera) "
        "ve Türkçe metin kanalları aynı tatbikatta koşturulmuş, ancak ölçülen erişim süreleri yalnız "
        "uydu/drone kanalından türetilen kapanmalara dayanmaktadır.")

    # ================================================================ 2.1
    deg("Bu çalışmada, hasar bölütleme bileşeni için SegFormer'ın MiT-B4 kodlayıcılı sürümü esas alınmıştır;",
        "Bu çalışmada hasar bölütleme bileşeninin mimari çekirdeği MiT kodlayıcısıdır: platformun varsayılan "
        "işletim modeli SegFormer MiT-B4'tür; tatbikatta ise Bölüm 4.2'deki protokolle eğitilen üç model "
        "(SegFormer MiT-B0, SegFormer MiT-B2 ve U-Net MiT-B4) koşturulmuştur;")

    # ================================================================ 3.1
    deg("Bölüm 4.5'teki senaryo değerlendirmesi bu beş kanaldan üçünü (yer seviyesi görüntü, Türkçe metin, "
        "acil çağrı) etkinleştirmektedir; uydu ve sismik katalog kanalları o koşuda sevk kararına girdi üretmemiştir.",
        "Bölüm 4.5'teki tatbikat bu beş kanaldan üçünü (uydu/drone görüntüsü, yer seviyesi görüntü, Türkçe metin) "
        "koşturmaktadır; erişim süresini değiştiren kapanmalar yalnız uydu/drone kanalından gelir, kamera ve metin "
        "kanalları ısı haritasına kayıt üretir. Acil çağrı ve sismik katalog kanalları tatbikatta kullanılmamıştır.")
    deg("Bu katmanın çıktıları ortak bir biçime indirgenir: konum, hasar veya aciliyet skoru, güven değeri ve zaman damgası.",
        "Bu katmanın çıktıları ortak bir biçime indirgenir: konum, hasar veya aciliyet skoru, güven değeri, zaman "
        "damgası ve kaynak tipi.")
    deg("çizge önbellekleri bölge bazında saklanabildiği için çevrim dışı çalışma mümkün olur.",
        "çizge önbellekleri bölge bazında saklanabilir.")
    deg("Uygulama yüzeyi bir REST hizmet katmanı ile bu katmanı kullanan web ve mobil istemcilerden oluşmaktadır;",
        "Uygulama yüzeyi bir REST hizmet katmanı ile bu katmanı kullanan web istemcilerinden oluşmaktadır;")

    # ================================================================ 3.2 / Tablo 2-4
    t2 = doc.tables[2]
    for ri, row in enumerate(t2.rows):
        k = row.cells[0].text
        if k.startswith('Esri World Imagery Wayback'):
            D.hucre_yaz(row.cells[3], 'Hatay, Gaziantep, Kahramanmaraş')
        elif k.startswith('Help.NGO'):
            D.hucre_yaz(row.cells[1], 'İnsansız hava aracı ortofotosu, üç sahne (Tablo 3–4)')
            D.hucre_yaz(row.cells[2], '9 ve 14 Şubat 2023 (deprem sonrası)')
            D.hucre_yaz(row.cells[3], 'Hatay (Defne/Elektrik), Gaziantep (İslahiye, Nurdağı)')
            D.hucre_yaz(row.cells[4], 'Özgün 3,4–7,9 cm/piksel; analizde z18 karolarıyla yaklaşık 0,48 m/piksel')
        elif k.startswith('KATE-CD'):
            D.hucre_yaz(row.cells[1], 'Deprem öncesi ve sonrası optik uydu görüntü çiftleri (Maxar Open Data, Airbus '
                        'Pléiades) ve elle çizilmiş ikili hasar maskeleri; bu çalışmada yalnız deprem öncesi görüntüleri '
                        'hasarsız (negatif) örnek olarak kullanılmıştır (Bölüm 4.2)')
            D.hucre_yaz(row.cells[4], '0,3–0,5 m; 512 × 512 piksel; 486 çift; sonrası görüntüleri KATE-PD ile özdeş '
                        'olduğundan yalnız öncesi görüntüleri alınmıştır')
        elif k.startswith('MEDBFS'):
            D.hucre_yaz(row.cells[1], 'Deprem öncesi (WorldView) ve sonrası (Jilin-1) optik uydu görüntüsü; piksel '
                        'düzeyinde ağır hasar ve yıkılma etiketleri; bu çalışmada eğitime katılmamış, yalnız dış '
                        'sınamada kullanılmıştır (Bölüm 4.2)')
            D.hucre_yaz(row.cells[4], '0,5 m (dış sınamada 0,75 m); 512 × 512 piksel; 438 deprem öncesi + 438 deprem sonrası görüntü')
    # KATE-PD ve Wikimedia satirlari
    kate_idx = [i for i, r in enumerate(t2.rows) if r.cells[0].text.startswith('KATE-CD')][0]
    D.satir_ekle(t2, kate_idx - 1 if False else kate_idx - 0, [
        'KATE-PD [56]', 'Deprem sonrası optik uydu görüntüsü (Maxar 193, Airbus Pléiades 639 karo) ve hasarlı bina '
        'maskeleri; bina hasarı bölütlemesinin eğitim, doğrulama ve sınama kümesi (Bölüm 4.2)',
        '2023 (deprem sonrası)', 'Adıyaman, Gaziantep, Hatay, Kahramanmaraş, Kilis, Malatya, Osmaniye',
        '0,375–0,5 m; 512 × 512 piksel; 832 karo', 'Hugging Face veri deposu (CSCRS/kate-pd); veri kartında lisans belirtilmemiş'])
    # satir_ekle kaynak satirin ARKASINA ekler; KATE-PD'yi KATE-CD'nin onune al
    rows = t2._tbl.findall(qn_tr())
    kp = [r for r in rows if 'KATE-PD [56]' in ''.join(x.text or '' for x in r.iter(qn_t()))][0]
    kc = [r for r in rows if ''.join(x.text or '' for x in r.iter(qn_t())).startswith('KATE-CD')][0]
    kc.addprevious(kp)
    ted = [i for i, r in enumerate(t2.rows) if r.cells[0].text.startswith('AFAD')][0]
    D.satir_ekle(t2, ted, [
        'Wikimedia Commons deprem fotoğrafları', 'Yer seviyesi fotoğraf (13 adet); kamera kanalının tatbikat girdisi (Bölüm 4.5)',
        'Şubat–Ağustos 2023 (deprem sonrası)', 'Çoğu Hatay/Antakya; dördünde il bilgisi yok', 'Özgün fotoğraf çözünürlüğü',
        'Kamu malı (5) ve CC BY-SA 4.0 (8); künye ve SHA-256 ile kayıtlı'])
    deg("Tablo 2. Çalışmada kullanılan dış (harici) veri kaynakları. Yer seviyesi görüntü ve acil çağrı kanalları "
        "senaryo içi üretilen veya toplanan verilerdir ve harici bir arşiv kaynağı taşımadıkları için bu tabloya dahil "
        "edilmemiştir. Son üç satır, model eğitimi ve değerlendirmesinde kullanılan kümeleri gösterir.",
        "Tablo 2. Çalışmada kullanılan dış (harici) veri kaynakları. Tatbikattaki ihbar metinleri elle yazılmış "
        "sentetik metinlerdir ve bu tabloya dahil edilmemiştir; acil çağrı kanalı tatbikatta kullanılmamıştır. Son "
        "dört satır, model eğitimi ve değerlendirmesinde kullanılan kümeleri gösterir.")

    deg("Üç sahne de 9 Şubat 2023'te (deprem sonrası) çekilmiş ve OpenAerialMap üzerinden açık lisansla yayımlanmıştır.",
        "Defne sahnesi 9 Şubat, İslahiye ve Nurdağı sahneleri 14 Şubat 2023'te (deprem sonrası) çekilmiş; üçü de "
        "Help.NGO tarafından OpenAerialMap üzerinden CC BY 4.0 lisansıyla yayımlanmıştır [64].")
    t3 = doc.tables[3]
    vals = {'Defne': ('Help.NGO (OpenAerialMap)', 'CC BY 4.0', '3,4 cm/piksel'),
            'İslahiye': ('Help.NGO (OpenAerialMap)', 'CC BY 4.0', '7,8 cm/piksel'),
            'Nurdağı': ('Help.NGO (OpenAerialMap); sağlayıcının "Low Quality" uyarısıyla', 'CC BY 4.0', '7,1 cm/piksel')}
    for row in t3.rows[1:]:
        for k, v in vals.items():
            if row.cells[0].text.startswith(k):
                for ci, x in enumerate(v, start=1):
                    D.hucre_yaz(row.cells[ci], x)
    D.hucre_yaz(t3.rows[0].cells[3], 'Özgün çözünürlük')
    deg("Görüntüler deprem sonrasına (9 Şubat 2023) aittir; depremden birkaç gün sonra bir miktar açma çalışması başlamış olabilir.",
        "Görüntüler deprem sonrasına (9–14 Şubat 2023) aittir; depremden üç ila sekiz gün sonra bir miktar açma "
        "çalışması başlamış olabilir. Model bu ortofotoları özgün çözünürlükte değil, eğitim verisinin çözünürlüğüne "
        "yakın z18 karolarıyla (yaklaşık 0,48 m/piksel) işler (Bölüm 3.3).")
    # Tablo 4: Defne penceresi katalogda 499,3 m
    t4 = doc.tables[4]
    pencere = {'Defne': '499,3 m', 'İslahiye': '499,7 m', 'Nurdağı': '500,5 m'}
    for row in t4.rows[1:]:
        for k, v in pencere.items():
            if row.cells[0].text.startswith(k):
                D.hucre_yaz(row.cells[5], v)

    # Tablo 5: EMSR AOI'leri artik kullanilmiyor
    t5 = doc.tables[5]
    for row in t5.rows:
        if row.cells[0].text.startswith('İlgi alanları'):
            D.hucre_yaz(row.cells[1], '[BELİRLENECEK]')

    # ================================================================ 3.3
    yaz("Uydu görüntüsü, bölütleme modelinin girdi boyutuna uygun biçimde",
        "Uydu görüntüsü, bölütleme modelinin girdi boyutuna uygun biçimde 512 × 512 piksellik model karolarına "
        "bölünmektedir; bu, sağlayıcıdan çekilen harita karosundan (256 × 256 piksel) farklı bir birimdir. Yakınlaştırma "
        "düzeyi z, görüntünün yer örnekleme aralığı eğitim verisinin aralığına (0,3–0,5 m/piksel) en yakın olacak biçimde "
        "seçilir: hedef aralık 0,4 m/piksel olmak üzere, enlem φ'deki yer örnekleme aralığı 156 543·cos(φ)/2^z m/piksel "
        "ile hesaplanır ve z = yuvarla(log₂(156 543·cos φ / 0,4)) olarak, [13, 20] aralığında alınır. Çalışma alanının enleminde (36–37°) bu koşul z = 18'i "
        "(yaklaşık 0,48 m/piksel) verir; tatbikattaki yaklaşık 500 m'lik mahalle kutusu 1280 × 1280 piksellik bir "
        "görüntüye karşılık gelir. Bir analizde getirilen karo sayısı en çok 8 × 8 = 64 ile sınırlıdır. Bu seçim, "
        "özgün çözünürlüğü santimetre düzeyinde olan drone ortofotolarının da modele eğitim ölçeğine yakın bir "
        "çözünürlükte verilmesini sağlar.")
    yaz("Karo kalite filtresi (bu sürümde uygulanmamıştır).",
        "Karo kalite filtresi (kodda tanımlı, çıkarım hattına bağlı değildir). Aşağıda tanımlanan kalite kontrol "
        "adımı kod düzeyinde yazılmış, ancak bu sürümde çıkarım hattından çağrılmamaktadır; tanım, tekrarlanabilirlik "
        "açısından burada verilmektedir. Filtre iki ölçüte dayanır: asgari yer örnekleme aralığı eşiği ve bulanıklık "
        "tespiti. Bulanıklık, gri seviyeye dönüştürülmüş karonun Laplacian yanıtının varyansı ile ölçülür. I(x, y) "
        "karonun gri seviye yoğunluğu ve L = ∇²I Laplacian yanıtı olmak üzere bulanıklık ölçütü Denklem (1) ile tanımlanır:")
    yaz("Denklem (1)'de M ve N karonun piksel boyutları",
        "Denklem (1)'de M ve N karonun piksel boyutları, L̄ Laplacian yanıtının ortalamasıdır. V < V₀ koşulunu sağlayan "
        "karolar bulanık kabul edilerek analiz dışı bırakılacak ve ilgili alan sonuç maskesinde \"değerlendirilmedi\" "
        "olarak işaretlenecektir. Kodda (tile_quality_score, road_damage/utils/inference.py) varsayılan eşik V₀ = 100 "
        "ve tamamlayıcı bir bulut sezgiseli (HSV'de parlaklığı 200'ün üstünde, doygunluğu 40'ın altında olan piksellerin "
        "oranı en çok 0,4) tanımlıdır. Filtrenin elenen karo oranı ve IoU_hasar üzerindeki etkisi, filtre hatta "
        "bağlandığında ve etiketli bir sınama karosu kümesiyle raporlanacaktır [GELECEK ÇALIŞMA].")

    # ================================================================ 3.4
    deg("Hasar bölütlemesi için esas alınan model, MiT-B4 kodlayıcı ile kullanılan SegFormer mimarisidir [19].",
        "Hasar bölütlemesi MiT kodlayıcılı modellerle yapılır: platformun varsayılan işletim modeli SegFormer [19] "
        "MiT-B4'tür; tatbikatta ise Bölüm 4.2'deki protokolle eğitilmiş SegFormer MiT-B0, SegFormer MiT-B2 ve "
        "U-Net MiT-B4 modelleri kullanılır (Bölüm 4.6).")
    deg("Tek sınıflı kurgu tercih edilmiştir; çünkü Tablo 3'te gösterildiği üzere yol hasarı derecelerinin ilgi alanları "
        "arasındaki dağılımı dengesizdir ve Kahramanmaraş'ta \"Destroyed\" sınıfı hiç gözlenmemektedir. Çok sınıflı bir "
        "kurgu, örneklem sayısı sıfır olan bir sınıfı öğrenmeye çalışmak anlamına gelirdi.",
        "Tek sınıflı kurgu tercih edilmiştir; çünkü eğitimde kullanılan açık kümeler (KATE-PD, KATE-CD) ikili hasar "
        "etiketi taşımakta, yol hasarını derecelendiren piksel etiketli bir Türkiye kümesi ise bulunmamaktadır "
        "(Bölüm 3.2.1).")
    deg("istemci bu kimlikle durum uç noktasını yoklayarak ilerleme oranını izler",
        "yeni konsol istemcisi ilerlemeyi ve tamamlanma bildirimini tek bir sunucu-gönderimli olay (SSE) akışından "
        "izler, sonucu tamamlanma bildirimi geldiğinde tek istekle alır (eski web istemcisi durum uç noktasını yoklar)")
    deg("Platformun işletim sürümünde kullanılan, kamuya açık olarak yayımlanmış",
        "Platformun varsayılan işletim modeli olan, kamuya açık olarak yayımlanmış")

    # ================================================================ 3.5 kamera
    deg("Her iki model için güven eşiği 0,20 olarak belirlenmiştir.",
        "Her iki model için uygulama katmanında 0,20 güven eşiği tanımlıdır; ancak model çağrısında ayrı bir eşik "
        "verilmediğinden kütüphanenin varsayılan 0,25 eşiği önce uygulanır ve fiilî eşik 0,25'tir.")
    deg("iki ayrı YOLOv8 [40] modeli konumlandırılmıştır.",
        "iki ayrı YOLOv8s [40] modeli konumlandırılmıştır.")

    # ================================================================ 3.6 metin
    yaz("Birinci aşama — biçimbilimsel çözümleme.",
        "Birinci aşama — metin normalleştirme. Türkçe sondan eklemeli bir dil olduğundan, yüzey biçimleri sözlük "
        "tabanlı eşlemeyi kırar. Önceki sürümde hatta çağrılan kök/ek ayrıştırıcının (Zemberek) çıktısı ne "
        "sınıflandırıcıya ne de anahtar sözcük kurallarına aktarılmadığından çağrı bu sürümde kaldırılmıştır; "
        "sınıflandırıcı ve anahtar sözcük kuralları küçük harfe çevrilmiş ham metin üzerinde çalışır.")
    yaz("İkinci aşama — sınıflandırma.",
        "İkinci aşama — sınıflandırma. BERTurk [32] tabanlı bir metin sınıflandırıcı, 13 alt kategori için sigmoid "
        "çıkışlı çok etiketli bir kurguda skor üretir; sevk hattı bu skorlardan en yüksek olanın etiketini alır. Alt "
        "kategoriler, sevk mantığının kullandığı beş üst kategoriye eşlenir (Tablo 9). Bu eşlemenin önünde bir anahtar "
        "sözcük kurtarma katmanı çalışır: metin dört kategoriden (Acil Yardım İhtiyacı, Enkaz Bildirimi, Lojistik "
        "İhtiyaç Bildirimi, Yol Kapanma Bilgisi) birine ait bir anahtar sözcük içeriyorsa kayıt, modelin kararından "
        "bağımsız olarak o kategoriye atanır. Anahtar sözcük bulunmaz ve en yüksek skor 0,30'un altında kalırsa kayıt "
        "\"Belirsiz\" kategorisine atanır; bu durumda varlık ismi tanıma ve coğrafi kodlama yine çalıştırılır ve kayıt "
        "düşük aciliyet düzeyiyle sevk zincirine aktarılır. Model \"Alakasız\" etiketini verdiğinde kayıt zincirden "
        "düşer. Bu tasarımın sonuçları Bölüm 4.5 ve Bölüm 5.3(c)'de tartışılmaktadır.")
    yaz("Düşük güvenli kayıtların doğrudan elenmesi yerine belirsiz duruma alınması",
        "Düşük güvenli kayıtların doğrudan elenmesi yerine belirsiz duruma alınması, önceki sürümün senaryo koşumunda "
        "gözlenen iki yanlış negatife (altı metin ihbarının ikisi, ikisi de medikal çağrı) yanıt olarak yapılmış bir "
        "tasarım değişikliğidir: o sürümde anahtar sözcükle kurtarma yalnızca yol ve ulaşım alanı için tanımlıydı. Bu "
        "sürümde kurtarma kümesi dört kategoriye genişletilmiştir (örneğin sağlık için \"ilaç\", \"nefes al\", "
        "\"kanama\", \"yaralı\"; enkaz için \"enkaz\", \"altında kaldı\", \"göçük\", \"ses geliyor\"; ulaşım için "
        "\"yol\", \"köprü\", \"ulaşım\", \"kapalı\", \"geçemiyor\"). Ancak bu genişletme yeni bir sorun doğurmuştur: "
        "kurtarma katmanı modelin kararından önce ve koşulsuz çalıştığı için modelin yüksek güvenle verdiği doğru "
        "etiketi ezebilmekte, iki harfli anahtar sözcükler (\"aç\") alt dizgi olarak \"araç\" veya \"ilaç\" gibi "
        "sözcüklerde de eşleşmekte ve döndürülen güven değeri gösterilen kategoriye ait olmamaktadır. Bu durumun ölçülen "
        "etkisi Bölüm 4.5'te verilmektedir.")
    D.hucre_yaz_bul = None
    t9 = doc.tables[9]
    for row in t9.rows:
        if row.cells[0].text.startswith('Alakasız'):
            D.hucre_yaz(row.cells[2], '— (zincirden düşer)')
    deg("* \"Lojistik\" alt kategorisi, metinde ulaşım alanına ait anahtar sözcük bulunup bulunmamasına göre iki farklı "
        "üst kategoriye ayrışır; tabloda bu nedenle iki satırda gösterilmiştir. Toplam alt kategori sayısı 13'tür.",
        "* \"Lojistik\" alt kategorisi, metinde ulaşım alanına ait anahtar sözcük bulunup bulunmamasına göre iki farklı "
        "üst kategoriye ayrışacak biçimde tasarlanmıştır; uygulamada ulaşım sözcüğü taşıyan metin kurtarma katmanında "
        "zaten \"Yol Kapanma Bilgisi\"ne atandığından, modelin \"Lojistik\" dediği kayıtlar \"Lojistik İhtiyaç "
        "Bildirimi\"ne (aciliyet 2) düşer. Toplam alt kategori sayısı 13'tür.")
    deg("Bölüm 4.5'te kaçırılan iki ihbarın hayati medikal çağrı olması, bu varsayımın sahada sınanması gerektiğini göstermektedir.",
        "Önceki sürümde kaçırılan iki ihbarın hayati medikal çağrı olması, bu varsayımın sahada sınanması gerektiğini "
        "göstermektedir.")

    # ================================================================ 3.7 risk
    deg("H_uzun son on yılda M ≥ 6,0 olaylarından türetilen uzun vadeli tehlike bileşeni,",
        "H_uzun uzun vadeli tehlike bileşeni (200 km içindeki M ≥ 6,0 olaylarının katalog boyunca yıllık oranı λ'dan "
        "türetilen, önümüzdeki on yılda en az bir M ≥ 6,0 olayı olasılığı: 1 − e^(−10λ)),")
    deg("Ağırlıklar uzman tanımlıdır ve öğrenilmiş değildir.",
        "Ağırlıklar (0,4 / 0,3 / 0,3) uzman tanımlıdır ve öğrenilmiş değildir; F_fay en yakın faya uzaklığa göre "
        "kademelidir (30 km altı 1,0; 80 km altı 0,7; 150 km altı 0,4; daha uzak 0,1).")
    deg("Ağırlıkların duyarlılığı ve kalibrasyonu gelecek çalışmada sınanacaktır (Bölüm 5.3(k)).",
        "Ağırlıkların duyarlılığı yalnızca kaba bir taramayla ölçülmüş (Bölüm 4.11), kalibrasyonu gelecek çalışmaya "
        "bırakılmıştır (Bölüm 5.3(k)).")

    # ================================================================ 3.9
    deg("Bu değerlendirmede kilitleme, iplik-güvenli (thread-safe) bir TaskLock sınıfıyla uygulanıp gerçek eşzamanlı "
        "isteklerle test edilmiştir (Bölüm 4.5); üretim REST API'sinde bu mekanizmanın bir HTTP uç noktasına (ör. HTTP "
        "409 Conflict yanıtı) bağlanması bu sürümde gerçekleştirilmemiştir.",
        "Kilitleme, önceki sürümün senaryo koşumunda iplik-güvenli (thread-safe) bir TaskLock sınıfıyla uygulanıp "
        "gerçek eşzamanlı isteklerle sınanmıştır; üretim REST API'sinde ekip talebi ucu aynı hedef için ikinci talebi "
        "HTTP 409 Conflict yanıtıyla reddeder. Bu uçta bu sürümde yetki denetimi bulunmamaktadır.")
    deg("Erişim, üç rollü bir yetkilendirme şeması (afetzede, saha ekibi, yönetici) olarak tasarlanmıştır;",
        "Erişim, dört rollü bir yetkilendirme şeması (afetzede, saha ekibi, koordinatör, yönetici) olarak tasarlanmıştır;")
    deg("Atama kuralı. Senaryoda atama her ilçe bölgesi için ayrı yapılır ve Algoritma 3'teki",
        "Atama kuralı. Tasarımda atama her ilçe bölgesi için ayrı yapılır ve Algoritma 2'deki")
    deg("Algoritma 3. Bölge içi rol-farkında açgözlü ekip ataması.", "Algoritma 2. Bölge içi rol-farkında açgözlü ekip ataması.")
    deg("hedef kilitleme bir eşzamanlılık denetimi, Algoritma 3 ise", "hedef kilitleme bir eşzamanlılık denetimi, Algoritma 2 ise")
    deg(" Bölüm 4.5'te raporlanan sonuçlar da bu dar iddiaya karşılık gelmektedir.", "")

    # ================================================================ 4 giris
    yaz("Ölçülen ve ölçülmeyenlerin ayrımı.",
        "Ölçülen ve ölçülmeyenlerin ayrımı. Bu sürümde ölçülen sonuçlar şunlardır: bina hasarı bölütlemesinde on "
        "mimarinin karşılaştırması (Bölüm 4.2, Tablo 13), kapanma katmanının bir mahallede kör etiketlemeyle ön "
        "doğrulaması (Bölüm 4.3, Tablo 14), kamera ve metin kanallarının tespit sayısı ve gecikmesi (Bölüm 4.5, "
        "Tablo 16), tatbikatın kıyas kolları, model seçimi, güzergâh varyantları, r_h ve ceza duyarlılığı (Bölüm "
        "4.6–4.9, Tablo 17–21) ile metin ve risk kıyaslamaları (Bölüm 4.10 ve 4.11). Bölüm 4.2'nin karşılaştırması yol "
        "hasarı değil bina hasarı görevinde yürütülmüştür; kodlayıcı sabitleme, mekânsal bölme ve sınama kümesinde "
        "hasarsız karo bulunması hedeflerinin üçü de bu koşumda karşılanmıştır, tek açık nokta çoklu tohum varyansının "
        "tüm mimarilerde ölçülmemiş olmasıdır (Bölüm 4.2, Sınırlılık). Bölüm 4.4'teki tablo (Tablo 15) büyük ölçüde "
        "iskelet hâlindedir; kalan hücreler [ÖLÇÜLECEK] ile işaretlenmiştir. Tatbikatta kapanma kümesi doğrudan kendi "
        "hasar modelimizin drone ortofotosu üzerindeki çıktısından türetilir; sentetik bir kapanma üreticisi "
        "kullanılmamaktadır. Tatbikatta yer gerçeği bir kapanma katmanı bulunmadığından dünya modelin kendi çıktısıdır; "
        "bu ayrım sonuçların yorumlanmasında belirleyicidir.")

    # ================================================================ 4.1.1 Tablo 12
    t12 = doc.tables[12]
    for row in t12.rows:
        k = row.cells[0].text
        if k.startswith('Ortalama hız'):
            D.hucre_yaz(row.cells[0], 'Kenar hızı v(e)')
            D.hucre_yaz(row.cells[1], 'maxspeed etiketinden, yoksa yol sınıfından (Bölüm 3.8)')
        elif k.startswith('Yeniden yönlenme cezası'):
            D.hucre_yaz(row.cells[0], 'Kapanma cezası t_u')

    # ================================================================ 4.2
    deg("ön denemelerde bu kümelerin eğitime eklenmesinin doğrulama/sınama başarımını iyileştirmediği gözlenmiştir (Tablo 21).",
        "ön denemelerde bu kümelerin eğitime eklenmesinin doğrulama/sınama başarımını iyileştirmediği gözlenmiştir "
        "(bu ön denemeler ayrıca tablo olarak raporlanmamıştır).")
    deg("Sınama kümesinde hem hasarlı hem hasarsız (negatif) karo bulunması,",
        "IoU_hasar 91 deprem sonrası sınama karosunda hesaplanmış, 41 negatif karo ise yanlış alarm sayımında ayrıca "
        "kullanılmıştır (Dış sınamalar). Sınama kümesinde hem hasarlı hem hasarsız (negatif) karo bulunması,")
    deg("Aynı gerekçeyle doğruluk ve özgüllük bu sürümde ayrıca raporlanmamıştır; ön denemelerde bu iki ölçüt de 0,98'in "
        "üzerinde yoğunlaşarak mimarileri ayırt etmemiştir.",
        "Aynı gerekçeyle doğruluk raporlanmamış, özgüllük ise yalnızca bütünlük için verilmiştir; iki ölçüt de 0,98'in "
        "üzerinde yoğunlaşarak mimarileri ayırt etmemektedir.")
    deg("aynı veri ve tarifle Türkiye 2023 depremleri için bildirilen bir literatür değeriyle (0,394 [22])",
        "aynı kodlayıcıyla Türkiye 2023 depremleri için farklı bir veri kümesinde bildirilen literatür değeriyle "
        "(U-Net/ResNet-18, 0,394 [22])")
    deg("SegFormer-B4 (esas alınan model) için Birleşik->Hatay", "SegFormer-B4 için Birleşik->Hatay")

    # ================================================================ 4.3 EMSR gerekcesi
    deg("Bu bir denemedir: tek etiketleyici, tek mahalle, küçük örneklem; sayılar bu sınırlarla okunmalıdır.",
        "Bu bir denemedir: tek etiketleyici, tek mahalle, küçük örneklem; sayılar bu sınırlarla okunmalıdır. Dış bir "
        "ürün olan Copernicus EMSR648 hasar derecelendirmesi [8] yer gerçeği olarak kullanılmamıştır: EMSR yolun "
        "kendisinin bozulmasını (köprü çökmesi, kaplama yarılması) derecelendirir, bu çalışmanın sorusu ise yanındaki "
        "binanın molozunun yolu kapatmasıdır. Aynı Defne/Elektrik kutusunda EMSR 62 yolun 62'sini \"görünür hasar yok\" "
        "olarak işaretlerken model 12 kapalı kenar bulmaktadır; bu tanım farkı nedeniyle EMSR'ın cetvel olarak "
        "kullanılması sistemi yanlış sebepten değerlendirirdi.")
    deg("ölçüt kapalı (molozun parçayı belirgin biçimde örtmesi), açık (sürekli geçilebilir koridor) ve kararsız (gölge/çözünürlük) olarak verilmiştir.",
        "ölçüt kapalı (molozun parçayı belirgin biçimde örtmesi), açık (sürekli geçilebilir koridor) ve kararsız "
        "(gölge, çözünürlük ya da OSM çizgisinin görünür sokağı izlememesi) olarak verilmiştir.")
    deg("görüntü depremden üç gün sonrasına ait olduğundan bir miktar açma çalışması başlamış olabilir; yalnız yukarıdan görünen moloz yakalanır.",
        "görüntü depremden üç gün sonrasına ait olduğundan bir miktar açma çalışması başlamış olabilir; yalnız "
        "yukarıdan görünen moloz yakalanır; modelin 12 kapanmasından yalnız 5'i uçuş izinin dolu olduğu yere düşer; "
        "kararsızlar hata sayılırsa kesinlik 3/5 = %60'a iner.")

    # ================================================================ 4.5
    deg("Kamera fotoğrafları 6 Şubat 2023 depreminin lisanslı gerçek kareleridir (künye ve SHA-256 ile kayıtlı);",
        "Kamera fotoğrafları 6 Şubat 2023 depremlerinin ardından (Şubat–Ağustos 2023) çekilmiş, Wikimedia Commons'ta "
        "kamu malı ya da CC BY-SA 4.0 lisansıyla yayımlanmış gerçek karelerdir (künye ve SHA-256 ile kayıtlı; çoğu "
        "Hatay/Antakya) ve gösterilen mahallenin fotoğrafı değildir;")
    yaz("[Toplanma alanı yönlendirmesi tatbikat zincirinde",
        "Toplanma alanı yönlendirmesi tatbikat zincirinde koşulmamış ve bu değerlendirmede raporlanmamıştır.")
    yaz("Metin kanalında bir eğilim gözlenmiştir.",
        "Metin kanalında bir eğilim gözlenmiştir. On sekiz ihbar metninden dokuzu yazarın niyetiyle örtüşmüş; "
        "örtüşmeyen dokuzun sekizi 'Acil Yardım İhtiyacı' sınıfına kaymıştır. Bu bir doğruluk ölçüsü değildir ('beklenen' "
        "etiket tek kişinin yargısıdır ve 'Enkaz Bildirimi' ile 'Acil Yardım İhtiyacı' sınırı gerçekten bulanıktır). "
        "Daha önemlisi, ölçüm modelin değil hattın davranışını göstermektedir: on sekiz çıktının on ikisini model değil "
        "anahtar sözcük kurtarma katmanı belirlemiştir. Acil Yardım'a kayan sekiz metnin beşi kurtarmadan (\"ekip "
        "lazım\", \"ambulans\", \"arama kurtarma\" ve iki metinde \"araç\" sözcüğü içindeki \"aç\" alt dizgisi), üçü "
        "ise Su, Elektrik ve Isınma alt kategorilerini Acil Yardım'a bağlayan eşleme tasarımından gelmektedir (Bölüm "
        "3.6, Tablo 9). Modelin 13 etiketli ham çıktısının başarımı bu ölçümle değerlendirilemez; F₁ iddiası için 13 "
        "etiket düzeyinde etiketli bir küme gereklidir.")

    # ================================================================ 4.10 / 4.11
    deg("Üçüncüsü, bu bulgu Bölüm 4.5'teki yanlış negatif gözlemiyle örtüşmektedir. Senaryo koşumunda sisteme giren altı "
        "metin ihbarının ikisi \"alakasız\" kategorisine atanmıştı ve her ikisi de medikal yardım çağrısıydı.",
        "Üçüncüsü, bu bulgu önceki sürümün senaryo koşumundaki yanlış negatif gözlemiyle örtüşmektedir: sisteme giren "
        "altı metin ihbarının ikisi \"alakasız\" kategorisine atanmıştı ve her ikisi de medikal yardım çağrısıydı.")
    deg("yedi günlük kayan pencere istatistikleri", "yedi olaylık kayan pencere istatistikleri")
    deg("(yedi günlük ortalama büyüklük)", "(son yedi olayın ortalama büyüklüğü)")
    deg("duyarlılıkları Tablo 21'de ölçülmüş ve raporlanmıştır ('Risk skoru ağırlıkları' satırı).",
        "duyarlılıkları yalnızca kaba bir taramayla ölçülmüştür: ağırlıkların ±0,1 ızgarası (9 kombinasyon), önceki "
        "sürümün sekiz Hatay/Kahramanmaraş konumunda kategori değişimini konuma göre 0–5/8 aralığında bulmuş, en "
        "duyarlı eksen fay ağırlığı olmuştur (konumların 7/8'i F_fay = 1,0 tavanındadır).")

    # ================================================================ 4.12
    deg("Şekil 21–24, sırasıyla yer seviyesi hasar tespiti, Türkçe afet metni analizi, deprem risk skorlaması ve uydu "
        "tabanlı yol hasarı/güzergâh modüllerinin gerçek çıktı görünümlerini sunmaktadır.",
        "Şekil 21–24, sırasıyla yer seviyesi hasar tespiti, deprem risk skorlaması, Türkçe afet metni analizi ve uydu "
        "tabanlı yol hasarı/güzergâh modüllerinin arayüz görünümlerini sunmaktadır. Şekil 21–23 modellerin gerçek "
        "çıktısını gösterir; Şekil 24 ise arayüzün yerleşik örnek veri (simülasyon) kipinde çekilmiştir ve gerçek bir "
        "analiz çıktısı değildir.")
    deg("Antakya üzerinde açık (yeşil) ve hasarlı/kapalı (kırmızı) yol ağı ile yıkım yoğunluğu ısı haritası katmanı.",
        "Antakya üzerinde açık (yeşil) ve hasarlı/kapalı (kırmızı) yol ağı ile yıkım yoğunluğu ısı haritası katmanı. "
        "Arayüzün örnek veri kipiyle çekilmiştir (simüle veri; gerçek analiz çıktısı değildir).")

    # ================================================================ 5
    sil("[Sentetik kapanma üreticisi bu sürümde kullanılmamaktadır;")
    deg("Bu çalışmanın hedeflediği fark, kapanma katmanının insan etiketlemesi yerine model çıktısından türetilmesidir; "
        "ancak bu hedef, Bölüm 5.3(a)'da belirtildiği üzere bu sürümde gerçekleştirilmemiştir.",
        "Bu çalışmanın farkı, kapanma katmanının insan etiketlemesi yerine model çıktısından türetilmesidir: tatbikattaki "
        "bütün kapanmalar kendi hasar modelimizin drone ortofotosu üzerindeki çıktısından gelmektedir; ancak katmanın "
        "doğruluğu yalnızca bir mahallede ön denemeyle sınanmıştır (Bölüm 4.3, 5.3(a)).")
    deg("öneri, kapanma katmanı gerçek çıkarımdan türetilene kadar tamamlanmış sayılamaz.",
        "öneri, kapanma katmanının doğruluğu daha geniş ve çok etiketleyicili bir yer gerçeğiyle sınanana kadar "
        "tamamlanmış sayılamaz.")
    yaz("(c) Metin sınıflandırıcısı, altı ihbarın ikisini eleyerek",
        "(c) Metin kanalında 13 alt kategoriyi beş üst kategoriye indiren katman modelin çıktısını bozmaktadır: anahtar "
        "sözcük kurtarması modelin kararından önce çalışıp onu koşulsuz ezmekte, iki harfli anahtar sözcükler alt "
        "dizgi olarak başka sözcüklerde de eşleşmekte ve gösterilen güven değeri gösterilen kategoriye ait "
        "olmamaktadır. Tatbikattaki 9/18 örtüşme bu katmanın davranışını yansıtır (on sekiz çıktının on ikisini "
        "anahtar sözcük belirlemiştir). Sınıflandırıcının beş üst kategori üzerindeki kesinlik, duyarlılık ve F₁ "
        "değerleri henüz ölçülmemiştir.")
    yaz("(d) Metin zincirinin yanıt süresinin bileşenlerine ayrıştırılması",
        "(d) Metin zincirinin yanıt süresi bileşenlerine ayrıştırılmamıştır. Tam hatta coğrafi kodlama iki metinde 3,3 "
        "ve 4,4 saniyeye çıkmıştır (önbellek ıskası, ağ çağrısı); bu nedenle tatbikat coğrafi kodlamayı çağırmaz. "
        "Ölçümler tek bir dizüstü bilgisayarda, grafik işlemci hızlandırması olmadan yapılmıştır.")
    deg(" Tatbikattaki görevlerin tamamı mahalle içindedir; mahalleler veya iller arası sevk sınanmamıştır. Görevler aynı "
        "mahalle ağını paylaştığından tam bağımsız değildir.",
        " Tatbikattaki görevlerin tamamı mahalle içindedir; mahalleler veya iller arası sevk sınanmamıştır.")
    deg("(k) Risk skorlaması bileşeni (Denklem 7) bu değerlendirmede kullanılmamıştır; ağırlıklar uzman tanımlıdır, "
        "duyarlılıkları ve olasılık kalibrasyonu ölçülmemiştir.",
        "(k) Risk skorlaması bileşeni (Denklem 7) bu değerlendirmede kullanılmamıştır; ağırlıklar uzman tanımlıdır, "
        "duyarlılıkları yalnızca sekiz konumda kaba bir taramayla ölçülmüş (Bölüm 4.11), olasılık kalibrasyonu "
        "yapılmamıştır.")
    deg(" Metin sınıflandırmada 13 alt kategorinin 5 üst kategoriye indirgenmesi modelin cevabını bozmakta, aciliyet "
        "taşıyan metinleri tek bir sınıfa çekmektedir.",
        " Metin kanalındaki kategori indirgeme sorunu (c)'de ayrıca ele alınmıştır.")
    deg("(r) Yol kapanması bu sürümde ayrı bir yol hasarı sınıfıyla değil, genel hasar bölütleme modelinin (KATE-CD + "
        "MEDBFS bina hasarı verisiyle eğitilmiş SegFormer)",
        "(r) Yol kapanması bu sürümde ayrı bir yol hasarı sınıfıyla değil, genel hasar bölütleme modellerinin (KATE-PD + "
        "KATE-CD bina hasarı verisiyle eğitilmiş SegFormer MiT-B0/B2 ve U-Net MiT-B4)")
    deg("Dördüncüsü, sevk katmanının turların döndürülmesinden gerçek bir rol-farkında atama eniyilemesine ve iplik-güvenli "
        "görev kilitlemenin üretim REST API'sine (HTTP 409) genişletilmesidir.",
        "Dördüncüsü, sevk katmanının turların döndürülmesinden gerçek bir rol-farkında atama eniyilemesine geçirilmesi "
        "ve üretim API'sindeki görev kilidinin (HTTP 409) yetki denetimiyle birlikte tatbikata bağlanmasıdır.")
    deg("metin sınıflandırmadaki 13→5 kategori indirgemesinin düzeltilmesidir.",
        "metin sınıflandırmadaki anahtar sözcük kurtarma katmanının modelin kararını ezmeyecek biçimde düzeltilmesidir.")

    # Tablo 24
    t24 = doc.tables[26]
    for row in t24.rows:
        k = row.cells[0].text
        if k.startswith('Risk skorlaması'):
            D.hucre_yaz(row.cells[3], row.cells[3].text.replace('(Tablo 22/23;', '(Tablo 23;'))
        elif k.startswith('Çok mimarili kıyaslama — yol'):
            D.hucre_yaz(row.cells[1], 'IoU_hasar (on mimari)')
        elif k.startswith('Çok mimarili kıyaslama — bina'):
            D.hucre_yaz(row.cells[1], 'IoU_hasar (on mimari)')
            D.hucre_yaz(row.cells[2], 'KATE-PD + KATE-CD [56]; dış sınama MEDBFS [57], QQB [67], HOTOSM [68]')
            D.hucre_yaz(row.cells[3], 'Ölçülmüştür (Tablo 13); kodlayıcı sabitlemesi ve mekânsal blok bölmesi '
                        'karşılanmış, çoklu tohum tekrarıyla yenilenecektir')

    # ================================================================ 6
    deg("Bölge tabanlı çizge mimarisi sayesinde sistem, il ölçeğinde tek bir çizgenin kurulamadığı koşullarda ilçe "
        "ölçeğinde işletilebilmektedir.",
        "Bölge tabanlı çizge mimarisi sayesinde sistem, il ölçeğinde tek bir çizgenin kurulamadığı koşullarda ilçe ve "
        "mahalle ölçeğinde işletilebilmektedir.")
    yaz("Ölçülemeyenler de aynı açıklıkla belirtilmelidir.",
        "Ölçülemeyenler de aynı açıklıkla belirtilmelidir. Yer seviyesi tespit başarımı (mAP) ve bölgeler arası "
        "genelleme bu sürümde ölçülmemiş; kapanma katmanının yer gerçeğine göre değerlendirilmesi yalnızca bir "
        "mahallede, tek etiketleyiciyle ve ön deneme olarak yapılabilmiştir. İlgili tablolarda (Tablo 15, Tablo 24) "
        "kalan hücreler [ÖLÇÜLECEK] işaretli bırakılmış; hiçbir hücre kestirimle doldurulmamıştır. Çok mimarili "
        "bölütleme kıyaslaması, metin sınıflandırma ve risk skorlaması ölçülmüş; ancak üçü de hedef protokolden ayrılan "
        "koşullarda üretildiği için (bölütlemede yol yerine bina hasarı görevi ve tek tohumlu koşum, metinde beş üst "
        "kategori yerine çok etiketli alt kategori kümesi, risk skorlamasında sabit ve kısa pencereli basit bir "
        "kümelenmeden arındırma yöntemi) nihai değerler değil ön bulgular olarak sunulmaktadır. Görev tabanlı "
        "değerlendirme üç model ve üç mahallede uygulanmıştır; mahalleler on bir ölçülen alandan, kıyas kollarının "
        "ayrıştığı alanlar olarak seçilmiştir. Metin kanalında anahtar sözcük kurtarma katmanının modelin kararını "
        "ezdiği gözlenmiştir; bu, kanal çeşitliliğinin ve katmanlar arası denetimin bir emniyet gereksinimi olduğunu "
        "düşündürmektedir.")
    deg("Protokolün ilk uygulaması, bugünkü bölütleme başarımının bu zincirde kazanç üretmediğini göstermiş ve algılama "
        "katmanının hangi yönde iyileştirilmesi gerektiğini ortaya koymuştur: küçük yol ağlarında yanlış kapanmalar, "
        "kaçırılan kapanmalardan daha pahalıdır.",
        "Protokolün ilk uygulaması, kapanmayı önceden bilmenin kazancının bölütleme başarımıyla değil kapanmaların ekip "
        "rotalarına düşmesiyle belirlendiğini göstermiş ve algılama katmanının hangi yönde iyileştirilmesi gerektiğini "
        "ortaya koymuştur: bugünkü model yanlış alarm üretmemekte, ancak gerçek kapanmaların çoğunu kaçırmaktadır; "
        "öncelik duyarlılığın artırılmasıdır.")
    deg("Yanlış negatif gözlemi, gecikme değişkenliği ve model tabanlı kapanmaların bugün kazanç üretmemesi birlikte "
        "değerlendirildiğinde,",
        "Metin kanalındaki kategori hataları, gecikme değişkenliği ve kapanma katmanının düşük duyarlılığı birlikte "
        "değerlendirildiğinde,")

    # ================================================================ arka bolumler
    deg("Help.NGO ve Portugese GRU ekiplerine teşekkür eder.", "Help.NGO ekibine ve kamera kanalında kullanılan "
        "fotoğrafları Wikimedia Commons'ta açık lisansla paylaşan fotoğrafçılara teşekkür eder.")
    deg("Senaryo kapsamındaki afetzede, ekip ve karargâh kayıtları sentetik olarak üretilmiştir; gerçek kişilere ait veri "
        "içermemektedir.",
        "Tatbikattaki afetzede ve ekip kayıtları ile ihbar metinleri sentetik olarak üretilmiştir; gerçek kişilere ait "
        "veri içermemektedir. Kamera kanalında kullanılan fotoğraflar Wikimedia Commons'ta kamu malı ya da CC BY-SA 4.0 "
        "lisansıyla yayımlanmış deprem fotoğraflarıdır.")
    deg("Bölütleme modellerinin eğitildiği iki küme de açık erişimlidir: KATE-CD [56] Hugging Face üzerinden "
        "yayımlanmaktadır (veri kartında lisans belirtilmemiştir); MEDBFS [57] Zenodo üzerinden CC BY-NC 4.0 lisansıyla "
        "sunulmaktadır (doi:10.5281/zenodo.11647956).",
        "Bölütleme modellerinin eğitildiği KATE-PD ve KATE-CD kümeleri [56] Hugging Face üzerinden yayımlanmaktadır "
        "(veri kartlarında lisans belirtilmemiştir); dış sınamada kullanılan MEDBFS [57] Zenodo üzerinden CC BY-NC 4.0 "
        "lisansıyla sunulmaktadır (doi:10.5281/zenodo.11647956). Kamera kanalındaki fotoğraflar Wikimedia Commons "
        "üzerinden kamu malı ya da CC BY-SA 4.0 lisansıyla erişilebilir.")
    deg("Canlı sistemin çalışma anındaki uydu karosu getirme bileşeni varsayılan olarak Google karo sunucusundan görüntü "
        "çekmektedir; bu, eğitim kaynaklarından ayrı bir kaynaktır ve raporlanan sonuçların parçası değildir.",
        "Canlı sistemin çalışma anındaki uydu karosu getirme bileşeni varsayılan olarak Google karo sunucusundan görüntü "
        "çekmektedir; bu, eğitim kaynaklarından ayrı bir kaynaktır ve raporlanan sonuçların parçası değildir. "
        "Tatbikattaki bütün analizler OpenAerialMap karolarıyla yapılmıştır.")

    # Kaynak [56] ve [64]
    deg("Veri kümesi: https://huggingface.co/datasets/CSCRS/kate-cd (Erişim tarihi: 16.09.2026).",
        "Veri kümeleri: https://huggingface.co/datasets/CSCRS/kate-cd ve https://huggingface.co/datasets/CSCRS/kate-pd "
        "(Erişim tarihi: 16.09.2026).")
    yaz("[64] Help.NGO,",
        "[64] Help.NGO, \"Hatay - Defne - Elektrik Mah.\"; \"Islahiye, Turkey - Feb 15, 2023\"; \"Nurdagi - Feb 15, 2023 "
        "(Warning Low Quality)\", İnsansız hava aracı ortofotoları, OpenAerialMap, görüntü kimlikleri "
        "63eb78a0ca43600005f4d920, 63eef9a4525f0700077ed49a, 63ef697b525f0700077ed4a2, lisans CC BY 4.0, "
        "https://map.openaerialmap.org (Erişim tarihi: 27.09.2026), (2023).")


def qn_tr():
    from docx.oxml.ns import qn
    return qn('w:tr')


def qn_t():
    from docx.oxml.ns import qn
    return qn('w:t')
