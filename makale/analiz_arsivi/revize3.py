# -*- coding: utf-8 -*-
"""Revizyon 3 (27 Eylul 2026): bos tablolarin metne cevrilmesi + tablo
yeniden numaralama, R0b duzeltmesi, yeni semalar (Sekil 2, 4, 6) ve Sekil 20,
kaynakca DOI duzeltmeleri. Yedekten baslar, yeniden calistirilabilir.
Kullanim: python revize3.py [--deneme]"""
import re
import shutil
import sys

sys.path.insert(0, r"C:\Users\utkuc\AppData\Local\Temp\claude\D--quakemind-final-QuakeMind\4779ac37-74b5-4aeb-8690-74d7792eedc5\scratchpad")
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from yardimci import Duzenleyici

SRC = r"D:\quakemind_final\QuakeMind\makale\arsiv\QuakeMind_Makale_YEDEK_pre_revize3_2026-09-27.docx"
DST = r"D:\quakemind_final\QuakeMind\makale\QuakeMind_Makale_Taslak_2026-09-16_sekillieklendi_v3_LOCAL_2026-09-21.docx"
if "--deneme" in sys.argv:
    DST = r"C:\Users\utkuc\AppData\Local\Temp\claude\D--quakemind-final-QuakeMind\4779ac37-74b5-4aeb-8690-74d7792eedc5\scratchpad\deneme3.docx"
SEK = r"D:\quakemind_final\QuakeMind\QuakeMindBackend\benchmark\makale_sekilleri"
RISK = r"D:\quakemind_final\QuakeMind\model_benchmarks\05_risk_skoru\local_run\results"

shutil.copyfile(SRC, DST)
doc = Document(DST)
D = Duzenleyici(doc)
deg, yaz, sil = D.degistir, D.yaz, D.sil


def tablo_sil_basligiyla(baslik_sub):
    cap = D.bul(baslik_sub)._p
    el = cap.getnext()
    while el is not None and el.tag != qn("w:tbl"):
        el = el.getnext()
    el.getparent().remove(el)
    cap.getparent().remove(cap)
    D.log.append("TABLO SIL " + baslik_sub[:40])


# =============================================================== 1) Tablo 5 -> metin
deg("yol hasarı eksenindeki uygulaması ise Tablo 5'te künyesi verilecek kümenin oluşturulmasına bağlı olarak",
    "yol hasarı eksenindeki uygulaması ise yol hasarını piksel düzeyinde etiketleyen bir kümenin oluşturulmasına "
    "bağlı olarak")
deg("Kümenin künyesi Tablo 5'te tanımlanan alanlarla raporlanacak; hiçbir alan kestirim değeriyle doldurulmayacaktır.",
    "Küme oluşturulduğunda kaynak görüntü ve sürümü, 512 × 512 piksellik karo sayısı, hasarlı piksel oranı, "
    "etiketleyici sayısı, [26] numaralı çalışmanın üç fazlı protokolünden uyarlanacak etiketleme kılavuzu, "
    "etiketleyiciler arası uyum (IoU) ve uzman doğrulaması raporlanacaktır; hiçbir alan kestirim değeriyle "
    "doldurulmayacaktır.")
tablo_sil_basligiyla("Tablo 5. Etiketli yol hasarı eğitim kümesinin künyesi.")
deg("Protokolün yol hasarına uygulanması, Tablo 5'te künyesi tanımlanan kümenin bu sürümde oluşturulmamış olması "
    "nedeniyle",
    "Protokolün yol hasarına uygulanması, yol hasarını etiketleyen kümenin bu sürümde oluşturulmamış olması nedeniyle")
deg("bu sürümde oluşturulmamıştır (Tablo 5); karşılaştırma", "bu sürümde oluşturulmamıştır (Bölüm 3.2.1); karşılaştırma")
deg("Tablo 5'te künyesi tanımlanan yol özgü etiketli küme oluşturulduğunda",
    "Yol hasarına özgü etiketli küme (Bölüm 3.2.1) oluşturulduğunda")

# =============================================================== 2) Tablo 15 -> metin
deg("Bölüm 4.4'teki tablo (Tablo 15) büyük ölçüde iskelet hâlindedir; kalan hücreler [ÖLÇÜLECEK] ile işaretlenmiştir.",
    "Bölüm 4.4'teki il-ayrık genelleme deneyi bu sürümde koşulmamıştır.")
yaz("Bölge içi ve bölgeler arası genelleme Tablo 15'te değerlendirilecektir.",
    "İl-ayrık eğitim (modelin yalnız bir ilin verisiyle eğitilip diğer ilde sınanması) bu sürümde koşulmamıştır; "
    "bu nedenle bölgeler arası genelleme kaybı ölçülmemiştir. Birleşik eğitimle il bazında sınama sonuçları Bölüm "
    "4.2'de verilmiştir: on mimarinin sınama IoU_hasar değerleri Hatay'da 0,267–0,425, Kahramanmaraş'ta 0,324–0,456 "
    "aralığındadır. İl-ayrık deney, eğitim kümesinin il etiketine göre yeniden bölünüp modellerin yeniden eğitilmesini "
    "gerektirmekte olup gelecek çalışmaya bırakılmıştır (Tablo 24).")
tablo_sil_basligiyla("Tablo 15. Bölge içi ve bölgeler arası bölütleme başarımı.")
sil("Bu tablonun değerleri Bölüm 3.10'da tanımlanan protokolle üretilecektir.")
deg("İlgili tablolarda (Tablo 15, Tablo 24) kalan hücreler [ÖLÇÜLECEK] işaretli bırakılmış; hiçbir hücre kestirimle "
    "doldurulmamıştır.",
    "Ölçülmeyen eksenler Tablo 24'te açıkça işaretlenmiş; hiçbir değer kestirimle doldurulmamıştır.")

# =============================================================== 3) Tablo 24 ve ilgili metin
t24 = [t for t in doc.tables if t.rows[0].cells[0].text.startswith("Değerlendirme ekseni")][0]
for row in t24.rows:
    k = row.cells[0].text
    if k.startswith("Yer seviyesi hasar tespiti"):
        D.hucre_yaz(row.cells[3], "Ölçülmedi (kutu düzeyinde etiketli saha kümesi yok)")
    elif k.startswith("Çok mimarili kıyaslama — yol"):
        D.hucre_yaz(row.cells[2], "Etiketli yol hasarı kümesi (oluşturulmadı; Bölüm 3.2.1)")
        D.hucre_yaz(row.cells[3], "Ölçülmedi (küme yok)")
    elif k.startswith("Bölgeler arası genelleme"):
        D.hucre_yaz(row.cells[3], "Ölçülmedi (il-ayrık eğitim koşulmadı; il bazında sınama Bölüm 4.2'de)")
deg("kalan eksenler (yer seviyesi hasar tespiti mAP, çok mimarili yol kıyaslaması, bölgeler arası genelleme) "
    "[ÖLÇÜLECEK] olarak işaretlidir.",
    "kalan eksenler (yer seviyesi hasar tespiti mAP, çok mimarili yol kıyaslaması, bölgeler arası genelleme) bu "
    "sürümde ölçülmemiştir.")
deg("ilgili eksen Tablo 24'te [ÖLÇÜLECEK] olarak işaretlenmiştir.", "ilgili eksen Tablo 24'te ölçülmedi olarak işaretlenmiştir.")
deg("raporlanacaktır [GELECEK ÇALIŞMA].", "gelecek çalışmada raporlanacaktır.")
deg("hatta bağlandığında ve etiketli bir sınama karosu kümesiyle gelecek çalışmada raporlanacaktır.",
    "hatta bağlandığında etiketli bir sınama karosu kümesiyle ölçülecektir.")

# =============================================================== 4) Risk: R0b duzeltmesi
deg("yalnızca fay mesafesine dayalı sıralama ve yalnızca geçmiş sismisite yoğunluğuna dayalı sıralama.",
    "yalnızca fay mesafesine dayalı sıralama ve yalnızca son yedi olayın ortalama büyüklüğüne dayalı sıralama.")
t23 = [t for t in doc.tables if t.rows[0].cells[0].text == "Model" and len(t.columns) > 5 and "Brier" in t.rows[0].cells[5].text][0]
r0b = [r for r in t23.rows if "(R0b)" in r.cells[0].text][0]
yeni = ["Naif — Yalnızca son 7 olayın ort. büyüklüğü (R0b)", "0,6143 ± 0,0520", "0,4021", "0,3424", "0,5629",
        "0,2046", "0,1616", "–"]
for c, v in zip(r0b.cells, yeni):
    D.hucre_yaz(c, v)
rf = [r for r in t23.rows if r.cells[0].text.startswith("Rastgele orman")][0]
rf._tr.addnext(r0b._tr)  # ROC-AUC sirasina gore RO'dan sonra
D.log.append("TABLO 23 R0b satiri guncellendi ve tasindi")
deg("öğrenilmiş modeller arasında en düşük değer olan LightGBM'in 0,3592'si ise yalnızca 0,061 üzerindedir. "
    "Yığınlanmış topluluk (0,3127) ve naif sismisite tabanı (0,2920) temel orana neredeyse eşit kalmaktadır.",
    "tekil öğrenilmiş modeller arasında en düşük değer olan LightGBM'in 0,3592'si ise yalnızca 0,061 üzerindedir. "
    "Yığınlanmış topluluk (0,3127) ve fay mesafesi tabanı (0,3074) temel orana neredeyse eşit kalmaktadır.")
deg("ROC-AUC, dengeli doğruluk ve MCC ölçütlerinin üçünde de en yüksek değeri lojistik regresyon vermektedir "
    "(sırasıyla 0,6370, 0,5784, 0,1551); CatBoost bu üç ölçütte de dördüncü sıradadır (0,5928, 0,5271, 0,0753).",
    "Öğrenilmiş modeller arasında ROC-AUC, dengeli doğruluk ve MCC ölçütlerinin üçünde de en yüksek değeri lojistik "
    "regresyon vermektedir (sırasıyla 0,6370, 0,5784, 0,1551); CatBoost bu ölçütlerde sırasıyla dördüncü, üçüncü ve "
    "beşinci sıradadır (0,5928, 0,5271, 0,0753).")
yaz("Beşincisi, naif tabanlarla karşılaştırma",
    "Beşincisi, naif tabanlarla karşılaştırma öğrenilmiş modellerin kattığı değerin sınırlı olduğunu göstermektedir. "
    "Yalnızca fay mesafesine dayalı sıralama (R0a) PR-AUC = 0,3074 (temel oranın yalnızca 0,009 üzerinde) ve MCC = "
    "0,0305 vermektedir; yakınlık sezgisi tek başına güçlü bir taban değildir. Buna karşılık yalnızca son yedi olayın "
    "ortalama büyüklüğüne dayalı sıralama (R0b) ROC-AUC = 0,6143 ± 0,0520, PR-AUC = 0,4021 ve MCC = 0,1616 "
    "vermektedir: MCC'de bütün öğrenilmiş modelleri, ROC-AUC ve PR-AUC'de ise esas alınan CatBoost'u geçmektedir. "
    "ROC-AUC'de yalnız lojistik regresyon, çok katmanlı algılayıcı ve rastgele orman bu tabanın üzerindedir ve en "
    "büyük fark (0,023) kat standart sapmalarının içinde kalmaktadır. Dolayısıyla bu kümede öğrenilmiş modellerin tek "
    "öznitelikli bir yakın dönem etkinlik sezgisine kattığı değer sınırlıdır. Önceki sürümde R0b, yedi olaylık "
    "penceredeki olay sayısıyla koşulmuştu; bu öznitelik 13 609 kaydın hepsinde 7'ye eşit (sabit) olduğundan taban "
    "tanım gereği rastgele düzey (ROC-AUC = 0,5000) veriyordu. Bu sürümde R0b tanımına uygun öznitelikle, aynı beş "
    "katlı ileri yürüyen bölmeyle yeniden ölçülmüş; aynı düzenekte R0a Tablo 23'teki değerleri birebir yeniden "
    "üretmiştir.")
deg("CatBoost'un üretim sisteminde kullanılmaya devam etmesi bu karşılaştırmadan bağımsız, önceden alınmış bir mühendislik "
    "kararıdır; bu bölümün bulgusu, kararın bu veri kümesinde ölçülen başarımla desteklenmediğidir.",
    "CatBoost'un üretim sisteminde kullanılmaya devam etmesi bu karşılaştırmadan bağımsız, önceden alınmış bir mühendislik "
    "kararıdır; bu bölümün bulgusu, kararın bu veri kümesinde ölçülen başarımla desteklenmediğidir (tek öznitelikli R0b "
    "tabanı dahi CatBoost'u ROC-AUC, PR-AUC ve MCC'de geçmektedir).")
deg("Yedi risk skorlama modelinin karşılaştırması Tablo 23'te verilmektedir.",
    "Yedi risk skorlama modelinin ve iki naif tabanın karşılaştırması Tablo 23'te verilmektedir.")

# =============================================================== 5) mobil istemci geri
deg("bu katmanı kullanan web istemcilerinden oluşmaktadır;",
    "bu katmanı kullanan web ve mobil (Flutter) istemcilerinden oluşmaktadır;")

# =============================================================== 6) semalar ve Sekil 20
D.resim_koy(D.basliktan_onceki_gorsel("Şekil 2. QuakeMind üç katmanlı"), SEK + r"\sekil_mimari.png", 15.5)
D.resim_koy(D.basliktan_onceki_gorsel("Şekil 4. Örtüşmeli karolama"), SEK + r"\sekil_karolama.png", 15.5)
D.resim_koy(D.basliktan_onceki_gorsel("Şekil 6. Üç düzeyli değerlendirme"), SEK + r"\sekil_protokol.png", 15.5)
cap20 = D.bul("Şekil 20. Esas alınan CatBoost modelinde")._p
gorseller = []
el = cap20.getprevious()
while el is not None and len(gorseller) < 2:
    if el.tag == qn("w:p") and el.findall(".//" + qn("w:drawing")):
        gorseller.append(Paragraph(el, doc))
    el = el.getprevious()
alt, ust = gorseller  # en yakin olan alttaki (SHAP)
D.resim_koy(ust, RISK + r"\shap_top_importance.png", 12.7)
D.resim_koy(alt, RISK + r"\shap_bottom_summary.png", 8.5)
deg("Şekil 20. Esas alınan CatBoost modelinde öznitelik önemi (üstte) ve SHAP değerlerinin özeti (altta).",
    "Şekil 20. Esas alınan CatBoost modelinde öznitelik önemi (üstte) ve SHAP değerlerinin özeti (altta; son kat). "
    "\"Son 7 olay\" öznitelikleri yedi günlük değil yedi olaylık kayan pencereden hesaplanır; pencere doluluğu her "
    "kayıtta 7 olduğundan etkisi sıfırdır.")

# =============================================================== 7) kaynakca
deg("10-17, doi:10.1109/CVPRW.2019.00332, (2019).",
    "10-17, https://openaccess.thecvf.com/content_CVPRW_2019/html/cv4gc/Gupta_Creating_xBD_A_Dataset_for_Assessing_"
    "Building_Damage_from_Satellite_CVPRW_2019_paper.html (Erişim tarihi: 27.09.2026), (2019).")
deg("doi:10.1109/TGRS.2024.3437470", "doi:10.1109/TGRS.2024.3438290")
deg("31: 6638-6648, doi:10.5555/3327757.3327770, (2018).",
    "31: 6638-6648, arXiv:1706.09516, doi:10.48550/arXiv.1706.09516, (2018).")
deg("30: 3146-3154, doi:10.5555/3294996.3295074, (2017).",
    "30: 3146-3154, https://papers.nips.cc/paper_files/paper/2017/hash/6449f44a102fde848669bdd9eb6b76fa-Abstract.html "
    "(Erişim tarihi: 27.09.2026), (2017).")
# bugun erisimi dogrulanan adresler
dogrulanan = ["[1] ", "[2] ", "[8] ", "[30] ", "[40] ", "[44] ", "[45] ", "[47] ", "[56] ", "[58] ", "[60] ", "[63] "]
for bas in dogrulanan:
    p = [q for q in D.paragraflar() if q.text.startswith(bas)][0]
    while "Erişim tarihi: 16.09.2026" in p.text:
        D._run_degistir(p, "Erişim tarihi: 16.09.2026", "Erişim tarihi: 27.09.2026")
p68 = [q for q in D.paragraflar() if q.text.startswith("[68] ")][0]
if "Erişim tarihi: 24.09.2026" in p68.text:
    D._run_degistir(p68, "Erişim tarihi: 24.09.2026", "Erişim tarihi: 27.09.2026")
D.log.append("ERISIM tarihleri: dogrulanan 13 kaynak 27.09.2026")


# =============================================================== 8) tablo yeniden numaralama
def yeni_no(n):
    if n < 5:
        return n
    if n in (5, 15):
        raise ValueError("silinen tabloya atif kaldi: Tablo %d" % n)
    return n - 1 if n < 15 else n - 2


def paragraf_konumsal(p, bas, son, yeni):
    runs = list(p.runs)
    pos = 0
    ilk = None
    for k, r in enumerate(runs):
        L = len(r.text)
        if ilk is None and pos + L > bas:
            ilk = (k, bas - pos)
        if pos + L >= son:
            kb, ob = ilk
            ks, os_ = k, son - pos
            if kb == ks:
                r.text = r.text[:ob] + yeni + r.text[os_:]
            else:
                runs[kb].text = runs[kb].text[:ob] + yeni + runs[ks].text[os_:]
                for j in range(kb + 1, ks + 1):
                    runs[j].text = ""
            return
        pos += L


desen = re.compile(r"Tablo(\s+)(\d+)(?:(\s*[–\-]\s*)(\d+))?")
sayac = 0
for p in list(D.paragraflar()):
    t = p.text
    eslesmeler = list(desen.finditer(t))
    for m in reversed(eslesmeler):
        a = yeni_no(int(m.group(2)))
        yeni = "Tablo" + m.group(1) + str(a)
        if m.group(4):
            yeni += m.group(3) + str(yeni_no(int(m.group(4))))
        if yeni != m.group(0):
            paragraf_konumsal(p, m.start(), m.end(), yeni)
            sayac += 1
D.log.append("TABLO NUMARALARI kaydirildi: %d atif" % sayac)

# yeni tablolarin bicimi ve kalinlik
D.kalinlik_duzelt()
doc.save(DST)
print("\n".join(D.log))
print("TOPLAM", len(D.log), "->", DST)
