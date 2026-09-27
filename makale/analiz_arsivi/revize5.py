# -*- coding: utf-8 -*-
"""Revizyon 5: Tablo 21 (risk) kalin hucrelerinin olculen en iyi degerlere
gore duzeltilmesi ve satir puntolarinin esitlenmesi."""
import re
import shutil
import sys

sys.path.insert(0, r"C:\Users\utkuc\AppData\Local\Temp\claude\D--quakemind-final-QuakeMind\4779ac37-74b5-4aeb-8690-74d7792eedc5\scratchpad")
from docx import Document
from docx.shared import Pt
from yardimci import Duzenleyici

SRC = r"D:\quakemind_final\QuakeMind\makale\arsiv\QuakeMind_Makale_YEDEK_pre_revize5_2026-09-27.docx"
DST = r"D:\quakemind_final\QuakeMind\makale\QuakeMind_Makale_Taslak_2026-09-16_sekillieklendi_v3_LOCAL_2026-09-21.docx"
shutil.copyfile(SRC, DST)
doc = Document(DST)
D = Duzenleyici(doc)

t = [t for t in doc.tables if t.rows[0].cells[0].text == "Model" and len(t.columns) > 5
     and "Brier" in t.rows[0].cells[5].text][0]


def sayi(s):
    m = re.match(r"\s*(\d+,\d+)", s)
    return float(m.group(1).replace(",", ".")) if m else None


ogrenilmis = [r for r in t.rows[1:] if not r.cells[0].text.startswith("Naif")]
# sutun: (indeks, en iyi = max mi)
sutunlar = {1: True, 2: True, 3: True, 4: True, 5: False, 6: True}
for r in t.rows[1:]:
    for c in r.cells:
        for p in c.paragraphs:
            for run in p.runs:
                run.font.size = Pt(8.5)
                run.bold = None
for ci, buyuk in sutunlar.items():
    degerler = [(sayi(r.cells[ci].text), r) for r in ogrenilmis]
    hedef = (max if buyuk else min)(degerler, key=lambda x: x[0])[1]
    for p in hedef.cells[ci].paragraphs:
        for run in p.runs:
            run.bold = True
    D.log.append("KALIN sutun %d -> %s" % (ci, hedef.cells[0].text[:20]))
D.degistir("Tablo 21. Deprem risk skorlama modellerinin karşılaştırması (n = 13 609, beş katlı ileri yürüyen bölme, "
           "pozitif oran 0,298).",
           "Tablo 21. Deprem risk skorlama modellerinin karşılaştırması (n = 13 609, beş katlı ileri yürüyen bölme, "
           "pozitif oran 0,298). Kalın değerler öğrenilmiş modeller arasında her ölçütteki en iyi değeri gösterir "
           "(Brier skorunda en düşük); naif tabanlar kalınlaştırılmamıştır.")
doc.save(DST)
print("\n".join(D.log))
