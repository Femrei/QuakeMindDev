# -*- coding: utf-8 -*-
"""Revizyon 14: kapanma yer gercegi artik INSAN etiketiyle (J-126 kapandi).

Kaynak: QuakeMindBackend/benchmark/kapanma_yer_gercegi/uyum.json
(araclar_kapanma_uyum.py; etiketler_insan_1.json + etiketler_defne_elektrik.json [Claude]).
Insan etiketi esas alinir; buyuk dil modeli ikinci gorus olarak ve etiketleyiciler arasi
uyum (Cohen kappa) ile raporlanir.
"""
import copy
import json
import shutil
import sys

sys.path.insert(0, r"C:\Users\utkuc\AppData\Local\Temp\claude\D--quakemind-final-QuakeMind\4779ac37-74b5-4aeb-8690-74d7792eedc5\scratchpad")
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm
from yardimci import Duzenleyici

UYUM = r"D:\quakemind_final\QuakeMind\QuakeMindBackend\benchmark\kapanma_yer_gercegi\uyum.json"
SRC = r"D:\quakemind_final\QuakeMind\makale\arsiv\QuakeMind_Makale_YEDEK_pre_revize14_2026-09-28.docx"
DST = r"D:\quakemind_final\QuakeMind\makale\QuakeMind_Makale_Taslak_2026-09-16_sekillieklendi_v3_LOCAL_2026-09-21.docx"
v = lambda x, b=2: f"{x:.{b}f}".replace(".", ",")

U = json.load(open(UYUM, encoding="utf-8"))
ins = next(x for k, x in U["etiketleyiciler"].items() if k.startswith("Insan"))
bdm = next(x for k, x in U["etiketleyiciler"].items() if k.startswith("Claude"))
uy = U["uyum"][0]
assert (ins["modelKapali"], ins["modelAcikOrnek"]) == ({"kapali": 4, "acik": 1, "kararsiz": 0},
                                                       {"kapali": 10, "acik": 13, "kararsiz": 5})
assert round(ins["duyarlilikKestirim"], 3) == 0.090 and round(ins["gercekKapanmaKestirim"]) == 44
assert (uy["ayniEtiket"], round(uy["kappa3"], 2), round(uy["kappa2KararsizHaric"], 2), uy["ikiliKesit"]) == \
       (14, 0.13, 0.37, 19)
IW = ins["kesinlikWilson95"]
BW = bdm["kesinlikWilson95"]
K3, K2 = v(uy["kappa3"]), v(uy["kappa2KararsizHaric"])

shutil.copyfile(SRC, DST)
doc = Document(DST)
D = Duzenleyici(doc)
R = D.degistir

# OZ / Abstract
R("Kapanma katmanının doğruluğu bir mahallede kör etiketlemeyle ayrıca sınanmış (etiketleme bir büyük dil "
  "modeliyle yapılmıştır; insan doğrulaması bekleniyor), karar verilebilir parçalarda yüksek kesinlik (3/3) ancak "
  "düşük duyarlılık (≈%11) ölçülmüştür.",
  "Kapanma katmanının doğruluğu bir mahallede bir insan etiketleyicinin kör etiketlemesiyle ayrıca sınanmış; "
  "kesinlik 4/5, duyarlılık ≈%9 ölçülmüştür (yüksek kesinlik, düşük duyarlılık).")
R("Blockage-layer accuracy was separately probed by blind labeling in one neighborhood (performed by a large "
  "language model; human verification is pending), yielding high precision (3/3) but low recall (≈11%).",
  "Blockage-layer accuracy was separately probed by blind human labeling in one neighborhood, yielding precision "
  "of 4/5 but low recall (≈9%).")

# 4.1
R("kapanma katmanının bir mahallede kör etiketlemeyle (yapay zekâ etiketleyici) ön doğrulaması",
  "kapanma katmanının bir mahallede kör insan etiketlemesiyle ön doğrulaması")

# 4.3 yontem
R("Etiketleme bu sürümde bir büyük dil modeli (Claude, Anthropic) tarafından, kesit görüntüleri üzerinden "
  "yapılmıştır; aynı kesitlerin bir insan etiketleyici tarafından etiketlenmesi ve etiketleyiciler arası uyumun "
  "raporlanması bekleyen bir iştir (etiketleme kiti ve uyum betiği depodadır).",
  "Etiketleme yazarlardan biri tarafından, 33 kesit üzerinden ve model çıktısı görülmeden yapılmıştır. Aynı "
  "kesitler ikinci görüş olarak bir büyük dil modeline (Claude, Anthropic) de etiketletilmiş ve iki etiketleyici "
  "arasındaki uyum ölçülmüştür (etiketleme sayfası ve uyum betiği depodadır).")
R("Bu bir denemedir: tek etiketleyici, tek mahalle, küçük örneklem;",
  "Bu bir denemedir: tek insan etiketleyici, tek mahalle, küçük örneklem;")
R("(Hatay / Defne – Elektrik Mah., U-Net MiT-B4; kör etiketleme, yapay zekâ etiketleyici, ön deneme).",
  "(Hatay / Defne – Elektrik Mah., U-Net MiT-B4; kör etiketleme; insan etiketleyici esas, büyük dil modeli ikinci "
  "görüş; ön deneme).")

# Tablo 13: insan + BDM sutunlari, uyum satiri
t13s = [t for t in doc.tables if t.rows[0].cells[0].text == "Ölçüt" and "Wilson" in
        " ".join(c.text for r in t.rows for c in r.cells)]
assert len(t13s) == 1 and len(t13s[0].columns) == 3
t13 = t13s[0]
g = t13._tbl.tblGrid.findall(qn("w:gridCol"))
g[1].addnext(copy.deepcopy(g[1]))
for tr in t13._tbl.findall(qn("w:tr")):
    tcs = tr.findall(qn("w:tc"))
    tcs[1].addnext(copy.deepcopy(tcs[1]))
satir = [
    ["Ölçüt", "İnsan etiketleyici (esas)", "Büyük dil modeli (ikinci görüş)", "Not"],
    ["Model KAPALI dedi → kapalı / açık / kararsız", "4 / 1 / 0", "3 / 0 / 2", "5 parça"],
    ["Model AÇIK dedi (örnek) → kapalı / açık / kararsız", "10 / 13 / 5", "6 / 14 / 8",
     "28 örnek parça (%24,8)"],
    ["Kesinlik (kararsız hariç)", f"4/5 = %80 (Wilson %{IW[0]*100:.0f}–{IW[1]*100:.0f})",
     f"3/3 = %100 (Wilson %{BW[0]*100:.0f}–{BW[1]*100:.0f})", "n = 5 ve 3; küçük örneklem"],
    ["Duyarlılık kestirimi", f"4/{ins['gercekKapanmaKestirim']:.0f} ≈ %{ins['duyarlilikKestirim']*100:.0f}",
     f"3/{bdm['gercekKapanmaKestirim']:.0f} ≈ %{bdm['duyarlilikKestirim']*100:.0f}",
     "payda: 118 parçadaki tahmini gerçek kapanma"],
    ["Yargılanamayan parça", f"{round(ins['yargilanamayan']*33)}/33 (%{ins['yargilanamayan']*100:.0f})",
     f"{round(bdm['yargilanamayan']*33)}/33 (%{bdm['yargilanamayan']*100:.0f})", "gölge / çözünürlük / geometri"],
]
for tr_, deg in zip(t13.rows, satir):
    for c, m in zip(tr_.cells, deg):
        D.hucre_yaz(c, m)
us = D.satir_ekle(t13, len(t13.rows) - 1, ["Etiketleyiciler arası uyum", "", "", "zayıf uyum"])
D.hucre_yaz(us.cells[1].merge(us.cells[2]), f"{uy['ayniEtiket']}/33 kesitte aynı etiket; Cohen κ = {K3} (üç sınıf), "
                                            f"{K2} (kararsız hariç, n = {uy['ikiliKesit']})")
for i, w in enumerate([4.6, 3.8, 3.8, 3.8]):
    t13.columns[i].width = Cm(w)
    for c in t13.columns[i].cells:
        c.width = Cm(w)
D.log.append("TABLO 13 insan + BDM + uyum")

# 4.3 yorum paragrafi
R("Model yanlış alarm üretmemektedir: kapalı dediği üç karar verilebilir parçanın üçü de gerçekten kapalıdır. "
  "Ancak çok azını bulmaktadır: mahallede tahminen 27 kapanma varken üçünü işaretlemiştir — yüksek kesinlik, "
  "düşük duyarlılık.",
  "Model az yanlış alarm üretmektedir: kapalı dediği beş parçanın dördü insan etiketinde de kapalı, biri açıktır. "
  "Ancak çok azını bulmaktadır: mahallede tahminen 44 kapanma varken dördünü işaretlemiştir — yüksek kesinlik, "
  "düşük duyarlılık.")
R("Sınırlar: tek etiketleyici ve bu etiketleyici bir yapay zekâ modelidir, insan görüşü yok; parçaların %30'u "
  "yargılanamamış;",
  f"Sınırlar: tek insan etiketleyici; ikinci görüş olarak etiketleyen büyük dil modeliyle uyum zayıftır (33 kesitin "
  f"{uy['ayniEtiket']}'ünde aynı etiket; Cohen κ üç sınıfta {K3}, kararsızlar hariç {K2}), bu da kapanmanın "
  "ortofotodan yargılanmasının öznel olduğunu gösterir; parçaların %15'i yargılanamamış;")
R("; kararsızlar hata sayılırsa kesinlik 3/5 = %60'a iner.",
  "; kararsızlar hata sayılsa da kesinlik değişmez (insan etiketinde modelin kapalı dediği parçalarda kararsız "
  "yoktur).")

# 5.1, 5.3, 5.4
R("Kapanma katmanının doğruluğu, bir mahallede kör etiketlemeyle (yapay zekâ etiketleyici) sınanmış ve yüksek "
  "kesinlik (3/3) ancak düşük duyarlılık (≈%11) bulunmuştur (Bölüm 4.3): model yanlış alarm üretmemekte ama gerçek "
  "kapanmaların çoğunu kaçırmaktadır.",
  "Kapanma katmanının doğruluğu, bir mahallede kör insan etiketlemesiyle sınanmış ve kesinlik 4/5 ancak düşük "
  "duyarlılık (≈%9) bulunmuştur (Bölüm 4.3): model az yanlış alarm üretmekte ama gerçek kapanmaların çoğunu "
  "kaçırmaktadır.")
R("tek (yapay zekâ) etiketleyiciyle ve küçük örneklemle sınanmıştır (Bölüm 4.3): yüksek kesinlik (3/3), düşük "
  "duyarlılık (≈%11).",
  f"tek insan etiketleyiciyle ve küçük örneklemle sınanmıştır (Bölüm 4.3): kesinlik 4/5, duyarlılık ≈%9; ikinci "
  f"görüş olan büyük dil modeliyle uyum zayıftır (κ = {K3}).")
R("tek (yapay zekâ) etiketleyicinin ötesinde, insan etiketleyicilerle ve etiketleyiciler arası uyum raporlanarak",
  "tek insan etiketleyicinin ötesinde, birden çok insan etiketleyiciyle ve etiketleyiciler arası uyum raporlanarak")
R("(yapay zekâ etiketleyici; eğitimdeki doğrulama bölmesi)", "(tek insan etiketleyici; eğitimdeki doğrulama bölmesi)")

# 6
R("kör etiketlemeyle (bir büyük dil modeliyle; insan doğrulaması bekleniyor) sınanması ve yüksek kesinlik (3/3) "
  "ile düşük duyarlılık (≈%11) bulunmasıdır.",
  "kör insan etiketlemesiyle sınanması ve kesinlik 4/5 ile düşük duyarlılık (≈%9) bulunmasıdır.")
R("tek (yapay zekâ) etiketleyiciyle ve ön deneme olarak yapılabilmiştir.",
  "tek insan etiketleyiciyle ve ön deneme olarak yapılabilmiştir.")
R("bugünkü model yanlış alarm üretmemekte, ancak gerçek kapanmaların çoğunu kaçırmaktadır;",
  "bugünkü model az yanlış alarm üretmekte, ancak gerçek kapanmaların çoğunu kaçırmaktadır;")

# Tablo 22
t22 = next(tb for tb in doc.tables if tb.rows[0].cells[0].text == "Değerlendirme ekseni")
for row in t22.rows:
    if row.cells[0].text.startswith("Kapanma katmanı"):
        D.hucre_yaz(row.cells[2], "Kör etiketleme (tek insan etiketleyici; ikinci görüş büyük dil modeli)")
        D.hucre_yaz(row.cells[3], f"Ön deneme (Tablo 13): kesinlik 4/5, duyarlılık ≈%9; etiketleyiciler arası "
                                  f"κ = {K3}")
D.log.append("TABLO 22 kapanma satiri")

doc.save(DST)
print("\n".join(D.log))
