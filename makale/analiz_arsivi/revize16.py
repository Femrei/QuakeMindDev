# -*- coding: utf-8 -*-
"""Revizyon 16: çoklu tohum (sınırlılık o) — üç mimaride üç tohum.

Kaynak: makale/analiz_arsivi/kaggle/sonuclar/
  birlesik_P2/<model>__tohum_42[__2].json    Tablo 12'nin kayıtları (artırım tohumlu, yeni kod)
  oturum1_coklu_tohum/P2__<model>__tohum_{7,2026}.json   U-Net, UPerNet (7, 2026), SegFormer-B4 (2026)
  oturum4_segformer_tohum7/P2__segformer_mitb4__tohum_7.json  SegFormer-B4 tohum 7, yeni kodla
Bütün kayıtlar aynı kod ayarında (artirim_tohumlu=True) olmak zorunda; betik bunu denetler.
revize15'in çıktısından (pre_revize16 yedeği) başlar. Sıralama cümleleri ölçülen değere göre kurulur.
"""
import glob
import json
import shutil
import statistics as st
import sys

sys.path.insert(0, r"C:\Users\utkuc\AppData\Local\Temp\claude\D--quakemind-final-QuakeMind\4779ac37-74b5-4aeb-8690-74d7792eedc5\scratchpad")
from docx import Document
from yardimci import Duzenleyici

KOK = r"D:\quakemind_final\QuakeMind"
SON = KOK + r"\makale\analiz_arsivi\kaggle\sonuclar"
SRC = KOK + r"\makale\arsiv\QuakeMind_Makale_YEDEK_pre_revize16_2026-09-28.docx"
DST = KOK + r"\makale\QuakeMind_Makale_Taslak_2026-09-16_sekillieklendi_v3_LOCAL_2026-09-21.docx"
v = lambda x, b=3: f"{x:.{b}f}".replace(".", ",")
oku = lambda y: json.load(open(y, encoding="utf-8"))
TOH = [42, 7, 2026]
AD = {"unet_mitb4": "U-Net", "upernet_mitb4": "UPerNet", "segformer_mitb4": "SegFormer-B4"}
TAB12 = {"unet_mitb4": 0.433, "upernet_mitb4": 0.425, "segformer_mitb4": 0.407}


def kayit(m, t):
    if t == 42:
        return oku(f"{SON}/birlesik_P2/{m}__tohum_42{'__2' if m == 'segformer_mitb4' else ''}.json")
    f = glob.glob(f"{SON}/oturum*/P2__{m}__tohum_{t}.json")
    assert len(f) == 1, (m, t, f)
    return oku(f[0])


K = {m: [kayit(m, t) for t in TOH] for m in AD}
for m, ks in K.items():
    for k, t in zip(ks, TOH):
        a = k["ayarlar"]
        assert k["model"] == m and k["tohum"] == t and a["artirim_tohumlu"] is True and a["girdi"] == "sonrasi"
        assert a["bolge_egitimi"] is None and a["donem"] == 50 and a["ogrenme_orani"] == 2e-4
        assert k["egitim"]["karo"]["egitim"] == 1013 and k["sinama_sonrasi_tumu"]["n_karo"] == 91
    assert round(ks[0]["sinama_sonrasi_tumu"]["iou_hasar"], 3) == TAB12[m]
IOU = {m: [k["sinama_sonrasi_tumu"]["iou_hasar"] for k in ks] for m, ks in K.items()}
ORT = {m: st.mean(x) for m, x in IOU.items()}
SS = {m: st.stdev(x) for m, x in IOU.items()}
SSMAX = max(SS.values())
sira = sorted(AD, key=lambda m: -ORT[m])
en, ikinci = sira[0], sira[1]
# en iyi her tohumda en yuksek mi?
her_tohumda = all(IOU[en][i] > max(IOU[m][i] for m in AD if m != en) for i in range(3))
ayrik = min(IOU[en]) > max(max(IOU[m]) for m in AD if m != en)
for m in sira:
    print(f"{m}: {[round(x, 4) for x in IOU[m]]} ort {ORT[m]:.4f} ss {SS[m]:.4f}")
print("en", en, "her tohumda", her_tohumda, "araliklar ayrik", ayrik)

ozet = ", ".join(f"{AD[m]} {v(ORT[m])} ± {v(SS[m])}" for m in AD)
if ayrik:
    iddia = (f"{AD[en]} üç tohumun her birinde öteki iki mimariden yüksektir ve tohum aralıkları örtüşmemektedir; "
             f"bu mimarinin üstünlüğü tohum oynamasına bağlanamaz.")
    iddia_o = f"{AD[en]}'in üstünlüğü üç tohumda da korunmuştur."
elif her_tohumda:
    iddia = (f"{AD[en]} üç tohumun her birinde en yüksek değeri vermiş, ancak tohum aralıkları örtüştüğünden fark "
             f"küçüktür.")
    iddia_o = f"{AD[en]} üç tohumda da en yüksektir, ancak fark tohum oynaması mertebesindedir."
else:
    iddia = ("Mimariler arasındaki fark tohum oynamasıyla aynı mertebededir ve sıralama tohuma göre değişmektedir; "
             "MiT-B4 ailesi içinde tek bir mimarinin üstünlüğü ileri sürülemez.")
    iddia_o = "Bu üç mimari arasındaki sıralama tohuma göre değiştiğinden MiT-B4 ailesi içindeki ince sıralama anlamlı değildir."

shutil.copyfile(SRC, DST)
doc = Document(DST)
D = Duzenleyici(doc)
R = D.degistir

ORT_MIN, ORT_MAX = min(ORT.values()), max(ORT.values())
ARALIK = {m: max(x) - min(x) for m, x in IOU.items()}
TABAN = (0.327, 0.339)   # Tablo 12, tek tohum
if ORT_MIN - TABAN[1] > 2 * SSMAX:
    kodlayici = (f"Kodlayıcı, kod çözücüden daha belirleyici görünmektedir: MiT-B4 ailesindeki beş mimari tek tohumda "
                 f"0,397-0,433 bandında toplanırken üç taban çizgisi 0,327-0,339 bandında kalmıştır; üç tohumla "
                 f"yinelenen üç MiT-B4 mimarisinin ortalamaları da ({v(ORT_MIN)}-{v(ORT_MAX)}) taban çizgilerinin "
                 f"üzerindedir ve fark tohumlar arası standart sapmanın (en çok {v(SSMAX)}) iki katını aşmaktadır. "
                 f"Aynı kod çözücünün (UPerNet) MiT-B4 ile üç tohum ortalamasında {v(ORT['upernet_mitb4'])}, "
                 f"ConvNeXt-S ile tek tohumda 0,327 vermesi bu gözlemi desteklemektedir.")
else:
    kodlayici = (f"Kodlayıcının kod çözücüden daha belirleyici olduğu düşünülebilir: MiT-B4 ailesindeki beş mimari tek "
                 f"tohumda 0,397-0,433 bandında toplanırken üç taban çizgisi 0,327-0,339 bandında kalmıştır. Ancak üç "
                 f"tohumla yinelenen üç MiT-B4 mimarisinin ortalamaları ({v(ORT_MIN)}-{v(ORT_MAX)}) ile taban "
                 f"çizgileri arasındaki fark, tohumlar arası standart sapmayla (en çok {v(SSMAX)}) aynı mertebededir "
                 f"ve taban çizgilerinin tohum oynaması ölçülmemiştir; bu gözlem ön bulgu olarak okunmalıdır. Aynı kod "
                 f"çözücü (UPerNet) MiT-B4 ile üç tohum ortalamasında {v(ORT['upernet_mitb4'])}, ConvNeXt-S ile tek "
                 f"tohumda 0,327 vermiştir.")
R("Belirleyici olan kod çözücü değil kodlayıcıdır: MiT-B4 ailesindeki beş mimari 0,397-0,433 bandında toplanırken, "
  "üç taban çizgisi 0,327-0,339 bandında kalmış; iki bant arasındaki fark (0,06-0,10) tek-tohum oynamasının (±0,02) "
  "belirgin biçimde üzerindedir. Aynı kod çözücünün (UPerNet) MiT-B4 ile 0,425, ConvNeXt-S ile 0,327 vermesi bu "
  "gözlemi doğrudan doğrulamaktadır.", kodlayici)

# 2) "bolutleme basarimini artirmak kazanci orantili artirmiyor" -> "sinama IoU farklari kazanci ongormuyor"
U_AR = ARALIK["unet_mitb4"]
R("bölütleme başarımını artırmanın ise bu kazancı orantılı biçimde artırmadığının gösterilmesidir.",
  "modeller arasındaki sınama IoU farklarının ise bu kazancı öngörmediğinin gösterilmesidir.")
R("while improving segmentation performance does not increase this gain proportionally.",
  "while differences in test IoU between models do not predict this gain.")
R("bölütleme başarımını artırmanın bu kazancı orantılı biçimde artırmadığı gösterilmiştir.",
  "modeller arasındaki sınama IoU farklarının bu kazancı öngörmediği gösterilmiştir.")
R("Çıkarım süresini 5,5 kat artırmak (1,91 → 10,50 sn) sınama IoU'sunu yalnız 0,051 yükseltmekte, operasyonel "
  "kazancı tek yönlü artırmamaktadır:",
  "Çıkarım süresini 5,5 kat artırmak (1,91 → 10,50 sn) sınama IoU'sunu tek tohumda yalnız 0,051 yükseltmekte ve "
  "operasyonel kazancı tek yönlü artırmamaktadır"
  + (f" (U-Net'in üç tohum arasındaki oynaması {v(U_AR)} olduğundan bu IoU farkı oynamanın içindedir)"
     if U_AR > 0.051 else "") + ":")
R("Bölütleme başarımını artırmak bu kazancı orantılı biçimde artırmamaktadır: çıkarım süresini 5,5 kat artıran en "
  "büyük model (U-Net MiT-B4, 10,50 sn) sınama IoU'sunu yalnız +0,051 yükseltmekte,",
  "Sınama IoU'sundaki farklar bu kazancı öngörmemektedir: çıkarım süresini 5,5 kat artıran en büyük model (U-Net "
  "MiT-B4, 10,50 sn) sınama IoU'sunu tek tohumda yalnız +0,051 yükseltmekte,")
R("Bölütleme başarımını artırmak operasyonel faydayı orantılı biçimde artırmamaktadır: İslahiye'de dengeli model "
  "(sınama IoU 0,409) en büyük modelden (0,433) yaklaşık dört kat fazla kazanç sağlamıştır (%33,7'e karşı %8,3).",
  "Sınama IoU'su operasyonel faydanın iyi bir göstergesi değildir: İslahiye'de dengeli model (sınama IoU 0,409) en "
  "büyük modelden (0,433) yaklaşık dört kat fazla kazanç sağlamıştır (%33,7'e karşı %8,3). Bu IoU değerleri tek "
  f"tohumludur; üç tohumla yinelendiğinde U-Net'in sınama IoU'su {v(min(IOU['unet_mitb4']))}-"
  f"{v(max(IOU['unet_mitb4']))} arasında değiştiğinden modeller arasında kesin bir bölütleme kalitesi sıralaması "
  "da kurulamamaktadır (Bölüm 4.2).")
R("İkincisi, bölütleme başarımını artırmanın bu kazancı orantılı biçimde artırmadığıdır:",
  "İkincisi, modeller arasındaki sınama IoU farklarının bu kazancı öngörmediğidir:")
R("Bu koşum, Bölüm 3.10'da tanımlanan hedef protokolden tek noktada ayrılmaktadır: tek tohumla (42) yürütülmüştür; "
  "çoklu tohum varyansı SegFormer-B4 dışındaki dokuz mimaride ölçülmemiştir. SegFormer-B4 için ikinci bir tohumla (7) "
  "ölçülen sınama IoU_hasar farkı ±0,02 mertebesindedir; MiT-B4 ailesi içindeki mimariler arası farklar (0,397-0,433 "
  "aralığında, yaklaşık 0,03 mertebesinde) bu bandın içinde kalabileceğinden, mimariler arası ince sıralamanın "
  "istatistiksel anlamlılığı bu sürümde gösterilmemiştir.",
  "Bu koşum, Bölüm 3.10'da tanımlanan hedef protokolden tek noktada ayrılmaktadır: on mimari tek tohumla (42) "
  "karşılaştırılmıştır. Tohum oynaması, sınamada en yüksek iki MiT-B4 mimarisi (U-Net, UPerNet) ile işletimdeki "
  "modelin mimarisi SegFormer-B4'te, aynı kod ve ayarlarla iki ek tohumla (7 ve 2026) ölçülmüştür. Üç tohumun sınama "
  f"IoU_hasar ortalaması ± standart sapması şöyledir: {ozet}. {iddia} Diğer yedi mimaride çoklu tohum "
  "ölçülmemiştir.")
R("† Tablo 12'deki değerler tek tohumla (42) elde edilmiştir; SegFormer-B4 için ayrıca ölçülen tohum-oynaması "
  "(tohum 7) sınama IoU_hasar'da ±0,02 mertebesindedir (Bölüm 4.2, Sınırlılık); bu bant diğer dokuz mimari için "
  "ayrıca ölçülmemiştir.",
  "† Tablo 12'deki değerler tek tohumla (42) elde edilmiştir. Üç tohumla (42, 7, 2026) yinelenen üç mimaride sınama "
  f"IoU_hasar ortalaması ± standart sapma: {ozet}; diğer yedi mimaride çoklu tohum ölçülmemiştir (Bölüm 4.2, "
  "Sınırlılık).")
R("Koşum tek tohumludur (42); SegFormer-B4 için ayrıca ölçülen tohum-oynaması ±0,02 mertebesindedir (Bölüm 4.2, "
  "Sınırlılık).",
  f"Koşum tek tohumludur (42); üç mimaride üç tohumla ölçülen standart sapma en çok {v(SSMAX)} kadardır (Tablo 12 "
  "dipnotu).")
R("(o) Bölüm 4.2'deki on mimari karşılaştırması tek tohumla (42) yürütülmüştür; SegFormer-B4 için ikinci bir tohumla "
  "(7) ölçülen oynama ±0,02 mertebesindedir, ancak bu bant diğer dokuz mimari için ayrıca ölçülmemiştir. MiT-B4 "
  "ailesi içindeki 0,397-0,433 aralığındaki ince farkların (0,03 mertebesinde) varyans içinde kalıp kalmadığı bu "
  "nedenle kesin değildir; bkz. Tablo 12 dipnotu.",
  "(o) Bölüm 4.2'deki on mimari karşılaştırması tek tohumla (42) yürütülmüştür; tohum oynaması yalnız üç mimaride "
  f"(U-Net, UPerNet, SegFormer-B4) üç tohumla ölçülmüş, standart sapma en çok {v(SSMAX)} bulunmuştur. {iddia_o} "
  "Diğer yedi mimaride çoklu tohum ölçülmemiştir; bkz. Tablo 12 dipnotu. Tatbikatta kullanılan üç model de tek "
  "eğitim tohumundan (42) gelmektedir; eğitim tohumunun tatbikat kazancına etkisi ölçülmemiştir (tatbikattaki beş "
  "tohum senaryo rastgeleliğine aittir).")
R("tek açık nokta çoklu tohum varyansının tüm mimarilerde ölçülmemiş olmasıdır",
  "tek açık nokta çoklu tohum varyansının yalnız üç mimaride ölçülmüş olmasıdır")
R("bölütlemede yol yerine bina hasarı görevi ve tek tohumlu koşum",
  "bölütlemede yol yerine bina hasarı görevi ve çoğu mimaride tek tohumlu koşum")
t23 = next(tb for tb in doc.tables if tb.rows[0].cells[0].text == "Değerlendirme ekseni")
for row in t23.rows:
    if row.cells[0].text.startswith("Çok mimarili kıyaslama — bina"):
        eski = row.cells[3].text
        assert eski.endswith("çoklu tohum tekrarıyla yenilenecektir"), eski
        D.hucre_yaz(row.cells[3], eski.replace("çoklu tohum tekrarıyla yenilenecektir",
                                               "üç mimaride üç tohumla yinelenmiştir (Tablo 12 dipnotu)"))
# OZ / Abstract: tek tohum degeri korunur, cok tohum sonucu eklenir
if en == "unet_mitb4" and (ayrik or her_tohumda):
    R("değeri U-Net ile elde edilmiştir (0,433).",
      f"değeri U-Net ile elde edilmiştir (0,433; üç tohum ortalaması {v(ORT[en])} ± {v(SS[en])}).")
    R("was obtained by U-Net (0.433).",
      f"was obtained by U-Net (0.433; three-seed mean {ORT[en]:.3f} ± {SS[en]:.3f}).")
elif ayrik:
    R("değeri U-Net ile elde edilmiştir (0,433).",
      f"değeri tek tohumda U-Net ile elde edilmiştir (0,433); üç tohumun ortalamasında ise {AD[en]} öndedir "
      f"({v(ORT[en])} ± {v(SS[en])}).")
    R("was obtained by U-Net (0.433).",
      f"was obtained by U-Net with a single seed (0.433); averaged over three seeds, {AD[en]} leads "
      f"({ORT[en]:.3f} ± {SS[en]:.3f}).")
else:
    R("değeri U-Net ile elde edilmiştir (0,433).",
      "değeri tek tohumda U-Net ile elde edilmiştir (0,433); üç tohumla yinelendiğinde MiT-B4 mimarileri arasındaki "
      "fark tohum oynaması mertebesinde kalmıştır.")
    R("was obtained by U-Net (0.433).",
      "was obtained by U-Net with a single seed (0.433); with three seeds, differences among MiT-B4 architectures "
      "stayed within seed variation.")
R("Kodlayıcısı MiT-B4'te sabitlenmiş beş mimari arasında en yüksek sınama IoU_hasar değerini U-Net vermiştir (0,433);",
  "Tek tohumda (42), kodlayıcısı MiT-B4'te sabitlenmiş beş mimari arasında en yüksek sınama IoU_hasar değerini "
  "U-Net vermiştir (0,433);")
# Tatbikattaki U-Net etiketi: "En iyi" tek tohuma dayaniyordu; olculen ozellik "en buyuk" (64,1 M parametre)
for eski, yeni in (("İslahiye'de dengeli model %2,06 hasarla %33,7 kazandırırken en iyi model",
                    "İslahiye'de dengeli model %2,06 hasarla %33,7 kazandırırken en büyük model"),
                   ("(Nurdağı, en iyi model:", "(Nurdağı, en büyük model:"),
                   ("İslahiye'de en iyi modelle 1,8 keşifle", "İslahiye'de en büyük modelle 1,8 keşifle"),
                   ("İslahiye'de en iyi modelin %1,34'üne", "İslahiye'de en büyük modelin %1,34'üne"),
                   ("Tablo 16'daki en iyi modelle (U-Net MiT-B4)", "Tablo 16'daki en büyük modelle (U-Net MiT-B4)")):
    R(eski, yeni)
n_hucre = 0
for tb in doc.tables:
    for row in tb.rows:
        for c in row.cells:
            if c.text in ("En iyi (U-Net MiT-B4)", "En iyi"):
                D.hucre_yaz(c, c.text.replace("En iyi", "Büyük"))
                n_hucre += 1
assert n_hucre == 4, n_hucre   # Tablo 16'da uc alan, Tablo 17'de bir satir
D.log.append(f"ETIKET En iyi -> Büyük ({n_hucre} hücre)")
# Tumuyle kalin kalmis tablo/sekil basliklari (eski Tablo 19 ve 21'den): etiket kalin, metin normal
import copy as _copy
import re as _re
n_bas = 0
for p in D.paragraflar():
    m = _re.match(r"(Tablo|Şekil) \d+\.", p.text)
    if not m or not p.runs:
        continue
    r0 = p.runs[0]
    if r0.bold and r0.text.startswith(m.group(0)) and len(r0.text) > len(m.group(0)):
        kalan = r0.text[len(m.group(0)):]
        r0.text = m.group(0)
        yeni_r = _copy.deepcopy(r0._element)
        r0._element.addnext(yeni_r)
        from docx.text.run import Run as _Run
        rr = _Run(yeni_r, p)
        rr.text = kalan
        rr.bold = False
        n_bas += 1
assert n_bas == 2, n_bas
D.log.append(f"BASLIK bicimi duzeltildi ({n_bas})")
D.log.append("COKLU TOHUM: " + ozet)
doc.save(DST)
print("\n".join(D.log))
