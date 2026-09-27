# -*- coding: utf-8 -*-
"""Revizyon 4 (27 Eylul 2026): Word okuma turunda bulunan duzeltmeler ve
tablo numaralari kaydirildiktan sonra bozulan Turkce ek uyumunun onarimi.
Kullanim: python revize4.py [--deneme]"""
import re
import shutil
import sys

sys.path.insert(0, r"C:\Users\utkuc\AppData\Local\Temp\claude\D--quakemind-final-QuakeMind\4779ac37-74b5-4aeb-8690-74d7792eedc5\scratchpad")
from docx import Document
from yardimci import Duzenleyici

SRC = r"D:\quakemind_final\QuakeMind\makale\arsiv\QuakeMind_Makale_YEDEK_pre_revize4_2026-09-27.docx"
DST = r"D:\quakemind_final\QuakeMind\makale\QuakeMind_Makale_Taslak_2026-09-16_sekillieklendi_v3_LOCAL_2026-09-21.docx"
if "--deneme" in sys.argv:
    DST = r"C:\Users\utkuc\AppData\Local\Temp\claude\D--quakemind-final-QuakeMind\4779ac37-74b5-4aeb-8690-74d7792eedc5\scratchpad\deneme4.docx"

shutil.copyfile(SRC, DST)
doc = Document(DST)
D = Duzenleyici(doc)
deg, yaz, sil = D.degistir, D.yaz, D.sil

# ------------------------------------------------------------------ tekrar / kalinti
deg(" Kapanma katmanı, kendi hasar modelimizin gerçek drone ortofotosu üzerindeki çıktısının OSM yol ağıyla piksel "
    "kesişiminden türetilmiştir.", "")
deg("; bu türetme uygulamanın üretim yolunun kullandığı kodun aynısıdır.",
    "; model çıkarımı uygulamanın üretim yolunun kullandığı kodun aynısıdır.")
deg("çıktısı yol ağı rasteriyle kesiştirilerek", "çıktısı yol ağı kenarlarıyla kesiştirilerek")
for tirnak in ('"', "“"):
    kapat = '"' if tirnak == '"' else "”"
    kalip = ' (' + tirnak + 'QuakeMindBackend/apps/road_damage/utils/inference.py' + kapat + ' ve ' + tirnak + 'network.py' + kapat + ')'
    if D.say(kalip):
        deg(kalip, "")
        break
deg("piksel düzeyi çıktısının OSM yol ağının rasterize edilmiş hâliyle kesiştirilmesiyle (Bölüm 3.4)",
    "piksel düzeyi çıktısının OSM yol kenarlarının merkez çizgisi boyunca örneklenmesiyle (Bölüm 3.4)")
deg("fotoğraflar aynı depremin gerçek kareleridir ama", "fotoğraflar aynı depremlerin ardından çekilmiş gerçek karelerdir ama")

# ------------------------------------------------------------------ "en cok 25 cift"
deg("mahalle başına 25 ekip–afetzede çifti (beş giriş ağzı × beş afetzede, beş turda döndürülerek)",
    "mahalle başına en çok 25 ekip–afetzede çifti (beş giriş ağzı × en çok beş afetzede, beş turda döndürülerek)")
deg("with 25 affected-individual/team pairs per neighborhood (five entry mouths × five individuals, rotated over five turns)",
    "with up to 25 affected-individual/team pairs per neighborhood (five entry mouths × up to five individuals, rotated "
    "over five turns)")
deg("mahallede (Defne/Elektrik, İslahiye, Nurdağı), mahalle başına 25 ekip–afetzede çifti, beş rastgelelik",
    "mahallede (Defne/Elektrik, İslahiye, Nurdağı), mahalle başına en çok 25 ekip–afetzede çifti, beş rastgelelik")
deg("(mahalle başına beş giriş ağzı × beş afetzede, beş turda 25 çift)",
    "(mahalle başına beş giriş ağzı × en çok beş afetzede, beş turda en çok 25 çift)")
deg("beş turda döndürülerek eşlenir (25 çift).", "beş turda döndürülerek eşlenir (en çok 25 çift).")
deg("(mahalle başına 25 çift)", "(mahalle başına en çok 25 çift)", hepsi=True)

# ------------------------------------------------------------------ olculmus sure
deg("Bölütleme çıkarımı, tek bir 512 × 512 pencerenin işlenmesi merkezî işlem biriminde saniyeler mertebesinde "
    "sürebildiği ve geniş yarıçaplı bir analizde yüzlerce pencere işlendiği için",
    "Bölütleme çıkarımı, yaklaşık 500 m'lik bir mahalle görüntüsünde bile merkezî işlem biriminde modele göre 1,9–10,5 "
    "saniye sürdüğü (Tablo 16) ve daha geniş bir analiz alanında pencere sayısı alanla birlikte arttığı için")

# ------------------------------------------------------------------ tablolar
t10 = [t for t in doc.tables if t.rows[0].cells[0].text == "Ölçüt" and len(t.columns) == 3
       and "Birincil" in t.rows[0].cells[2].text][0]
for row in t10.rows:
    if row.cells[0].text.startswith("IoU_hasar"):
        D.hucre_yaz(row.cells[2], "Anlamsal bölütleme (hasar maskesi)")
t22 = [t for t in doc.tables if t.rows[0].cells[0].text.startswith("Değerlendirme ekseni")][0]
for row in t22.rows:
    if row.cells[0].text.startswith("Yol bölütlemesi"):
        D.hucre_yaz(row.cells[0], "Kapanma katmanı (yol)")
    elif row.cells[0].text.startswith("Risk skorlaması"):
        D.hucre_yaz(row.cells[3], row.cells[3].text + "; tek öznitelikli naif taban (R0b) ROC-AUC = 0,6143, "
                    "PR-AUC = 0,4021, MCC = 0,1616 ile CatBoost'u geçmektedir")

# ------------------------------------------------------------------ 4.4 / 4.6 / 4.7
deg("iki il arasında çift yönlü bir genelleme deneyi kurgulanmıştır.",
    "iki il arasında çift yönlü bir genelleme deneyi tasarlanmıştır.")
sil("Karşılaştırma sonuçları Tablo 15'de özetlenmektedir.")
deg("ekip kaç kapanmaya fiilen denk gelirse kazanç o kadar büyüktür.",
    "ekip kaç kapanmaya fiilen denk gelirse kazanç o kadar büyüktür. Bu ilişkinin bir bölümü yapısaldır: kazancın büyük "
    "bölümü keşif başına beyan edilen cezadan geldiğinden (aşağıda, Farkın ayrıştırılması), keşif sayısı ile kazanç "
    "arasındaki bağ ölçülen yol farkından çok bu cezayı yansıtır.")
deg("az sayıda keşif de görev başına yüksek bir kazanç üretebilmektedir.",
    "az sayıda keşif de görev başına yüksek bir kazanç üretebilmektedir. Kazancın büyük bölümü keşif başına beyan edilen "
    "cezadan geldiğinden bu ilişki kısmen yapısaldır (Bölüm 4.6).")

# ------------------------------------------------------------------ 4.11
deg("(ortalama, standart sapma, en büyük değer, olay sayısı)",
    "(ortalama, standart sapma, en büyük değer ve olay sayısı; sonuncusu her kayıtta 7 olduğundan bilgi taşımaz)")

# ------------------------------------------------------------------ 5 ve 6
deg("Bu, görsel tespit başarımı ile operasyonel sonuç arasındaki kopukluğu somut biçimde göstermektedir.",
    "Bu, görsel tespit başarımı ile operasyonel sonuç arasındaki kopukluğu somut biçimde göstermektedir. Kazancın büyük "
    "bölümü keşif başına beyan edilen cezadan geldiğinden, bu gözlemin büyüklüğü cezanın değerinden bağımsız değildir "
    "(Bölüm 4.9).")
deg("kapanmanın rotaya düşmediği alanlarda fark yoktur. Bu büyüklük,",
    "kapanmanın rotaya düşmediği alanlarda fark yoktur. Kaba kestirime göre bu kazancın büyük bölümü keşif başına "
    "beyan edilen cezadan geldiğinden sonuç, cezanın değerine duyarlılığıyla birlikte bant olarak okunmalıdır. Bu "
    "büyüklük,")
deg("nihai değerler değil ön bulgular olarak sunulmaktadır.",
    "nihai değerler değil ön bulgular olarak sunulmaktadır. Risk skorlamasında tek öznitelikli bir naif taban (son yedi "
    "olayın ortalama büyüklüğü) üretimdeki CatBoost modelini ROC-AUC, PR-AUC ve MCC'de geçmiştir; bu, risk bileşeninin "
    "yeniden ele alınması gerektiğini göstermektedir.")
deg("yol ağı verisini sağlayan OpenStreetMap katkıcılarına ve görev tabanlı değerlendirmede",
    "yol ağı verisini sağlayan OpenStreetMap katkıcılarına, görev tabanlı değerlendirmede")


# ------------------------------------------------------------------ Turkce ek uyumu
BIRLER = ["", "bir", "iki", "üç", "dört", "beş", "altı", "yedi", "sekiz", "dokuz"]
ONLAR = ["", "on", "yirmi", "otuz", "kırk", "elli", "altmış", "yetmiş", "seksen", "doksan"]
UNLU = "aeıioöuü"
SERT = "fstkçşhp"


def okunus(n):
    return BIRLER[n % 10] if n % 10 else ONLAR[(n // 10) % 10]


def ek(n, eski):
    w = okunus(n)
    son_unlu = [c for c in w if c in UNLU][-1]
    on = son_unlu in "eiöü"
    yuvarlak = son_unlu in "oöuü"
    unluyle = w[-1] in UNLU
    sert = w[-1] in SERT
    A = "e" if on else "a"
    I = ("ü" if yuvarlak else "i") if on else ("u" if yuvarlak else "ı")
    Dh = "t" if sert else "d"
    m = re.fullmatch(r"[dt][ea](ki|n)?", eski)
    if m:
        return Dh + A + (m.group(1) or "")
    if re.fullmatch(r"y?[ea]", eski):
        return ("y" if unluyle else "") + A
    if re.fullmatch(r"n?[iıuü]n", eski):
        return ("n" if unluyle else "") + I + "n"
    if re.fullmatch(r"y?[iıuü]", eski):
        return ("y" if unluyle else "") + I
    return None


def konumsal(p, bas, son, yeni):
    runs = list(p.runs)
    pos, ilk = 0, None
    for k, r in enumerate(runs):
        L = len(r.text)
        if ilk is None and pos + L > bas:
            ilk = (k, bas - pos)
        if pos + L >= son:
            kb, ob = ilk
            if kb == k:
                r.text = r.text[:ob] + yeni + r.text[son - pos:]
            else:
                runs[kb].text = runs[kb].text[:ob] + yeni + runs[k].text[son - pos:]
                for j in range(kb + 1, k + 1):
                    runs[j].text = ""
            return
        pos += L


desen = re.compile(r"(?:Tablo|Şekil)\s+(?:\d+\s*[–\-]\s*)?(\d+)(['’])([a-zçğıöşü]+)")
duzeltilen, taninmayan = 0, set()
for p in list(D.paragraflar()):
    for m in reversed(list(desen.finditer(p.text))):
        yeni_ek = ek(int(m.group(1)), m.group(3))
        if yeni_ek is None:
            taninmayan.add(m.group(0))
            continue
        if yeni_ek != m.group(3):
            konumsal(p, m.start(3), m.end(3), yeni_ek)
            duzeltilen += 1
D.log.append("EK UYUMU duzeltildi: %d; taninmayan: %s" % (duzeltilen, sorted(taninmayan)))

D.kalinlik_duzelt()
doc.save(DST)
print("\n".join(D.log))
print("TOPLAM", len(D.log), "->", DST)
