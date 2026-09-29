# -*- coding: utf-8 -*-
"""Revizyon 17: §4.4 çoklu tohum — il dışarıda bırakma deneyi üç tohumla (42, 7, 2026).

Kaynak: makale/analiz_arsivi/kaggle/sonuclar/
  birleşik:  birlesik_P2/unet_mitb4__tohum_42.json, birlesik_P2/segformer_mitb4__tohum_42__2.json (Tablo 12),
             oturum1_coklu_tohum/P2__<model>__tohum_{7,2026}.json, oturum4_segformer_tohum7/P2__segformer_mitb4__tohum_7.json
  hariç:     oturum2/3 (tohum 42), oturum5/6 (tohum 7, 2026): <P2_haric_hatay|P2_haric_maras>__<model>__tohum_<n>.json
Birleşik ve il dışarıda bırakma koşuları tohum başına eşleştirilir (aynı tohum: aynı başlangıç ve artırım sırası).
Güncel makaleden (revize16 çıktısı) başlar; Tablo 14'ün hücrelerini ve §4.4 cümlelerini yeniden yazar.
Yedek: makale/arsiv/QuakeMind_Makale_YEDEK_pre_revize17_2026-09-29.docx (yoksa güncel dosyadan alınır).
"""
import glob
import json
import os
import shutil
import statistics as st
import sys

BURASI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BURASI)
from docx import Document
from yardimci import Duzenleyici

KOK = os.path.abspath(os.path.join(BURASI, "..", ".."))
SON = os.path.join(BURASI, "kaggle", "sonuclar")
DST = os.path.join(KOK, "makale", "QuakeMind_Makale_Taslak_2026-09-16_sekillieklendi_v3_LOCAL_2026-09-21.docx")
SRC = os.path.join(KOK, "makale", "arsiv", "QuakeMind_Makale_YEDEK_pre_revize17_2026-09-29.docx")
v = lambda x, b=3: f"{x:.{b}f}".replace(".", ",")
r3 = lambda x: float(f"{x:.3f}")
fark = lambda x: ("−" if x < 0 else "+") + v(abs(x))
oku = lambda y: json.load(open(y, encoding="utf-8"))
TOH = [42, 7, 2026]
AL = {"Hatay": "sinama_hatay", "Kahramanmaraş": "sinama_kahramanmaras"}
AD = {"unet_mitb4": "U-Net (MiT-B4)", "segformer_mitb4": "SegFormer-B4"}
DENEY = {"Hatay": "P2_haric_hatay", "Kahramanmaraş": "P2_haric_maras"}
BOLGE = {"Hatay": "haric:Hatay", "Kahramanmaraş": "haric:Kahramanmaraş"}
KARO = {"Hatay": 521, "Kahramanmaraş": 798}
SAYI = {1: "bir", 2: "iki", 3: "üç", 4: "dört", 5: "beş", 6: "altı", 7: "yedi", 8: "sekiz", 9: "dokuz", 10: "on",
        11: "on bir", 12: "on iki"}
SAYIDA = {8: "sekizinde", 9: "dokuzunda", 10: "onunda", 11: "on birinde", 12: "on ikisinde"}


def tek(desen):
    f = glob.glob(desen)
    assert len(f) == 1, (desen, f)
    return oku(f[0])


def birlesik(m, t):
    if t == 42:
        return oku(os.path.join(SON, "birlesik_P2", f"{m}__tohum_42{'__2' if m == 'segformer_mitb4' else ''}.json"))
    return tek(os.path.join(SON, "oturum*", f"P2__{m}__tohum_{t}.json"))


def haric(m, il, t):
    return tek(os.path.join(SON, "oturum*", f"{DENEY[il]}__{m}__tohum_{t}.json"))


def denetle(k, m, t, bolge, karo):
    a = k["ayarlar"]
    assert k["model"] == m and k["tohum"] == t, (m, t)
    assert a["artirim_tohumlu"] is True and a["girdi"] == "sonrasi" and a["donem"] == 50, (m, t)
    assert a["bolge_egitimi"] == bolge, (m, t, a["bolge_egitimi"])
    assert k["egitim"]["karo"]["egitim"] == karo, (m, t, k["egitim"]["karo"]["egitim"])
    assert k["sinama_hatay"]["n_karo"] == 34 and k["sinama_kahramanmaras"]["n_karo"] == 32
    assert k["sinama_sonrasi_tumu"]["n_karo"] == 91


B = {m: [birlesik(m, t) for t in TOH] for m in AD}
H = {(m, i): [haric(m, i, t) for t in TOH] for m in AD for i in AL}
for m in AD:
    for k, t in zip(B[m], TOH):
        denetle(k, m, t, None, 1013)
    for i in AL:
        for k, t in zip(H[(m, i)], TOH):
            denetle(k, m, t, BOLGE[i], KARO[i])
iou = lambda ks, a: [k[a]["iou_hasar"] for k in ks]
ORT, SS = st.mean, st.stdev

SONUC = {}
for (m, i), hs in H.items():
    j = next(x for x in AL if x != i)
    SONUC[(m, i)] = dict(
        j=j,
        kayip=r3(ORT(iou(B[m], AL[i]))) - r3(ORT(iou(hs, AL[i]))),     # tabloda gorunen ortalamalardan
        oran=(ORT(iou(B[m], AL[i])) - ORT(iou(hs, AL[i]))) / ORT(iou(B[m], AL[i])),
        d_gor=[h - b for h, b in zip(iou(hs, AL[i]), iou(B[m], AL[i]))],
        ote=r3(ORT(iou(hs, AL[j]))) - r3(ORT(iou(B[m], AL[j]))))
    SONUC[(m, i)]["dusen"] = sum(x < 0 for x in SONUC[(m, i)]["d_gor"])
for k, s in SONUC.items():
    print(k, "kayip", round(s["kayip"], 4), "oran", round(s["oran"], 3), "dusen", s["dusen"],
          "d_gor", [round(x, 4) for x in s["d_gor"]], "ote", round(s["ote"], 4))
KMIN = min(s["kayip"] for s in SONUC.values())
KMAX = max(s["kayip"] for s in SONUC.values())
OMIN = min(s["oran"] for s in SONUC.values())
OMAX = max(s["oran"] for s in SONUC.values())
TOPLAM = 3 * len(SONUC)
DUSEN = sum(s["dusen"] for s in SONUC.values())
kesin = DUSEN >= 0.9 * TOPLAM
YARIM = ("unet_mitb4", "Hatay")          # egitim kumesinin yariya indigi ve oteki ilin de dustugu durum
OTEKI_DIGER = [abs(s["ote"]) for k, s in SONUC.items() if k != YARIM]
assert SONUC[YARIM]["ote"] < -0.02 and max(OTEKI_DIGER) < 0.02, (SONUC[YARIM]["ote"], OTEKI_DIGER)
print("KMIN", KMIN, "KMAX", KMAX, "dusen", DUSEN, "/", TOPLAM, "kesin", kesin)

if not os.path.exists(SRC):
    shutil.copyfile(DST, SRC)
shutil.copyfile(SRC, DST)
doc = Document(DST)
D = Duzenleyici(doc)
R = D.degistir

# ---------------------------------------------------------------- Tablo 14
t14 = next(t for t in doc.tables if t.rows[0].cells[0].text == "Mimari" and
           t.rows[0].cells[1].text == "Eğitimden çıkarılan il")
satir = {}
for m in AD:
    b_tum = iou(B[m], "sinama_sonrasi_tumu")
    satir[(m, None)] = [f"{v(ORT(iou(B[m], AL[j])))} ± {v(SS(iou(B[m], AL[j])))}" for j in AL] + \
                       [f"{v(ORT(b_tum))} ± {v(SS(b_tum))}"]
    for i in AL:
        hs = H[(m, i)]
        hucre = []
        for j in AL:
            o = ORT(iou(hs, AL[j]))
            hucre.append(f"{v(o)} ± {v(SS(iou(hs, AL[j])))} ({fark(r3(o) - r3(ORT(iou(B[m], AL[j]))))})")
        h_tum = iou(hs, "sinama_sonrasi_tumu")
        satir[(m, i)] = hucre + [f"{v(ORT(h_tum))} ± {v(SS(h_tum))}"]
n = 0
for row in t14.rows[1:]:
    m = next(k for k, a in AD.items() if row.cells[0].text == a)
    il = None if row.cells[1].text.startswith("—") else row.cells[1].text
    for c, deger in zip(row.cells[3:6], satir[(m, il)]):
        D.hucre_yaz(c, deger)
    n += 1
assert n == 6, n
D.log.append("TABLO 14 uc tohum")
R("(sınama IoU_hasar, tohum 42).",
  "(sınama IoU_hasar; üç tohum, 42, 7 ve 2026; ortalama ± standart sapma).")

# ---------------------------------------------------------------- yontem cumlesi
R("Deney, sınamada en yüksek IoU_hasar değerini veren U-Net (MiT-B4) ile",
  "Deney, tek tohumda sınamada en yüksek IoU_hasar değerini veren U-Net (MiT-B4) ile")
R("Tablo 12 ile aynı ayarlar, yazılım sürümleri ve tohumla (42) koşulmuştur; sonuçlar Tablo 14'tedir.",
  "Tablo 12 ile aynı ayarlar ve yazılım sürümleriyle, üç tohumla (42, 7, 2026) koşulmuştur. Karşılaştırma tabanı "
  "olan birleşik eğitim de aynı üç tohumla yinelenmiştir (Bölüm 4.2); birleşik ve il dışarıda bırakma koşuları tohum "
  "başına eşleştirilmiştir. Sonuçlar Tablo 14'tedir.")

# ---------------------------------------------------------------- bulgu paragrafi
u_m, s_m = SONUC[("unet_mitb4", "Kahramanmaraş")], SONUC[("segformer_mitb4", "Kahramanmaraş")]
tek_tohum = "0,05–0,10"   # revize15'in tohum 42 araligi (Tablo 14'un onceki surumu)
cumleler = [
    f"Üç tohumun ortalamasında, eğitimde görülmemiş ilde IoU_hasar birleşik eğitime göre {v(KMIN, 2)}–{v(KMAX, 2)} "
    f"düşmüştür (%{OMIN * 100:.0f}–{OMAX * 100:.0f}). Birleşik ve il dışarıda bırakma koşuları aynı tohumla "
    f"eşleştirildiğinde görülmemiş ildeki düşüş {SAYI[TOPLAM]} eşleşmenin {SAYIDA[DUSEN]} görülmektedir.",
    f"Aynı modelin öteki ildeki ortalama değişimi ±{v(max(OTEKI_DIGER))} içinde kalmaktadır; tek istisna, eğitim "
    f"kümesinin yarıya indiği (1 013 → 521 karo) U-Net'te Hatay'ın çıkarıldığı durumdur (öteki ilde "
    f"{fark(SONUC[YARIM]['ote'])}), burada kaybın bir kısmı veri azalmasına bağlanabilir.",
    f"Kahramanmaraş çıkarıldığında eğitim kümesi yalnız %21 küçüldüğünden (798 karo) bu durumdaki kayıplar (U-Net "
    f"{fark(-u_m['kayip'])}, SegFormer-B4 {fark(-s_m['kayip'])}) büyük ölçüde görülmemiş ilin etkisine bağlanabilir.",
    f"Tohum 42'deki tek koşunun verdiği aralık ({tek_tohum}) ile üç tohumun ortalamaları arasındaki fark, tek tohumlu "
    f"bir genelleme ölçümünün ne kadar oynak olduğunu göstermektedir."]
D.yaz("Eğitimde görülmemiş ilde IoU_hasar birleşik eğitime göre", " ".join(cumleler))
R("Sınırlar: deney tek tohumludur ve il başına 32–34 sınama karosuna dayanır;",
  "Sınırlar: deney üç tohumludur ve il başına 32–34 sınama karosuna dayanır;")

# ---------------------------------------------------------------- oteki yerler
R("genelleme deneyinde görülmemiş ilde IoU_hasar 0,05–0,10 düşmüştür (tek tohum; tohum oynaması mertebesinde).",
  f"genelleme deneyinde görülmemiş ilde IoU_hasar üç tohumun ortalamasında {v(KMIN, 2)}–{v(KMAX, 2)} düşmüştür"
  + ("." if kesin else " (kısmen tohum oynaması mertebesinde)."))
R("dropped by 0.05–0.10 (single seed; within seed variation).",
  f"dropped by {KMIN:.2f}–{KMAX:.2f} on average over three seeds"
  + ("." if kesin else " (partly within seed variation)."))
R("Bölüm 4.4'teki genelleme deneyi iki mimaride ve tek tohumla koşulmuştur.",
  "Bölüm 4.4'teki genelleme deneyi iki mimaride ve üç tohumla koşulmuştur.")
R("bölgeler arası genelleme iki mimaride tek tohumla ölçülmüştür.",
  "bölgeler arası genelleme iki mimaride üç tohumla ölçülmüştür.")
R("bölgeler arası genelleme ise iki mimaride ve tek tohumla ölçülebilmiştir;",
  "bölgeler arası genelleme ise iki mimaride ve üç tohumla ölçülmüştür;")
t23 = next(tb for tb in doc.tables if tb.rows[0].cells[0].text == "Değerlendirme ekseni")
for row in t23.rows:
    if row.cells[0].text.startswith("Bölgeler arası genelleme"):
        D.hucre_yaz(row.cells[3], f"Ölçülmüştür (Tablo 14): görülmemiş ilde IoU_hasar üç tohumun ortalamasında "
                                  f"{v(KMIN, 2)}–{v(KMAX, 2)} düşük (U-Net ve SegFormer-B4)")
D.log.append("TABLO 23 genelleme satiri")
metin = "\n".join(p.text for p in D.paragraflar())
for eski in ("tek tohum; tohum oynaması", "single seed; within seed", "deney tek tohumludur",
             "iki mimaride tek tohumla", "iki mimaride ve tek tohumla", "(sınama IoU_hasar, tohum 42)"):
    assert eski not in metin, eski
doc.save(DST)
print("\n".join(D.log))
