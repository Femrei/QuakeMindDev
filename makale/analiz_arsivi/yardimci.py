# -*- coding: utf-8 -*-
"""revize2.py icin docx yardimcilari: bicimi koruyan metin degisimi, paragraf
ekleme/silme/tasima, gorsel degistirme, tablo bicimi ve kalinlik duzeltmesi."""
import copy
import re
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Cm
from docx.text.paragraph import Paragraph

M = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'


class Duzenleyici:
    def __init__(self, doc):
        self.doc = doc
        self.log = []

    # ---------------------------------------------------------------- bulma
    def paragraflar(self, tablolar_dahil=True):
        body = self.doc.element.body
        for p in body.iter(qn('w:p')):
            yield Paragraph(p, self.doc)

    def bul(self, sub, nth=0, tablolar_dahil=True):
        hits = [p for p in self.paragraflar() if sub in p.text]
        if len(hits) <= nth:
            raise KeyError("BULUNAMADI: " + sub[:70])
        return hits[nth]

    def say(self, sub):
        return sum(1 for p in self.paragraflar() if sub in p.text)

    # ------------------------------------------------------ metin degisimi
    def degistir(self, eski, yeni, nth=0, hepsi=False):
        """Paragraf icinde eski->yeni; run bicimini korur (degisim baslangic
        run'inin bicimini alir)."""
        n = 0
        hedefler = [p for p in self.paragraflar() if eski in p.text]
        if not hedefler:
            raise KeyError("DEGISTIR BULUNAMADI: " + eski[:70])
        if not hepsi:
            hedefler = [hedefler[nth]]
        for p in hedefler:
            while eski in p.text:
                self._run_degistir(p, eski, yeni)
                n += 1
                if not hepsi:
                    break
        self.log.append("DEG x%d  %s" % (n, eski[:50]))
        return n

    @staticmethod
    def _run_degistir(p, eski, yeni):
        runs = [r for r in p.runs]
        metin = ''.join(r.text for r in runs)
        i = metin.index(eski)
        j = i + len(eski)
        pos = 0
        bas = son = None
        for k, r in enumerate(runs):
            L = len(r.text)
            if bas is None and pos + L > i:
                bas = (k, i - pos)
            if pos + L >= j:
                son = (k, j - pos)
                break
            pos += L
        kb, ob = bas
        ks, os_ = son
        if kb == ks:
            t = runs[kb].text
            runs[kb].text = t[:ob] + yeni + t[os_:]
        else:
            runs[kb].text = runs[kb].text[:ob] + yeni + runs[ks].text[os_:]
            for k in range(kb + 1, ks + 1):
                runs[k].text = ''

    def yaz(self, sub, yeni, nth=0):
        """Paragrafin butun metnini degistirir (ilk run bicimi kalir)."""
        p = self.bul(sub, nth)
        rs = p.runs
        if rs:
            rs[0].text = yeni
            for r in rs[1:]:
                r._element.getparent().remove(r._element)
        else:
            p.add_run(yeni)
        self.log.append("YAZ  " + sub[:50])
        return p

    # --------------------------------------------------- paragraf islemleri
    def sil(self, sub, nth=0):
        p = self.bul(sub, nth)
        p._p.getparent().remove(p._p)
        self.log.append("SIL  " + sub[:50])

    def sonrasina_ekle(self, capa_el, metin, ornek_p, kalin_bas=None):
        """capa_el ogesinin arkasina ornek_p bicimli yeni paragraf ekler."""
        yeni = OxmlElement('w:p')
        if ornek_p._p.pPr is not None:
            yeni.append(copy.deepcopy(ornek_p._p.pPr))
        capa_el.addnext(yeni)
        p = Paragraph(yeni, self.doc)
        rpr = None
        for r in ornek_p.runs:
            if r._element.rPr is not None:
                rpr = r._element.rPr
                break
        def _run(t, kalin):
            r = p.add_run(t)
            if rpr is not None:
                r._element.insert(0, copy.deepcopy(rpr))
            r.bold = kalin
            return r
        if kalin_bas and metin.startswith(kalin_bas):
            _run(kalin_bas, True)
            _run(metin[len(kalin_bas):], False)
        else:
            _run(metin, False)
        self.log.append("EKLE " + metin[:50])
        return p

    def tasi(self, sub, capa_sub):
        p = self.bul(sub)
        capa = self.bul(capa_sub)
        capa._p.addnext(p._p)
        self.log.append("TASI %s -> %s" % (sub[:30], capa_sub[:30]))

    # ------------------------------------------------------------ gorseller
    @staticmethod
    def gorsel_paragrafi_mi(p):
        return bool(p._p.findall('.//' + qn('w:drawing')))

    def resim_koy(self, p, yol, genislik_cm):
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        for d in p._p.findall('.//' + qn('w:drawing')):
            d.getparent().getparent().remove(d.getparent())
        p.add_run().add_picture(yol, width=Cm(genislik_cm))
        self.log.append("RESIM %s" % yol.split('\\')[-1].split('/')[-1])

    def basliktan_onceki_gorsel(self, baslik_sub):
        cap = self.bul(baslik_sub)._p
        el = cap.getprevious()
        while el is not None:
            if el.tag == qn('w:p') and el.findall('.//' + qn('w:drawing')):
                return Paragraph(el, self.doc)
            if el.tag == qn('w:tbl'):
                return el  # tablo kapsayici
            el = el.getprevious()
        raise KeyError("gorsel yok: " + baslik_sub)

    # --------------------------------------------------------------- tablo
    @staticmethod
    def hucre_yaz(hucre, metin):
        ps = hucre.paragraphs
        p = ps[0]
        rs = p.runs
        if rs:
            rs[0].text = metin
            for r in rs[1:]:
                r._element.getparent().remove(r._element)
        else:
            p.add_run(metin)
        for q in ps[1:]:
            q._p.getparent().remove(q._p)

    def satir_ekle(self, tablo, sonra_idx, degerler):
        kaynak = tablo.rows[sonra_idx]._tr
        yeni = copy.deepcopy(kaynak)
        kaynak.addnext(yeni)
        from docx.table import _Row
        satir = _Row(yeni, tablo)
        for c, v in zip(satir.cells, degerler):
            self.hucre_yaz(c, v)
        self.log.append("SATIR+ " + degerler[0][:30])
        return satir

    @staticmethod
    def tablo_bicimle(tablo, punto=8.5):
        """Ozgun tablolarla ayni bicim: %100 genislik, ortali, baslik satiri
        gri (E8E8E8) ve kalin, 1 pt paragraf sonrasi bosluk."""
        tbl = tablo._tbl
        pr = tbl.tblPr
        w = pr.find(qn('w:tblW'))
        if w is None:
            w = OxmlElement('w:tblW'); pr.append(w)
        w.set(qn('w:w'), '5000'); w.set(qn('w:type'), 'pct')
        jc = pr.find(qn('w:jc'))
        if jc is None:
            jc = OxmlElement('w:jc'); pr.append(jc)
        jc.set(qn('w:val'), 'center')
        for ri, row in enumerate(tablo.rows):
            for c in row.cells:
                if ri == 0:
                    tcPr = c._tc.get_or_add_tcPr()
                    shd = tcPr.find(qn('w:shd'))
                    if shd is None:
                        shd = OxmlElement('w:shd'); tcPr.append(shd)
                    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), 'E8E8E8')
                for p in c.paragraphs:
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.space_before = Pt(0)
                    for r in p.runs:
                        r.font.size = Pt(punto)
                        if ri == 0:
                            r.bold = True

    # ----------------------------------------------------- kalinlik duzelt
    def kalinlik_duzelt(self, ozel=None):
        """Tamamen kalin kalmis govde paragraflarini duzeltir: yalniz etiket
        ('Tablo N.' / 'Sekil N.') ya da bas cumle kalin kalir."""
        ozel = ozel or {}
        n = 0
        for i, p in enumerate(self.doc.paragraphs):
            if p.style.name.startswith('Heading') or len(p.text) < 60:
                continue
            rs = [r for r in p.runs if r.text.strip()]
            if not rs or not all(r.bold for r in rs):
                continue
            t = p.text
            if i in (0, 1, 12):  # baslik ve Ingilizce baslik
                continue
            kural = None
            for anahtar, v in ozel.items():
                if t.startswith(anahtar):
                    kural = v
            if kural is None and t.startswith('Algoritma '):
                kural = -1  # algoritma basliklari ozgun metinde tamamen kalin
            if kural is None and p.style.name == 'List Number':
                m = re.match(r'^(.{3,260}?[.?])\s', t)
                kural = len(m.group(1)) if m else 0
            if kural == -1:
                continue
            if kural is None:
                m = re.match(r'^((?:Tablo|Şekil) \d+\.)', t)
                if m:
                    kural = len(m.group(1))
                elif re.match(r'^\([a-z]\) ', t):
                    kural = 0
                else:
                    m = re.match(r'^(.{3,90}?[.?])\s', t)
                    kural = len(m.group(1)) if m else 0
            # tek run'a indir, sonra bol
            rs0 = p.runs[0]
            for r in p.runs[1:]:
                r._element.getparent().remove(r._element)
            rs0.text = t
            if kural == 0:
                rs0.bold = False
            else:
                bas, geri = t[:kural], t[kural:]
                rs0.text = bas
                rs0.bold = True
                yeni = copy.deepcopy(rs0._element)
                rs0._element.addnext(yeni)
                from docx.text.run import Run
                r2 = Run(yeni, p)
                r2.text = geri
                r2.bold = False
            n += 1
        self.log.append("KALINLIK duzeltildi: %d paragraf" % n)


def alanlari_duzlestir(doc):
    """REF/SEQ gibi karmasik alanlari onbellekteki sonuc metnine indirger.
    Onceki revizyon bazi basliklarin yer imlerini bosaltmisti; Word alanlari
    guncelleyince 'Sekil 'te' gibi numarasiz metin cikiyordu. Sonuc metni
    (python-docx'in gordugu, dogrulanmis numara) yerinde kalir."""
    body = doc.element.body
    n = 0
    for p in list(body.iter(qn('w:p'))):
        runs = [r for r in p.iter(qn('w:r'))]
        durum = []  # yigin: 'instr' ya da 'sonuc'
        sil = []
        for r in runs:
            fc = r.find(qn('w:fldChar'))
            if fc is not None:
                t = fc.get(qn('w:fldCharType'))
                if t == 'begin':
                    durum.append('instr'); n += 1
                elif t == 'separate' and durum:
                    durum[-1] = 'sonuc'
                elif t == 'end' and durum:
                    durum.pop()
                sil.append(r)
                continue
            if durum and durum[-1] == 'instr':
                sil.append(r)
        for r in sil:
            r.getparent().remove(r)
    for f in list(body.iter(qn('w:fldSimple'))):
        par = f.getparent()
        for c in list(f):
            f.addprevious(c)
        par.remove(f)
        n += 1
    return n
