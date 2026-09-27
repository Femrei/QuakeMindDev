# -*- coding: utf-8 -*-
"""Revizyon 11: deney ortami -- tatbikat ikinci makinede yeniden kosuldu.

Tablo 11'e "Tatbikat yeniden kosumu" sutunu (AMD Ryzen 5 5500; surumler
benchmark/reports/tatbikat_b0/pip_freeze.txt), ilk sutun basligi gecikme
olcumune daraltildi; Bolum 4.1'deki ortam cumlesi duzeltildi.
"""
import copy
import shutil
import sys

sys.path.insert(0, r"C:\Users\utkuc\AppData\Local\Temp\claude\D--quakemind-final-QuakeMind\4779ac37-74b5-4aeb-8690-74d7792eedc5\scratchpad")
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm
from yardimci import Duzenleyici

SRC = r"D:\quakemind_final\QuakeMind\makale\arsiv\QuakeMind_Makale_YEDEK_pre_revize11_2026-09-28.docx"
DST = r"D:\quakemind_final\QuakeMind\makale\QuakeMind_Makale_Taslak_2026-09-16_sekillieklendi_v3_LOCAL_2026-09-21.docx"
shutil.copyfile(SRC, DST)
doc = Document(DST)
D = Duzenleyici(doc)

t = doc.tables[11]
tbl = t._tbl
assert t.rows[0].cells[1].text == "Tatbikat koşumu ve gecikme ölçümü"
# ikinci sutunun kopyasi ucuncu sutun olarak eklenir; birlesik not satirinin
# yayilimi bir artar
g = tbl.tblGrid.findall(qn("w:gridCol"))
g[1].addnext(copy.deepcopy(g[1]))
for tr in tbl.findall(qn("w:tr")):
    tcs = tr.findall(qn("w:tc"))
    if len(tcs) == 3:
        tcs[1].addnext(copy.deepcopy(tcs[1]))
    else:
        span = tcs[0].tcPr.find(qn("w:gridSpan"))
        span.set(qn("w:val"), str(int(span.get(qn("w:val"))) + 1))

yeni = {
    "Bileşen": ("Gecikme ve çıkarım süresi ölçümü (Tablo 14, 16)",
                "Tatbikat koşumu (Bölüm 4.6–4.9; düzeltilmiş kod, 28 Eylül 2026)"),
    "Ortam": (None, "Masaüstü bilgisayar (yerel)"),
    "İşlemci": (None, "AMD Ryzen 5 5500 (6 çekirdek, 12 iş parçacığı)"),
    "Grafik işlemci": (None, "Hesaplamada kullanılmadı (AMD Radeon RX 6600)"),
    "Sistem belleği": (None, "16 GB"),
    "Depolama": (None, "Kayıt altına alınmamıştır"),
    "İşletim sistemi": (None, "Windows 11 Pro, 64 bit (sürüm 10.0.26200)"),
    "Python / PyTorch": (None, "Python 3.12.10 / PyTorch 2.11.0 (yalnız CPU)"),
    "Temel kitaplıklar": (None, "NetworkX 3.6.1; OSMnx 2.1.0; pyrosm 0.13.1; segmentation_models_pytorch 0.5.0; "
                                "Ultralytics 8.4.57; Transformers 5.8.0"),
    "Kenar hızı v(e)": (None, "aynı"),
    "Kapanma cezası t_u": (None, "aynı"),
}
for row in t.rows:
    c = row.cells
    anahtar = c[0].text
    if anahtar in yeni:
        ilk, ikinci = yeni[anahtar]
        if ilk:
            D.hucre_yaz(c[1], ilk)
        D.hucre_yaz(c[2], ikinci)
    elif anahtar.startswith("Tatbikat ortamındaki sürümler"):
        D.hucre_yaz(c[0], "Tatbikat ve ölçüm ortamlarındaki sürümler, koşumların yürütüldüğü arka uç sanal "
                          "ortamlarından okunmuştur. Tatbikat sayıları benzetim içindeki ulaşım süreleridir; ikinci "
                          "makinede aynı model ağırlıklarıyla (SHA-256 doğrulandı) hasar oranı, kapalı kenar ve "
                          "kapsama değerleri ilk makinedekiyle birebir aynı çıkmıştır. Eğitim oturumlarında grafik "
                          "işlemci dışındaki bilgiler kaydedilmemiştir.")
    else:
        raise KeyError(anahtar)
for i, w in enumerate([2.5, 4.3, 4.3, 4.9]):
    t.columns[i].width = Cm(w)
    for c in t.columns[i].cells:
        c.width = Cm(w)
D.log.append("TABLO 11 ikinci makine sutunu")

D.degistir("Tatbikat koşumları ve Bölüm 4.5'teki gecikme ölçümleri, grafik işlemci hızlandırması kullanılmadan tek bir "
           "dizüstü bilgisayarda yürütülmüştür; çıkarım modelleri de bu koşumlarda merkezî işlemci üzerinde "
           "çalışmıştır.",
           "Bölüm 4.5'teki gecikme ölçümleri ve Tablo 16'daki çıkarım süreleri, grafik işlemci hızlandırması "
           "kullanılmadan bir dizüstü bilgisayarda (Intel Core i7-12700H) ölçülmüştür. Bölüm 4.6–4.9'daki tatbikat "
           "sayıları, tatbikat kodunda bulunan üç kusurun düzeltilmesinden sonra ikinci bir masaüstü bilgisayarda "
           "(AMD Ryzen 5 5500) yeniden üretilmiştir (Bölüm 5.3(b)); bu sayılar benzetim içindeki ulaşım "
           "süreleridir ve aynı model ağırlıklarıyla hasar oranı, kapalı kenar ve kapsama değerleri ilk "
           "makinedekiyle birebir aynı çıkmıştır. Her iki makinede de çıkarım modelleri merkezî işlemci üzerinde "
           "çalışmıştır.")
# Izgara tablolarinda puntosu tanimsiz kalan run'lar Normal stilinin buyuk
# puntosunu aliyordu (Tablo 2, 11, 22: "Kaggle Notebooks" vb.); tablonun
# baskin puntosuna esitlenir.
import collections
from docx.shared import Pt
from docx.text.run import Run
n = 0
for tb in doc.tables:
    if tb.style is None or tb.style.name != "Table Grid":
        continue
    runs = [Run(r, None) for r in tb._tbl.iter(qn("w:r"))]
    boy = collections.Counter(r.font.size.pt for r in runs if r.font.size and r.text.strip())
    if not boy:
        continue
    baskin = boy.most_common(1)[0][0]
    for r in runs:
        if r.font.size is None and r.text.strip():
            r.font.size = Pt(baskin)
            n += 1
D.log.append(f"PUNTO {n} run tablonun baskin puntosuna")

doc.save(DST)
print("\n".join(D.log))
