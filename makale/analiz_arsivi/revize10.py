# -*- coding: utf-8 -*-
"""Revizyon 10: tatbikat sayilari duzeltilmis kodla (J-112/113/114) yeniden
uretildi -- B0 paketi, 28 Eylul 2026, QuakeMindBackend/benchmark/reports/tatbikat_b0/.

Tablo degerleri JSON'dan okunur (elle yazim yok); metindeki bantlar asagida
verilerle sinanir (assert). Sekiller: Sekil 5, 10, 11 (a, b), 12 yeniden uretildi.
"""
import copy
import json
import shutil
import sys

sys.path.insert(0, r"C:\Users\utkuc\AppData\Local\Temp\claude\D--quakemind-final-QuakeMind\4779ac37-74b5-4aeb-8690-74d7792eedc5\scratchpad")
sys.path.insert(0, r"D:\quakemind_final\QuakeMind\QuakeMindBackend")
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph
from yardimci import Duzenleyici
from araclar_tatbikat_ozet import alan_ozeti

KOK = r"D:\quakemind_final\QuakeMind\QuakeMindBackend"
SON = KOK + r"\benchmark\reports\tatbikat_b0"
SEK = KOK + r"\benchmark\makale_sekilleri"
SRC = r"D:\quakemind_final\QuakeMind\makale\arsiv\QuakeMind_Makale_YEDEK_pre_revize10_2026-09-28.docx"
DST = r"D:\quakemind_final\QuakeMind\makale\QuakeMind_Makale_Taslak_2026-09-16_sekillieklendi_v3_LOCAL_2026-09-21.docx"


def oku(ad):
    return json.load(open(f"{SON}\\{ad}", encoding="utf-8"))


def v(x, b=2):
    return f"{x:.{b}f}".replace(".", ",")


def alan_kisa(a):
    return "Defne" if a.startswith("Defne") else a


# ------------------------------------------------------------------ veri
MODEL = {"segformer_mitb0_p2_t42": "B0", "segformer_mitb2_p2_t42": "B2", "unet_mitb4_p2_t42": "UNET"}
OZ = {}
for f in ("sonuc_b0.json", "sonuc_b2.json", "sonuc_unet.json"):
    for s in oku(f):
        r = alan_ozeti(s)
        r["_hasar"] = s.get("hasarGecerli")
        OZ[(alan_kisa(r["alan"]), MODEL[r["model"]])] = r

gecerli = [r for r in OZ.values() if r["gecerliTohum"]]
kaz = [r["kazanc"][0] for r in gecerli]
nur = [OZ[("Nurdağı", m)]["kazanc"][0] for m in ("B0", "B2", "UNET")]
pay = [r["cezaPayi"][0] for r in gecerli]
pay_nur = [OZ[("Nurdağı", m)]["cezaPayi"][0] for m in ("B0", "B2", "UNET")]
BANT = f"%{min(kaz):.0f}–{max(kaz):.0f}"
BANT_NUR = f"%{min(nur):.0f}–{max(nur):.0f}"
PAY = f"%{min(pay):.0f}–{max(pay):.0f}"
PAY_NUR = f"%{min(pay_nur):.0f}–{max(pay_nur):.0f}"
assert (BANT, BANT_NUR, PAY, PAY_NUR) == ("%8–60", "%54–60", "%68–78", "%73–76"), (BANT, BANT_NUR, PAY, PAY_NUR)
U = OZ[("Nurdağı", "UNET")]
assert (v(U["bilmeyen"][0]), v(U["telsiz"][0]), v(U["quakemind"][0])) == ("3,82", "3,24", "1,51")
I2, IU = OZ[("İslahiye", "B2")]["kazanc"][0], OZ[("İslahiye", "UNET")]["kazanc"][0]
assert (v(I2, 1), v(IU, 1)) == ("33,7", "8,3")

# ------------------------------------------------------------------ belge
shutil.copyfile(SRC, DST)
doc = Document(DST)
D = Duzenleyici(doc)
T = doc.tables


def hucre(c, metin):
    D.hucre_yaz(c, metin)


# ---- Kisaltmalar, Tablo 11, Tablo 22
for row in T[0].rows:
    if row.cells[0].text == "t_u":
        hucre(row.cells[1], "Kapanmaya varan ekibin keşif başına kaybettiği değerlendirme süresi (1,0 dk; geri "
                            "dönüş manevrası çizgede dönüş maliyeti olarak ayrıca ödenir)")
for row in T[11].rows:
    if row.cells[0].text.startswith("Kapanma cezası t_u"):
        hucre(row.cells[1], "değerlendirme 1,0 dk (manevra çizgede dönüş maliyetiyle); dönüş 0/3/10/25 sn; r_h = 0")
for row in T[24].rows:
    if row.cells[0].text.startswith("Görev tabanlı değerlendirme"):
        hucre(row.cells[3], f"Ölçülmüştür (Tablo 15): kazanç alan/modele göre {BANT}, Nurdağı {BANT_NUR}; "
                            "kapanma rotaya düşmezse fark yok")
D.log.append("TABLO kisaltma t_u, Tablo 11, Tablo 22")

# ---- Tablo 15 (12 sutun: kesif iki kola ayrildi, ceza payi eklendi)
t15 = T[15]
assert t15.rows[0].cells[0].text == "Alan" and len(t15.columns) == 11
# add_column bicimsiz bos hucre uretir; son sutunun kopyasi eklenir (yazi
# boyutu, hizalama ve hucre ozellikleri korunur)
_grid = t15._tbl.tblGrid
_grid.append(copy.deepcopy(_grid.findall(qn("w:gridCol"))[-1]))
for _tr in t15._tbl.findall(qn("w:tr")):
    _tr.append(copy.deepcopy(_tr.findall(qn("w:tc"))[-1]))
bas = ["Alan", "Model", "Hasar", "Kapalı k.", "Afetzede", "Bilmeyen", "Telsizle paylaşan", "QuakeMind",
       "Kazanç", "Keşif (bilmeyen / telsiz)", "Ceza payı", "Küme"]
for c, m in zip(t15.rows[0].cells, bas):
    hucre(c, m)
sira = [(a, m) for a in ("Defne", "İslahiye", "Nurdağı") for m in ("B0", "B2", "UNET")]
for row, key in zip(t15.rows[1:], sira):
    r = OZ[key]
    tek = r["kume"].startswith("1/")
    hucre(row.cells[2], f"%{v(100 * r['_hasar'])}")
    hucre(row.cells[3], str(r["kapaliKenar"]))
    hucre(row.cells[4], str(r["afetzede"]))
    if not r["gecerliTohum"]:
        for i, m in zip(range(5, 11), ["—", "—", "—", "araç rotası yok", "—", "—"]):
            hucre(row.cells[i], m)
    else:
        for i, k in zip((5, 6, 7), ("bilmeyen", "telsiz", "quakemind")):
            hucre(row.cells[i], f"{v(r[k][0])} ± {v(r[k][1])}")
        hucre(row.cells[8], f"%{v(r['kazanc'][0], 1)}" + ("" if tek else f" ± {v(r['kazanc'][1], 1)}"))
        hucre(row.cells[9], f"{v(r['kesifBilmeyen'][0], 1)} / {v(r['kesifTelsiz'][0], 1)}")
        hucre(row.cells[10], f"%{r['cezaPayi'][0]:.0f}" + ("" if tek else f" ± {r['cezaPayi'][1]:.0f}"))
    hucre(row.cells[11], r["kume"])
gen = [1.55, 1.55, 1.05, 0.95, 1.05, 1.55, 1.55, 1.55, 1.45, 1.5, 1.05, 0.8]
for i, w in enumerate(gen):
    t15.columns[i].width = Cm(w)
    for c in t15.columns[i].cells:
        c.width = Cm(w)
D.log.append("TABLO 15 yeniden (12 sutun)")

# ---- Tablo 16
t16 = T[16]
assert t16.rows[0].cells[0].text == "Model"
for row, m in zip(t16.rows[1:], ("B0", "B2", "UNET")):
    for i, a in zip((4, 5, 6), ("Nurdağı", "İslahiye", "Defne")):
        r = OZ[(a, m)]
        if not r["gecerliTohum"]:
            hucre(row.cells[i], "araç rotası yok")
        else:
            tek = r["kume"].startswith("1/")
            hucre(row.cells[i], f"%{v(r['kazanc'][0], 1)}" + ("" if tek else f" ± {v(r['kazanc'][1], 1)}"))
D.log.append("TABLO 16 kazanc")

# ---- Tablo 17 (guzergah varyantlari, bes ekip)
G = oku("guzergah_nurdagi_unet_t42.json")[0]["guzergahVaryantlari"]
g, o = G["gorevler"], G["olcum"]
t17 = T[18]
assert t17.rows[0].cells[0].text == "Görev" and len(t17.rows) == 6
m = lambda x: f"{x:.0f} m"
okunus = [
    f"kapanma rotayı değiştiriyor: metrede kısa ({m(g[0]['griM'])[:-2]} → {m(g[0]['maviM'])}), sürede uzun "
    f"({v(g[0]['griDk'], 1)} → {v(g[0]['maviDk'], 1)} dk)",
    f"kapanma {g[1]['maviM'] - g[1]['griM']:.0f} m dolaştırıyor ({v(g[1]['griDk'], 1)} → {v(g[1]['maviDk'], 1)} dk)",
    "yol kapalı değil ama sarı enkazdan uzak duruyor",
    f"kapanma {g[3]['maviM'] - g[3]['griM']:.0f} m dolaştırıyor ({v(g[3]['griDk'], 1)} → {v(g[3]['maviDk'], 1)} dk)",
    f"metrede kısa ({m(g[4]['griM'])[:-2]} → {m(g[4]['maviM'])}), sürede uzun "
    f"({v(g[4]['griDk'], 1)} → {v(g[4]['maviDk'], 1)} dk)",
]
assert g[0]["maviM"] < g[0]["griM"] and g[0]["maviDk"] > g[0]["griDk"]
assert abs(g[2]["maviM"] - g[2]["griM"]) < 0.5 and g[4]["maviDk"] > g[4]["griDk"]
satirlar = [[f"Ekip {i + 1}", m(x["griM"]), m(x["maviM"]), m(x["sariM"]), okunus[i]] for i, x in enumerate(g)]
for row, s in zip(t17.rows[1:4], satirlar[:3]):
    for c, t in zip(row.cells, s):
        hucre(c, t)
for c, t in zip(t17.rows[4].cells, satirlar[3]):
    hucre(c, t)
D.satir_ekle(t17, 4, satirlar[4])
son = t17.rows[-1]
assert son.cells[0].text.startswith("Ortalama uzama")
hucre(son.cells[2], f"{v(o['ortalamaUzamaM'], 1)} m (%{v(o['ortalamaUzamaYuzde'], 1)})")
hucre(son.cells[4], f"süre: {v(o['ortalamaUzamaDk'])} dk (%{v(o['ortalamaUzamaSureYuzde'], 1)}); "
                    f"en büyük {v(o['enBuyukUzamaDk'], 1)} dk")
assert o["uzamaHesabindakiCift"] == 25 and o["ulasilamayanCift"] == 0
# Okunusu sutunu dardi (bes esit sutun); toplam genislik ayni (16,0 cm)
for i, w in enumerate([2.4, 2.3, 2.3, 2.3, 6.7]):
    t17.columns[i].width = Cm(w)
    for c in t17.columns[i].cells:
        c.width = Cm(w)
D.log.append("TABLO 17 bes ekip")

# ---- Tablo 18 (r_h)
RH = {alan_kisa(s["alan"]): s["rhKiyasi"] for s in oku("sonuc_unet.json")}
t18 = T[19]
assert t18.rows[0].cells[1].text.startswith("r_h")
for i, row in enumerate(t18.rows[1:]):
    a = ("Defne", "İslahiye", "Nurdağı")[i // 4]
    x = RH[a][i % 4]
    assert float(row.cells[1].text) == x["yaricapM"]
    hucre(row.cells[2], str(x["kapaliKenar"]))
    hucre(row.cells[3], str(x["ortakGorev"]))
    hucre(row.cells[4], "—" if x.get("bilmeyenDk") is None else v(x["bilmeyenDk"]))
    hucre(row.cells[5], "—" if x.get("quakemindDk") is None else v(x["quakemindDk"]))
    hucre(row.cells[6], "—" if x.get("gorunurKazanc") is None else f"%{x['gorunurKazanc']:.0f}")
D.log.append("TABLO 18 r_h")

# ---- Tablo 19 (duyarlilik, U-Net tohum 42)
B_, Q_ = "Kapanmayı bilmeyen ekip", "QuakeMind"


def duy(ad):
    d = {}
    for s in oku(ad):
        k = s["kollarTohum"]["42"]["varyantlar"]
        b, q = k[B_]["ortalamaSureDk"], k[Q_]["ortalamaSureDk"]
        d[alan_kisa(s["alan"])] = (b, q, 100 * (b - q) / b)
    return d


DUY = [("ANA — değerlendirme 1,0 dk, dönüş maliyetli", duy("sonuc_unet.json")),
       ("Dönüş maliyeti kapalı (ceza 1,42 dk: değerlendirme + 25 sn manevra)", duy("duy_donussuz.json")),
       ("Değerlendirme 0,5 dk", duy("duy_deg_0.5.json")),
       ("Değerlendirme 2,0 dk", duy("duy_deg_2.0.json")),
       ("ESKİ SABİT — 20 dk (kullanılmıyor)", duy("duy_deg_20.json"))]
t19 = T[20]
assert t19.rows[0].cells[0].text == "Koşu"
for row, (ad, d) in zip(t19.rows[1:], DUY):
    hucre(row.cells[0], ad)
    for c, a in zip(row.cells[1:], ("Defne", "İslahiye", "Nurdağı")):
        b, q, k = d[a]
        hucre(c, f"{v(b)} → {v(q)} (%{k:.0f})")
ANA, DON, D05, D20, ESK = (x[1] for x in DUY)
assert [f"{ANA[a][2]:.0f}" for a in ("Defne", "İslahiye", "Nurdağı")] == ["18", "6", "59"]
assert [f"{DON[a][2]:.0f}" for a in ("Defne", "İslahiye", "Nurdağı")] == ["24", "9", "64"]
assert [f"{D05[a][2]:.0f}–{D20[a][2]:.0f}" for a in ("Defne", "İslahiye", "Nurdağı")] == ["12–28", "4–10", "47–71"]
artis = [100 * (ANA[a][1] / DON[a][1] - 1) for a in ("Defne", "İslahiye", "Nurdağı")]
assert f"%{min(artis):.0f}–{max(artis):.0f}" == "%16–28", artis
assert (v(ESK["Nurdağı"][0], 1), v(ESK["Nurdağı"][1], 1), f"{ESK['Nurdağı'][2]:.0f}") == ("34,1", "1,5", "96")
D.log.append("TABLO 19 duyarlilik")

# ---- Sekiller
for alt, dosya, cm in (("Şekil 5. Kapanma katmanının çizge", "sekil_tatbikat_nurdagi.png", None),
                       ("Şekil 10. Üç kanal tek haritada", "sekil_uc_kanal_nurdagi.png", None),
                       ("Şekil 12. QuakeMind kolunun kazancı", "sekil_kesif_kazanc.png", None)):
    p = D.basliktan_onceki_gorsel(alt)
    ext = p._p.find(".//" + qn("wp:extent"))
    genislik = int(ext.get("cx")) / 360000
    D.resim_koy(p, f"{SEK}\\{dosya}", genislik)
# Sekil 11: iki panel gorunmez duzen tablosunda
t_s11 = T[17]
for c, dosya in zip(t_s11.rows[0].cells, ("sekil_kiyas_kollari.png", "sekil_duyarlilik.png")):
    p = [x for x in c.paragraphs if x._p.findall(".//" + qn("w:drawing"))][0]
    ext = p._p.find(".//" + qn("wp:extent"))
    D.resim_koy(p, f"{SEK}\\{dosya}", int(ext.get("cx")) / 360000)

# ------------------------------------------------------------------ metin
R = D.degistir
# OZ
R("ölçülen bant alan ve modele göre %7–66, en sağlam alanda (Nurdağı, üç model × beş tohum) %59–66'dır",
  f"ölçülen bant alan ve modele göre {BANT}, en sağlam alanda (Nurdağı, üç model × beş tohum) {BANT_NUR}'tır")
R("Kapanmaya varan ekibin kaybettiği süre, önceki sürümdeki 20 dakikalık sabit yerine değerlendirme ve manevradan "
  "oluşan 1,42 dakikaya indirilmiştir; buna karşın kaba bir kestirime göre en sağlam alanda kazancın büyük bölümü "
  "hâlâ bu beyan edilmiş cezadan gelmektedir,",
  "Kapanmaya varan ekibin kaybettiği süre, önceki sürümdeki 20 dakikalık sabit yerine 1,0 dakikalık değerlendirme "
  "süresine indirilmiş, geri dönüş manevrası yol çizgesinde fiilen yapılan dönüş olarak fiyatlanmıştır; buna karşın "
  f"farkın {PAY}'i hâlâ bu beyan edilmiş süreden gelmektedir,")
# Abstract
R("the measured band is 7–66% depending on area and model, and 59–66% in the most robust area",
  "the measured band is 8–60% depending on area and model, and 54–60% in the most robust area")
R("The time lost by a team reaching a blockage was reduced from the previous 20-minute constant to 1.42 minutes "
  "(assessment plus manoeuvre); a rough estimate nevertheless attributes most of the gain in the most robust area "
  "to this declared penalty,",
  "The time lost by a team reaching a blockage was reduced from the previous 20-minute constant to a 1.0-minute "
  "assessment time, with the turn-back manoeuvre priced as the actual turn on the road graph; 68–78% of the "
  "difference nevertheless still comes from this declared time,")
# 3.8 cizge kurulumu
R("(Türkiye PBF'inden çıkarılan yaklaşık 756.623 kenar, 718.564 düğüm ve 205 tesis)",
  "(30 Mayıs 2026'da indirilen Türkiye PBF dosyasından çıkarılan 755.236 kenar, 717.341 düğüm ve 201 tesis; üç "
  "tatbikat alanında kapalı kenar ve kapsama değerleri önceki çıkarımla aynıdır)")
# 3.8 A* sezgiseli
R("Süre tabanlı aramada sezgisel h(n)/v_max olarak kullanılır (v_max = 22,2 m/s ≈ 80 km/sa). Bu sezgisel, "
  "çizgede hızı v_max'ı aşan kenar bulunmadığında kabul edilebilirdir ve A* eniyi çözümü korur; hız tablosunda "
  "otoyol sınıfı 90 km/sa olduğundan ve maxspeed etiketi 130 km/sa'ya kadar kabul edildiğinden bu koşul her "
  "çizgede kendiliğinden sağlanmaz, mahalle çizgelerinde 80 km/sa'yı aşan kenar bulunup bulunmadığı bu sürümde "
  "ayrıca raporlanmamıştır.",
  "Süre tabanlı aramada sezgisel h(n)/v_max olarak kullanılır; v_max çizgedeki en yüksek kenar hızıdır ve her "
  "çizge için ayrıca hesaplanır. Böylece sezgisel kalan süreyi hiçbir zaman aşmaz, yani kabul edilebilirdir ve "
  "A* eniyi çözümü korur. Taslağın önceki hâlinde sabit v_max = 80 km/sa kullanılıyordu; hız tablosunda otoyol "
  "sınıfı 90 km/sa olduğundan ve maxspeed etiketi 130 km/sa'ya kadar kabul edildiğinden bu sabit her çizgede "
  "kabul edilebilir değildi (yapay bir sınamada eniyi 41,5 sn'lik yol yerine 44,4 sn'lik yol seçiliyordu).")
R("ve tek yönlü kenarlara ters girilemez (bölgede 3.083 tek yönlü kenar). Bunlar beyan edilmiş tasarım "
  "parametreleridir, ölçüm değildir; dört sınama davranışı doğrulamaktadır (testler/test_donus.py).",
  "ve tek yönlü kenarlara ters girilemez (bölge dosyasında 43.821 tek yönlü kenar; oneway = −1 etiketli, yani "
  "çizim yönünün tersine tek yön olan kenarlar ters yönde kurulur, bu bölgede böyle kenar yoktur). Bunlar beyan "
  "edilmiş tasarım parametreleridir, ölçüm değildir; yedi sınama davranışı doğrulamaktadır (testler/test_donus.py).")
# Algoritma 1 notu ve ceza
R("c_kapanma = değerlendirme 1,0 dk + manevra 25 sn = 1,42 dk. Satır 3'teki yeniden planlama geri dönüşü ayrıca "
  "25 sn'lik U dönüşü olarak fiyatlar; manevra bu nedenle iki kez sayılmaktadır (Bölüm 5.3(b)).",
  "c_kapanma = değerlendirme 1,0 dk. Geri dönüş manevrası sabit olarak eklenmez: satır 3'teki yeniden planlama, "
  "ekibin fiilen yaptığı dönüşü (çoğunlukla 25 sn'lik U dönüşü) çizgenin dönüş maliyetiyle fiyatlar. Taslağın "
  "önceki hâlinde 25 sn'lik manevra ayrıca sabite de ekleniyor, yani iki kez sayılıyordu (Bölüm 5.3(b)). Dönüş "
  "maliyeti kapatıldığında (duyarlılık koşusu, Bölüm 4.9) manevra sabit olarak eklenir: c_kapanma = 1,0 dk + 25 sn.")
R("bu sürümde ceza ayrıştırılmış ve beyan edilmiştir: değerlendirme 1,0 dk + manevra 25 sn = 1,42 dk.",
  "bu sürümde ceza ayrıştırılmış ve beyan edilmiştir: değerlendirme 1,0 dk; geri dönüş manevrası çizgede fiilen "
  "yapılan dönüş olarak fiyatlanır.")
R("keşif başına beyan edilen cezanın (t_u = 1,42 dk) toplamıdır.",
  "keşif başına beyan edilen değerlendirme süresinin (t_u = 1,0 dk) toplamıdır.")
# 4.6
R("kazancın büyük bölümü keşif başına beyan edilen cezadan geldiğinden (aşağıda, Farkın ayrıştırılması), keşif "
  "sayısı ile kazanç arasındaki bağ ölçülen yol farkından çok bu cezayı yansıtır. Nurdağı'nda üç modelin üçü de beş "
  "farklı tohum kümesiyle koşabilmiş ve kazanç %59–66 bandında kalmıştır;",
  "farkın büyük bölümü keşif başına beyan edilen değerlendirme süresinden geldiğinden (aşağıda, Farkın "
  "ayrıştırılması), keşif sayısı ile kazanç arasındaki bağ ölçülen yol farkından çok bu süreyi yansıtır. "
  f"Nurdağı'nda üç modelin üçü de beş farklı tohum kümesiyle koşabilmiş ve kazanç {BANT_NUR} bandında kalmıştır;")
R("İslahiye'de dengeli model %2,06 hasarla %40,9 kazandırırken en iyi model %1,34 hasarla yalnız %9,9 kazandırmıştır.",
  f"İslahiye'de dengeli model %2,06 hasarla %{v(I2, 1)} kazandırırken en iyi model %1,34 hasarla yalnız "
  f"%{v(IU, 1)} kazandırmıştır.")
R("Keşif: ortak görevlerde yaşanan toplam kapanma keşfi (tohum ortalaması). Kol ortalamaları 0,1 dk'ya "
  "yuvarlandığından küçük farklarda (ör. İslahiye) kazanç birkaç puan belirsizdir.",
  "Keşif: ortak görevlerde yaşanan toplam kapanma keşfi, kapanmayı bilmeyen ekip / telsizle paylaşan filo (tohum "
  "ortalaması). Ceza payı: kapanmayı bilmeyen ekibin ödediği toplam değerlendirme süresinin iki kol arasındaki "
  "toplam süre farkına oranı. Küme 1/5 satırlarında beş tohum aynı kurulumu verdiğinden kazanç ve ceza payının "
  "standart sapması yazılmamıştır.")
R("ile kapanmaya her varışta yüklenen beyan edilmiş ceza (1,42 dk). Tablo 15'teki değerlerden yapılan kaba bir "
  "kestirime göre (keşif sayısının kapanmayı bilmeyen ekibe ait olduğu ve 25 ortak görev varsayımıyla) Nurdağı'nda "
  "(U-Net MiT-B4) cezanın toplam farktaki payı yaklaşık %83'tür (44,4 × 1,42 ≈ 62,9 dk; 25 × 3,04 ≈ 76,0 dk). Yani "
  "20 dakikalık sabit kaldırılmış olsa da kazancın büyük bölümü hâlâ kapanmaya varış anındaki beyan edilmiş cezadan "
  "gelmektedir; bu nedenle sonuç, cezanın değerine duyarlılığıyla birlikte bant olarak verilmiştir (Bölüm 4.9, "
  "Tablo 19). Kod her görev için cezadan gelen süreyi ayrıca yazmaktadır; alan ve model bazında kesin ayrıştırma bu "
  "alanla yeniden koşularak raporlanacaktır [ÖLÇÜLECEK].",
  "ile kapanmaya her varışta yüklenen beyan edilmiş değerlendirme süresi (1,0 dk). Tablo 15'in ceza payı sütunu bu "
  "ayrımı ölçer: her tohumda, kapanmayı bilmeyen ekibin ortak görevlerde ödediği toplam değerlendirme süresinin "
  f"iki kol arasındaki toplam süre farkına oranı. Pay {PAY} aralığındadır (Nurdağı'nda {PAY_NUR}); yani keşif "
  "başına yalnız 1,0 dakikalık bir değerlendirme süresiyle bile farkın büyük bölümü bu beyan edilmiş süreden, "
  "kalanı çizgede fiilen fazladan sürülen yol ve dönüşlerden gelmektedir. 20 dakikalık sabit kaldırılmış olsa da "
  "sonuç bu nedenle değerlendirme süresine duyarlılığıyla birlikte bant olarak verilmiştir (Bölüm 4.9, Tablo 19).")
R("(Nurdağı, en iyi model: kapanmayı bilmeyen 4,58 dk → telsizle paylaşan 3,80 dk → QuakeMind 1,54 dk).",
  f"(Nurdağı, en iyi model: kapanmayı bilmeyen {v(U['bilmeyen'][0])} dk → telsizle paylaşan "
  f"{v(U['telsiz'][0])} dk → QuakeMind {v(U['quakemind'][0])} dk); paylaşım keşif sayısını da düşürür "
  f"({v(U['kesifBilmeyen'][0], 1)} → {v(U['kesifTelsiz'][0], 1)}).")
R("(b) QuakeMind kolunun kazancının beyan edilen değerlendirme süresine (0,5–2,0 dk; toplam ceza 0,92–2,42 dk) "
  "duyarlılığı (segformer_mitb4_dice30, tohum 42; Tablo 19);",
  "(b) QuakeMind kolunun kazancının beyan edilen değerlendirme süresine (0,5–2,0 dk) duyarlılığı (U-Net MiT-B4, "
  "tohum 42; Tablo 19);")
# 4.7
R("ortalama uzama 101,7 m (%32,8) iken süre cinsinden 0,43 dk (%61,6), en büyük süre uzaması 1,8 dk'dır",
  f"ortalama uzama {v(o['ortalamaUzamaM'], 1)} m (%{v(o['ortalamaUzamaYuzde'], 1)}) iken süre cinsinden "
  f"{v(o['ortalamaUzamaDk'])} dk (%{v(o['ortalamaUzamaSureYuzde'], 1)}), en büyük süre uzaması "
  f"{v(o['enBuyukUzamaDk'], 1)} dk'dır")
R("Ekip 4 kayıtta raporlanmamıştır; uzama ortalaması yalnızca araçla tamamlanan rotalardan alınmıştır.",
  "Uzama ortalaması 25 ekip–afetzede çiftinin tamamından alınmıştır (hepsi araçla tamamlanmıştır).")
R("Nurdağı'nda 29–44 keşifle (tohum ortalaması) kazanç %59–66'ya çıkarken İslahiye'de en iyi modelle 1,8 keşifle "
  "%10'da kalmıştır.",
  f"Nurdağı'nda 29–44 keşifle (kapanmayı bilmeyen ekip, tohum ortalaması) kazanç {BANT_NUR}'a çıkarken "
  f"İslahiye'de en iyi modelle 1,8 keşifle %{IU:.0f}'de kalmıştır.")
R("(Tablo 15'in sekiz satırı; hata çubukları beş tohum üzerinden standart sapma).",
  "(Tablo 15'in sekiz satırı; yatay eksen kapanmayı bilmeyen ekibin keşif sayısı; hata çubukları beş tohum "
  "üzerinden standart sapma).")
R("Kazancın büyük bölümü keşif başına beyan edilen cezadan geldiğinden bu ilişki kısmen yapısaldır (Bölüm 4.6).",
  f"Farkın {PAY}'i keşif başına beyan edilen değerlendirme süresinden geldiğinden bu ilişki kısmen yapısaldır "
  "(Bölüm 4.6).")
# 4.8
rh = {a: RH[a] for a in RH}
k = lambda a, i: f"{rh[a][i]['gorunurKazanc']:.0f}"
og = lambda a, i: rh[a][i]["ortakGorev"]
R("r_h büyüdükçe 'kazanç' yükselir (%41→56, %9→59, %64→81) ama bu bir iyileşme değildir. Ortak görev sayısı çöker "
  "(6→3, 20→5, 25→7)",
  f"r_h büyüdükçe 'kazanç' yükselir (%{k('Defne', 0)}→{k('Defne', 1)}, %{k('İslahiye', 0)}→{k('İslahiye', 1)}, "
  f"%{k('Nurdağı', 0)}→{k('Nurdağı', 1)}) ama bu bir iyileşme değildir. Ortak görev sayısı çöker "
  f"({og('Defne', 0)}→{og('Defne', 1)}, {og('İslahiye', 0)}→{og('İslahiye', 1)}, "
  f"{og('Nurdağı', 0)}→{og('Nurdağı', 1)})")
# 4.9
R("Bu koşu Tablo 15'teki üç modelden biriyle değil, platformun varsayılan işletim modeliyle (segformer_mitb4_dice30; "
  "karar eşiği 0,4, son işleme seviye 2) yapılmıştır; bu nedenle ANA satırındaki süreler Tablo 15'teki değerlerle "
  "aynı değildir ve tablo mutlak süreleri değil, sabitlerin göreli etkisini okumak için verilmiştir. Bu modelin "
  "sınama IoU'su künyesinde kayıtlı değildir.",
  "Bu koşu Tablo 15'teki en iyi modelle (U-Net MiT-B4) yapılmıştır; ANA satırı Tablo 15'in tohum 42'deki "
  "değeridir (Tablo 15 beş tohumun ortalamasını verdiğinden İslahiye ve Nurdağı'nda sayılar küçük farklarla "
  "ayrışır).")
R("Tablo 19. Beyan edilen ceza süresine ve dönüş maliyetine duyarlılık (segformer_mitb4_dice30, tohum 42)",
  "Tablo 19. Beyan edilen değerlendirme süresine ve dönüş maliyetine duyarlılık (U-Net MiT-B4, tohum 42)")
R("Dönüş maliyeti (sağdan trafik kuralı) mutlak süreleri Defne ve Nurdağı'nda yaklaşık %30 artırır; göreli kazancı "
  "bu iki alanda neredeyse değiştirmez (%28↔%29, %58↔%56), İslahiye'de ise %13'ten %27'ye çıkarır;",
  "Dönüş maliyeti (sağdan trafik kuralı) QuakeMind kolunun mutlak sürelerini %16–28 artırır; dönüş maliyeti "
  "kapatıldığında ceza, değerlendirme ile sabit bir manevradan oluştuğundan (1,42 dk) göreli kazanç üç alanda da "
  f"birkaç puan yükselir (%{ANA['Defne'][2]:.0f}→{DON['Defne'][2]:.0f}, %{ANA['İslahiye'][2]:.0f}→"
  f"{DON['İslahiye'][2]:.0f}, %{ANA['Nurdağı'][2]:.0f}→{DON['Nurdağı'][2]:.0f});")
R("Değerlendirme süresi 0,5–2,0 dk aralığında (toplam ceza 0,92–2,42 dk) ölçülü bir etki yapar (kazanç Defne %24–41, "
  "İslahiye %13–19, Nurdağı %50–67)",
  "Değerlendirme süresi 0,5–2,0 dk aralığında ölçülü bir etki yapar (kazanç Defne %12–28, İslahiye %4–10, "
  "Nurdağı %47–71)")
R("aynı alan ve kapanmalarla Nurdağı '22,6 dk → 1,4 dk, %94 iyileşme' diye okunur.",
  "aynı alan ve kapanmalarla Nurdağı '34,1 dk → 1,5 dk, %96 iyileşme' diye okunur.")
# 5.1
R("ölçülen kazanç alan ve modele göre %7–66, en sağlam alanda (Nurdağı, üç model × beş tohum) %59–66'dır;",
  f"ölçülen kazanç alan ve modele göre {BANT}, en sağlam alanda (Nurdağı, üç model × beş tohum) {BANT_NUR}'tır;")
R("Kapanmaya varan ekibin kaybettiği süre önceki sürümdeki 20 dakikalık sabitten 1,42 dakikaya indirilmiştir; buna "
  "karşın kaba kestirime göre Nurdağı'nda kazancın yaklaşık %83'ü hâlâ bu beyan edilmiş cezadan gelmektedir "
  "(Bölüm 4.6),",
  "Kapanmaya varan ekibin kaybettiği süre önceki sürümdeki 20 dakikalık sabitten 1,0 dakikalık değerlendirme "
  "süresine indirilmiş, geri dönüş manevrası çizgede fiilen yapılan dönüş olarak fiyatlanmıştır; buna karşın "
  f"ölçülen ceza payına göre farkın {PAY}'i hâlâ bu beyan edilmiş süreden gelmektedir (Bölüm 4.6),")
R("(Nurdağı, en iyi model: 4,58 dk → 3,80 dk).",
  f"(Nurdağı, en iyi model: {v(U['bilmeyen'][0])} dk → {v(U['telsiz'][0])} dk).")
R("yaklaşık dört kat fazla kazanç sağlamıştır (%40,9'a karşı %9,9).",
  f"yaklaşık dört kat fazla kazanç sağlamıştır (%{v(I2, 1)}'e karşı %{v(IU, 1)}).")
# 5.3(b)
D.yaz("(b) Kapanmaya varan ekibin kaybettiği süre bir model parametresidir;",
      "(b) Kapanmaya varan ekibin kaybettiği süre bir model parametresidir; bu sürümde değerlendirme 1,0 dk olarak "
      "beyan edilmiş, geri dönüş manevrası çizgede fiilen yapılan dönüş olarak fiyatlanmıştır. Ölçülen ceza payına "
      f"göre farkın {PAY}'i bu beyan edilmiş süreden gelmektedir (Bölüm 4.6); değerlendirme süresine duyarlılık "
      "Bölüm 4.9'da taranmış ve alanların sıralamasının değişmediği gösterilmiştir, yine de mutlak süreler ve "
      "kazanç yüzdesi bu beyan edilmiş sabite bağlıdır. Taslağın önceki hâlinde tatbikat kodunda üç kusur bulunmuş "
      "ve düzeltilmiştir: kapanmada geri dönen ekibin manevrası iki kez sayılıyor, OSM'deki oneway = −1 etiketi "
      "ileri yön olarak yorumlanıyor ve A* sezgiseli sabit bir hız üst sınırı kullanıyordu (oneway = −1 etiketli "
      "kenar bu bölgenin verisinde bulunmadığından ikinci kusurun sonuca etkisi yoktur). Bu bölümdeki bütün "
      "tatbikat sayıları düzeltilmiş kodla yeniden üretilmiştir; düzeltme kazançları düşürmüş (Nurdağı bandı "
      f"%59–66'dan {BANT_NUR}'a, Defne'de %41'den %18–28'e), orta boy modelin İslahiye'deki üstünlüğünü "
      "değiştirmemiştir. Telsizle paylaşan filo kolu, tek başına öğrenen ekibe göre kazancın bir bölümünü yakalar.")
# 6
R("kapanmanın ekip rotalarına düştüğü alanlarda erişim süresi %7–66 kısalmış, en sağlam alanda (Nurdağı) %59–66 "
  "kazanç ölçülmüştür;",
  f"kapanmanın ekip rotalarına düştüğü alanlarda erişim süresi {BANT} kısalmış, en sağlam alanda (Nurdağı) "
  f"{BANT_NUR} kazanç ölçülmüştür;")
R("Kaba kestirime göre bu kazancın büyük bölümü keşif başına beyan edilen cezadan geldiğinden sonuç,",
  f"Ölçülen ceza payına göre farkın {PAY}'i keşif başına beyan edilen değerlendirme süresinden geldiğinden sonuç,")

doc.save(DST)
print("\n".join(D.log))
